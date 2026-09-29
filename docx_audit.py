"""Inspect generated DOCX structure without an inverse Word-to-Markdown parser.

Changelog (term variables):
    - Compare tagged term values with each other and with the generation
      snapshot (`terms`); mismatches and missing controls are warnings only.
    - Serialize reports explicitly so documents without terms keep their JSON.

Changelog (hardening):
    - Index styles once per audit and separate render validation from observations.
    - Treat placeholder-like text as warnings and validate native numbering references.
"""

import argparse
import json
import logging
import re
from dataclasses import asdict, dataclass, field
from typing import Any, Dict, List, Optional, Tuple, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.parts.numbering import NumberingPart
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph
from lxml import etree

from term_variables import TERM_PROPERTY_PREFIX, TERM_TAG_PREFIX

logger = logging.getLogger(__name__)


@dataclass
class DocumentAudit:
    """Structural observations and actionable errors, not a visual QA verdict."""

    paragraphs: int = 0
    tables: int = 0
    headings: int = 0
    images: int = 0
    numbered_paragraphs: int = 0
    issues: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    pagination_marked_paragraphs: int = 0
    visual_review: str = "not_performed"


@dataclass
class TermAudit:
    """Consistency of tagged term values after editing, not their financial validity.

    Attributes:
        mismatched: Key -> distinct current values in document order, when the
            key's controls disagree (for example, only one place was edited).
        changed: Key -> generated snapshot value and the first current value
            that differs from it, for copying back into the `terms:` YAML.
        missing: Snapshot keys none of whose controls remain in the document.
        indicative: Keys with a current value containing a bracketed `[...]` part.
    """

    mismatched: Dict[str, List[str]] = field(default_factory=dict)
    changed: Dict[str, Dict[str, str]] = field(default_factory=dict)
    missing: List[str] = field(default_factory=list)
    indicative: List[str] = field(default_factory=list)


# Text nodes of a control's current value; None means the node's own text.
_TERM_TEXT: Dict[str, Optional[str]] = {
    qn("w:t"): None, qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n",
}
_REMOVED_CONTENT = frozenset({qn("w:del"), qn("w:moveFrom")})
_TERM_TAG_PATH = qn("w:sdtPr") + "/" + qn("w:tag")
_PLACEHOLDER_PATH = qn("w:sdtPr") + "/" + qn("w:showingPlcHdr")
_INDICATIVE_RE = re.compile(r"\[[^\[\]]*\]")


_PAGINATION_FLAGS = (
    ("keep_together", "keepLines"),
    ("keep_with_next", "keepNext"),
    ("page_break_before", "pageBreakBefore"),
)


class _StyleLookup:
    """Audit-local style index and inherited pagination cache."""

    def __init__(self, doc: DocxDocument) -> None:
        self.styles = {style.style_id: style for style in doc.styles}
        self.default = next((
            style for style in self.styles.values()
            if style.type == WD_STYLE_TYPE.PARAGRAPH
            and style.element.get(qn("w:default")) in {"1", "true", "on"}
        ), None)
        self.defaults: Dict[str, bool] = {}
        for _, tag in _PAGINATION_FLAGS:
            nodes = doc.styles.element.xpath("./w:docDefaults/w:pPrDefault/w:pPr/w:" + tag)
            self.defaults[tag] = bool(
                nodes and nodes[0].get(qn("w:val"), "1") not in {"0", "false", "off"}
            )
        self.flags: Dict[Tuple[Optional[str], str], bool] = {}

    def paragraph_style(self, paragraph: Paragraph) -> Optional[ParagraphStyle]:
        """Resolve a paragraph style, including the default for unknown IDs."""
        style = self.styles.get(paragraph._p.style)
        if style is None or style.type != WD_STYLE_TYPE.PARAGRAPH:
            style = self.default
        return style

    def inherited_flag(self, paragraph: Paragraph, attribute: str, tag: str) -> bool:
        """Resolve each style's ancestry only once for each pagination control."""
        style = self.paragraph_style(paragraph)
        key = (style.style_id if style is not None else None, attribute)
        if key not in self.flags:
            value = None
            visited = set()
            while value is None and style is not None and style.style_id not in visited:
                visited.add(style.style_id)
                value = getattr(style.paragraph_format, attribute)
                style = self.styles.get(style.element.basedOn_val)
            self.flags[key] = bool(value) if value is not None else self.defaults[tag]
        return self.flags[key]


def _effective_pagination_flag(
    paragraph: Paragraph, attribute: str, xml_tag: str, styles: _StyleLookup,
) -> bool:
    """Resolve local > paragraph style ancestry > document default controls."""
    value = getattr(paragraph.paragraph_format, attribute)
    return bool(value) if value is not None else styles.inherited_flag(paragraph, attribute, xml_tag)


def inspect_document(doc: DocxDocument) -> DocumentAudit:
    """Inspect structure and observations, including text inside tables.

    Args:
        doc: Open python-docx document to inspect without modifying it.

    Returns:
        Structural issues and descriptive counts, warnings and review status.
        Inconsistent or missing term values are warnings, never issues.
    """
    styles = _StyleLookup(doc)
    paragraph_styles = (styles.paragraph_style(p) for p in doc.paragraphs)
    result = DocumentAudit(
        paragraphs=len(doc.paragraphs),
        tables=len(doc.tables),
        headings=sum(style is not None and style.name.startswith("Heading ") for style in paragraph_styles),
        images=len(doc.inline_shapes),
        numbered_paragraphs=len(doc.element.xpath(".//w:pPr/w:numPr")),
        issues=inspect_document_issues(doc),
    )
    for element in doc.element.xpath(".//w:p"):
        paragraph = Paragraph(element, doc)
        if any(
            _effective_pagination_flag(paragraph, attribute, tag, styles)
            for attribute, tag in _PAGINATION_FLAGS
        ):
            result.pagination_marked_paragraphs += 1
    normal = doc.styles["Normal"].paragraph_format
    if normal.keep_together or normal.keep_with_next or normal.page_break_before:
        result.warnings.append(
            "Normal style has pagination constraints: these can show nonprinting square "
            "marks and create large page gaps. Keep such controls on necessary headings only."
        )
    # Text alone cannot distinguish renderer fallback markers from quoted user
    # examples. Real rendering failures are recorded by the renderer itself.
    for text in doc.element.xpath(".//w:t/text()"):
        if "[Render Error:" in text or text.startswith(("[Image:", "[Diagram:")):
            result.warnings.append("Possible rendering placeholder (may be literal user text): " + text)
    terms = inspect_terms(doc)
    if terms is not None:
        result.warnings.extend(_term_warnings(terms))
    return result


def inspect_terms(doc: DocxDocument) -> Optional[TermAudit]:
    """Compare tagged term values with each other and with the generation snapshot.

    Only the tagged values are compared; numbers typed outside a control and
    the financial validity of the terms are not checked.

    Args:
        doc: Open python-docx document to inspect without modifying it.

    Returns:
        Term results when the body has term controls or the package has a term
        snapshot; None otherwise, so a report without terms is unchanged.
    """
    current = _current_term_values(doc)
    generated = _generated_term_values(doc)
    if not current and not generated:
        return None
    result = TermAudit()
    for key, values in sorted(current.items()):
        distinct = list(dict.fromkeys(values))
        if len(distinct) > 1:
            result.mismatched[key] = distinct
        if any(_INDICATIVE_RE.search(value) for value in values):
            result.indicative.append(key)
    for key, value in sorted(generated.items()):
        if key not in current:
            result.missing.append(key)
            continue
        differing = [text for text in current[key] if text != value]
        if differing:
            result.changed[key] = {"generated": value, "current": differing[0]}
    return result


def _current_term_values(doc: DocxDocument) -> Dict[str, List[str]]:
    """Read each remaining term control's current text in document order."""
    values: Dict[str, List[str]] = {}
    for control in doc.element.body.iter(qn("w:sdt")):
        tag = control.find(_TERM_TAG_PATH)
        name = tag.get(qn("w:val"), "") if tag is not None else ""
        if not name.startswith(TERM_TAG_PREFIX) or _is_removed(control):
            continue
        content = control.find(qn("w:sdtContent"))
        text = ""
        # Word refills a cleared control with placeholder text and flags it.
        if content is not None and not _shows_placeholder(control):
            text = _current_text(content, control)
        values.setdefault(name[len(TERM_TAG_PREFIX):], []).append(text)
    return values


def _shows_placeholder(control: Any) -> bool:
    """Whether a content control displays placeholder text instead of a value."""
    flag = control.find(_PLACEHOLDER_PATH)
    return flag is not None and flag.get(qn("w:val"), "true") not in {"0", "false", "off"}


def _current_text(content: Any, control: Any) -> str:
    """Join text, tabs and breaks, including insertions but not deletions or moves away."""
    parts: List[str] = []
    for node in content.iter(*_TERM_TEXT):
        if _is_removed(node, stop=control):
            continue
        replacement = _TERM_TEXT[node.tag]
        parts.append((node.text or "") if replacement is None else replacement)
    return "".join(parts)


def _is_removed(node: Any, stop: Any = None) -> bool:
    """Whether a deletion or move-away ancestor (below `stop`) hides the node."""
    for ancestor in node.iterancestors():
        if ancestor is stop:
            return False
        if ancestor.tag in _REMOVED_CONTENT:
            return True
    return False


def _generated_term_values(doc: DocxDocument) -> Dict[str, str]:
    """Read the `ibrep.term.` snapshot written when the document was generated."""
    try:
        part = doc.part.package.part_related_by(RT.CUSTOM_PROPERTIES)
    except KeyError:
        return {}
    values: Dict[str, str] = {}
    for prop in etree.fromstring(part.blob):
        if not isinstance(prop.tag, str):
            continue  # comments and processing instructions
        name = prop.get("name") or ""
        if name.startswith(TERM_PROPERTY_PREFIX):
            values[name[len(TERM_PROPERTY_PREFIX):]] = "".join(prop.itertext())
    return values


def _term_warnings(terms: TermAudit) -> List[str]:
    """Human-readable warnings for inconsistent and missing term values."""
    warnings = [
        f"Term {key!r} has inconsistent values in the document: "
        + ", ".join(repr(value) for value in values)
        for key, values in terms.mismatched.items()
    ]
    warnings.extend(
        f"Term {key!r} was generated but none of its content controls remain in the document"
        for key in terms.missing
    )
    return warnings


def audit_to_dict(result: DocumentAudit, terms: Optional[TermAudit] = None) -> Dict[str, Any]:
    """Serialize an audit report; documents without terms keep the previous schema.

    Args:
        result: Structural audit from inspect_document.
        terms: Term results from inspect_terms, if any.

    Returns:
        The audit fields, plus `terms` only when term results exist.
    """
    data = asdict(result)
    if terms is not None:
        data["terms"] = asdict(terms)
    return data


def inspect_document_issues(doc: DocxDocument) -> List[str]:
    """Validate structure without computing descriptive audit observations.

    Args:
        doc: Open python-docx document to validate before saving.

    Returns:
        Structural errors suitable for strict render validation.
    """
    issues: List[str] = []
    for index, section in enumerate(doc.sections, 1):
        width, left, right = section.page_width, section.left_margin, section.right_margin
        if width is not None and left is not None and right is not None and width <= left + right:
            issues.append(f"Section {index} has no printable width")
    references = doc.element.xpath(".//w:footnoteReference")
    if references:
        try:
            part = doc.part.part_related_by(RT.FOOTNOTES)
            notes = etree.fromstring(part.blob)
            ids = {node.get(qn("w:id")) for node in notes}
            for reference in references:
                if reference.get(qn("w:id")) not in ids:
                    issues.append("Footnote reference has no definition")
        except KeyError:
            issues.append("Footnote definitions part is missing")
    for hyperlink in doc.element.xpath(".//w:hyperlink"):
        relation = hyperlink.get(qn("r:id"))
        if relation and relation not in doc.part.rels:
            issues.append("Hyperlink relationship is missing")
    issues.extend(_numbering_issues(doc))
    return issues


def _numbering_issues(doc: DocxDocument) -> List[str]:
    """Validate numbering definitions and numPr references across XML parts."""
    issues: List[str] = []
    definitions = {}
    try:
        numbering = cast(NumberingPart, doc.part.part_related_by(RT.NUMBERING)).element
    except KeyError:
        numbering = None
    if numbering is not None:
        abstracts = {node.get(qn("w:abstractNumId")) for node in numbering.findall(qn("w:abstractNum"))}
        definitions = {node.get(qn("w:numId")): node for node in numbering.findall(qn("w:num"))}
        for num_id, node in definitions.items():
            abstract = node.find(qn("w:abstractNumId"))
            abstract_id = abstract.get(qn("w:val")) if abstract is not None else None
            if abstract_id is None or abstract_id not in abstracts:
                issues.append(f"Numbering numId {num_id} has no abstract definition ({abstract_id})")
    for part in doc.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            if not part.content_type.endswith("+xml"):
                continue
            root = etree.fromstring(part.blob)
        for properties in root.iter(qn("w:numPr")):
            for reference in properties.findall(qn("w:numId")):
                num_id = reference.get(qn("w:val"))
                # Word uses zero to explicitly suppress inherited numbering.
                if num_id == "0":
                    continue
                if numbering is None:
                    issue = "Numbering definitions part is missing"
                elif num_id not in definitions:
                    issue = f"Numbering reference numId {num_id} has no definition"
                else:
                    continue
                if issue not in issues:
                    issues.append(issue)
    return issues


def audit_file(path: str) -> DocumentAudit:
    """Open an output file and inspect its structure."""
    return inspect_document(Document(path))


def main() -> None:
    """Print a JSON audit result and return nonzero on structural errors only."""
    parser = argparse.ArgumentParser(description="Inspect generated Word document structure")
    parser.add_argument("input_file")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    try:
        document = Document(args.input_file)
        result = inspect_document(document)
        terms = inspect_terms(document)
    except Exception as exc:
        logger.error("Cannot inspect DOCX: %s", exc)
        raise SystemExit(1) from exc
    logger.info(json.dumps(audit_to_dict(result, terms), ensure_ascii=False, indent=2))
    raise SystemExit(1 if result.issues else 0)


if __name__ == "__main__":
    main()
