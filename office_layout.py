"""Office document composition and editable Word numbering.

Changelog (hardening):
    - Scope native list instances to their parent item and override the actual level.
    - Reuse memo title/metadata typography for cover-free IB reports, adding subtitles.
"""

from typing import Dict, Optional, Tuple

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from document_model import DocumentMetadata, DocumentModel, ElementType, Heading
from document_profiles import string_list
from render_styles import STYLE


def add_text(doc, text: str, bold: bool = False, keep_next: bool = False, style_name: Optional[str] = None):
    """Add a consistently styled office paragraph."""
    paragraph = doc.add_paragraph(style=style_name or STYLE.STYLE_IB_BODY)
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(text)
    run.font.name = STYLE.BODY_FONT
    if style_name is None:
        run.font.size = STYLE.BODY_SIZE
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), STYLE.KOREAN_FONT)
    return paragraph


def setup_letter_styles(doc) -> None:
    """Create reusable letter styles without blanket pagination constraints."""
    for name, size, before, after in [
        ("Office Metadata", STYLE.BODY_SIZE.pt, 0, 2),
        ("Office Subject", STYLE.BODY_SIZE.pt, 0, 16),
        ("Office Attachment", 10, 0, 1),
        ("Office Signatory", 15, 18, 12),
        ("Office Contact", 9, 0, 1),
    ]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[STYLE.STYLE_IB_BODY]
        style.font.name, style.font.size = STYLE.BODY_FONT, Pt(size)
        style.font.color.rgb = STYLE.NAVY
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), STYLE.KOREAN_FONT)
        fmt = style.paragraph_format
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.keep_together, fmt.keep_with_next = False, False
        fmt.tab_stops.add_tab_stop(Pt(42))
        if name in {"Office Metadata", "Office Subject", "Office Contact"}:
            fmt.left_indent, fmt.first_line_indent = Pt(42), Pt(-42)


def letter_appendix_index(model: DocumentModel) -> Optional[int]:
    """Resolve a unique explicit H1 boundary before composing any output."""
    if model.metadata.profile != "office-letter":
        return None
    heading = model.metadata.extra.get("letter", {}).get("appendix_heading")
    if heading is None:
        return None
    matches = [
        index for index, element in enumerate(model.elements)
        if element.element_type == ElementType.HEADING_1
        and isinstance(element.content, Heading) and element.content.text == heading
    ]
    if len(matches) != 1 or heading == model.metadata.title:
        raise ValueError("letter.appendix_heading must match one distinct body H1")
    if not any(
        element.element_type not in {ElementType.HEADING_1, ElementType.HEADING_2, ElementType.HEADING_3}
        for element in model.elements[:matches[0]]
    ):
        raise ValueError("letter.appendix_heading must follow the letter body")
    return matches[0]


def render_office_opening(doc, metadata: DocumentMetadata) -> bool:
    """Render profile metadata at the start of the body when there is no cover.

    Args:
        doc: Document with request-scoped styles already initialized.
        metadata: Profile metadata, including an optional IB report subtitle.

    Returns:
        Whether a title was inserted and a matching body H1 should be skipped.
    """
    name = metadata.profile
    if name == "plain" or (name == "ib-report" and not metadata.title.strip()):
        return False
    extra = metadata.extra
    sender = extra.get("sender", {})
    if name == "office-letter":
        organization = add_text(doc, sender["organization"], bold=True, keep_next=True)
        organization.style = doc.styles["Title"]
        organization.alignment = WD_ALIGN_PARAGRAPH.CENTER
        organization.runs[0].font.size = Pt(18)
        organization.paragraph_format.space_after = Pt(24)
        for label, key in [("수신", "recipients"), ("참조", "cc")]:
            values = string_list(extra, key)
            if values:
                add_text(doc, "{}\t{}".format(label, ", ".join(values)), keep_next=True, style_name="Office Metadata")
        add_text(doc, "제목\t" + metadata.title, bold=True, keep_next=True, style_name="Office Subject")
        return True
    title = add_text(doc, metadata.title, bold=True, keep_next=True)
    title.style = doc.styles["Title"]
    title.runs[0].font.size = STYLE.H1_SIZE
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(12)
    if name == "ib-report" and metadata.subtitle.strip():
        add_text(doc, metadata.subtitle, keep_next=True)
    rows = []
    for label, value in [
        ("작성일", extra.get("date")),
        ("작성부서", sender.get("department")),
        ("작성자", metadata.analyst),
        ("일시", extra.get("meeting_time")),
        ("장소", extra.get("location")),
    ]:
        if value:
            rows.append(f"{label}: {value}")
    attendees = string_list(extra, "attendees")
    if attendees:
        rows.append("참석자: " + ", ".join(attendees))
    for row in rows:
        add_text(doc, row, keep_next=True)
    return True


def render_office_closing(doc, metadata: DocumentMetadata) -> None:
    """Render attachments and sender without manufacturing a signature."""
    if metadata.profile != "office-letter":
        return
    attachments = string_list(metadata.extra, "attachments")
    if attachments:
        label = add_text(doc, "붙임", style_name="Office Attachment", keep_next=True)
        label.paragraph_format.space_before = Pt(8)
        numbering = NativeNumbering(doc)
        for index, item in enumerate(attachments, 1):
            ending = "  끝." if index == len(attachments) else ""
            paragraph = add_text(doc, item.rstrip().rstrip(".") + "." + ending, style_name="Office Attachment")
            numbering.apply(paragraph, 0)
            paragraph.paragraph_format.left_indent = Pt(30)
            paragraph.paragraph_format.first_line_indent = Pt(-14)
            paragraph.paragraph_format.tab_stops.clear_all()
            paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(30))
    else:
        # Attach the closing marker to the last actual body paragraph when possible.
        last = doc.paragraphs[-1] if doc.paragraphs else None
        if last is not None and last._p.getnext() is not None and last._p.getnext().tag == qn("w:sectPr"):
            if not last.text.rstrip().endswith("끝."):
                last.add_run("  끝.")
        else:
            add_text(doc, "끝.")
    sender = metadata.extra["sender"]
    paragraph = add_text(doc, sender.get("signatory") or sender["organization"], bold=True, style_name="Office Signatory")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    issued = " ".join(str(metadata.extra[k]) for k in ("document_no", "date") if metadata.extra.get(k))
    if issued:
        add_text(doc, "시행\t" + issued, style_name="Office Contact")
    contact = " · ".join(sender[k] for k in ("department", "contact") if sender.get(k))
    if contact:
        add_text(doc, "담당\t" + contact, style_name="Office Contact")


class NativeNumbering:
    """One numbering state per document; supports Korean multilevel lists."""

    def __init__(self, doc) -> None:
        self.doc = doc
        self.ids: Dict[Tuple[int, str], int] = {}

    def reset(self) -> None:
        """Start a fresh list after a non-list block."""
        self.ids.clear()

    def apply(self, paragraph, level: int, bullet: bool = False, start: int = 1) -> None:
        """Attach native numbering, restarting descendants under each new parent.

        Args:
            paragraph: Destination Word paragraph.
            level: Zero-based nesting level, clamped to Word's supported range.
            bullet: Whether the list uses bullets instead of numbers.
            start: First value for a newly encountered ordered-list scope.
        """
        level = min(max(level, 0), 8)
        # A new item at this level ends any child list owned by the previous item.
        self.ids = {key: value for key, value in self.ids.items() if key[0] <= level}
        kind = "bullet" if bullet else "number"
        key = (level, kind)
        if key not in self.ids:
            self.ids[key] = self._create(bullet, start, level)
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = level
        num_pr.get_or_add_numId().val = self.ids[key]

    def _create(self, bullet: bool, start: int, start_level: int) -> int:
        numbering = self.doc.part.numbering_part.element
        abstract_ids = [
            int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
        abstract_id = max(abstract_ids + [-1]) + 1
        num_id = max(num_ids + [0]) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(9):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            number_format = "bullet" if bullet else ("ganada" if level % 3 == 1 else "decimal")
            label = "•" if bullet else ("(%{})" if level % 3 == 2 else "%{}.").format(level + 1)
            for tag, value in [
                ("start", "1"),
                ("numFmt", number_format),
                ("lvlText", label),
                ("lvlJc", "left"),
            ]:
                element = OxmlElement("w:" + tag)
                element.set(qn("w:val"), value)
                lvl.append(element)
            p_pr = OxmlElement("w:pPr")
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), str(360 * (level + 1)))
            indent.set(qn("w:hanging"), "360")
            p_pr.append(indent)
            lvl.append(p_pr)
            abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            numbering.insert(list(numbering).index(first_num), abstract)
        else:
            numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        if not bullet and (start != 1 or start_level > 0):
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), str(start_level))
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), str(start))
            override.append(start_override)
            num.append(override)
        numbering.append(num)
        return num_id
