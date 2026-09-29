"""Term-sheet profile composition: house boilerplate, opening block and layout helpers.

Called from `IBDocumentRenderer.render` (the single composition path) when the
resolved profile is `term-sheet`. Design: docs/term-sheet-design-20260929.md.
This module depends only on the model, the inline `md_parser.TextParser`, styles,
YAML and python-docx primitives; it must not import `ib_renderer` (the renderer
injects its run-rendering callback, so run output stays identical everywhere).

Changelog (A1 foundation):
    - Validate immutable house boilerplate with presence-based frontmatter precedence.

Changelog (house style):
    - NEW: `style:` house/frontmatter options (cover page, logo, boxed disclaimer,
      header/footer rules, label colour, page number format, dark header,
      open-sided tables); the defaults keep the standard layout.

Changelog (A2 rendering):
    - Schema-ordered single cell fills shared with generic merged-table emission.
    - Per-line runs with marker hanging indents; row-split estimation; fixed label
      grid, table frame, caption/unit line, note, source and 4pt spacer.
    - Opening block, per-line body paragraphs, accent headings, document defaults,
      two-row confirmation box and per-section header/footer.
"""

import base64
import math
import re
import unicodedata
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple, Union

import yaml
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tc
from docx.shared import Emu, Length, Mm, Pt, RGBColor, Twips

from document_model import DocumentMetadata, Heading, Paragraph, TextRun
from md_parser import TextParser
from render_styles import STYLE

RenderRuns = Callable[..., None]
"""Injected `TextRenderer.render_runs(paragraph, runs, default_color=, font_name=, font_size=)`."""

ROW_SPLIT_THRESHOLD = 12
CELL_MARGIN_VERTICAL = 70  # twips (3.5pt) above and below cell text
CELL_MARGIN_HORIZONTAL = 100  # twips (5pt) left and right of cell text
_BORDER_SIZE = "4"  # eighths of a point: 0.5pt
_BORDER_EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")
_CELL_SPACE_AFTER = Pt(1.5)
_CELL_LINE_SPACING = 1.05
_SPACER_HEIGHT = Pt(4)
_CONFIRMATION_GAP = Pt(6)  # air between a preceding note and the confirmation box
_BOLD_ALLOWANCE = 1.1  # bold header glyphs are wider than the em estimate
# Successor elements used to insert children in ECMA-376 schema order.
_TCPR_AFTER_SHD = (
    "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
    "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
)
_TBLPR_AFTER_IND = (
    "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
    "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)
_TBLPR_AFTER_MARGINS = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
_TRPR_AFTER_CANT_SPLIT = (
    "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del",
    "w:trPrChange",
)
_PPR_AFTER_KINSOKU = (
    "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", *_PPR_AFTER_KINSOKU)
_RPR_AFTER_SZ_CS = (
    "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl",
    "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
)
_HEADING_RULE_SIZE = "8"  # 1pt accent rule under Heading 2
_CHECKBOX = "□"
_CHECKBOX_LAYOUT_MM = 4.5  # hanging indent of confirmation check items
_COVER_TITLE_SPACE = Mm(60)  # sets the cover title block about a third down the page
_COVER_DISCLAIMER_SPACE = Mm(30)
_HOUSE_KEYS = frozenset({"prepared_by", "disclaimer", "confidential_label", "confirmation", "style"})
_CONFIRMATION_KEYS = frozenset({"intro", "items", "signature"})
_STYLE_CHOICES = {
    "cover": ("inline", "page"),
    "disclaimer": ("rules", "box"),
    "page_number_align": ("right", "center"),
    "table_header": ("light", "dark"),
    "table_sides": ("closed", "open"),
}
_STYLE_FLAGS = ("header_rule", "footer_rule")
_HEX_COLOR_RE = re.compile(r"#[0-9A-Fa-f]{6}")
_PAGE_FIELD_RE = re.compile(r"(\{page\}|\{pages\})")


@dataclass(frozen=True)
class ConfirmationText:
    """Validated confirmation wording without mutable item collections."""

    intro: str = ""
    items: Tuple[str, ...] = ()
    signature: str = ""


@dataclass(frozen=True)
class TermSheetStyle:
    """House layout choices (`style:`); the defaults are the profile's standard layout.

    Attributes:
        cover: `inline` opens the first page with the title block; `page` gives
            the title block, logo and disclaimer a cover page of their own.
        logo: Absolute image path or `data:` URI, centred in the opening.
        logo_width_mm: Logo width.
        disclaimer: `rules` (thin rules above and below) or `box` (a frame).
        header_rule: Accent rule under the header.
        footer_rule: Accent rule over the footer.
        label_color: `#RRGGBB` of the confidentiality label; empty keeps grey.
        page_number: Footer page text with `{page}` and optional `{pages}`.
        page_number_align: `right` or `center` in the footer.
        table_header: `light` (tinted, accent text) or `dark` (accent fill, white text).
        table_sides: `closed` draws the outer left and right borders; `open` omits them.
    """

    cover: str = "inline"
    logo: str = ""
    logo_width_mm: float = 40.0
    disclaimer: str = "rules"
    header_rule: bool = False
    footer_rule: bool = False
    label_color: str = ""
    page_number: str = "{page} / {pages}"
    page_number_align: str = "right"
    table_header: str = "light"
    table_sides: str = "closed"


@dataclass(frozen=True)
class TermSheetTexts:
    """Resolved boilerplate for one render; empty labels intentionally remain empty."""

    prepared_by: str
    disclaimer: str
    confidential_label: str = "Strictly Confidential"
    confirmation: Optional[ConfirmationText] = None
    style: TermSheetStyle = TermSheetStyle()


def validate_style(value: Any, base_dir: Optional[Path] = None) -> Dict[str, Any]:
    """Validate a `style:` mapping and resolve a relative logo path.

    Args:
        value: The mapping from a house file or frontmatter.
        base_dir: Folder a relative `logo` path is resolved against (the house
            file's or the Markdown file's); None keeps it as written.

    Returns:
        The supplied settings, with `logo` made absolute when possible.

    Raises:
        ValueError: A key is unknown or a value is malformed.
    """
    if not isinstance(value, dict):
        raise ValueError("style must be a mapping")
    known = set(_STYLE_CHOICES) | set(_STYLE_FLAGS) | {"logo", "logo_width_mm", "label_color", "page_number"}
    unknown = set(value) - known
    if unknown:
        raise ValueError("Unknown style settings: " + ", ".join(sorted(map(str, unknown))))
    settings = dict(value)
    for key, choices in _STYLE_CHOICES.items():
        if key in settings and settings[key] not in choices:
            raise ValueError(f"style.{key} must be one of: " + ", ".join(choices))
    for key in _STYLE_FLAGS:
        if key in settings and not isinstance(settings[key], bool):
            raise ValueError(f"style.{key} must be true or false")
    width = settings.get("logo_width_mm")
    if width is not None and (isinstance(width, bool) or not isinstance(width, (int, float)) or not 5 <= width <= 150):
        raise ValueError("style.logo_width_mm must be a number from 5 to 150")
    color = settings.get("label_color")
    if color is not None and (not isinstance(color, str) or not _HEX_COLOR_RE.fullmatch(color)):
        raise ValueError("style.label_color must be #RRGGBB")
    page = settings.get("page_number")
    if page is not None and (
        not isinstance(page, str)
        or "{page}" not in page
        or re.sub(r"\{page\}|\{pages\}", "", page).count("{")
        or "\n" in page
    ):
        raise ValueError("style.page_number must be one line with {page} and optional {pages}")
    logo = settings.get("logo")
    if logo is not None:
        if not isinstance(logo, str) or not logo.strip():
            raise ValueError("style.logo must be an image path or a data: URI")
        if not logo.lower().startswith("data:") and base_dir is not None and not Path(logo).is_absolute():
            settings["logo"] = str((base_dir / logo).resolve())
    return settings


def _confirmation_text(value: Any) -> ConfirmationText:
    """Convert a supported confirmation mapping into an immutable payload."""
    if not isinstance(value, dict) or not value:
        raise ValueError("confirmation must be a mapping with intro, items or signature")
    unknown = set(value) - _CONFIRMATION_KEYS
    if unknown:
        raise ValueError("Unknown confirmation settings: " + ", ".join(sorted(map(str, unknown))))
    for key in ("intro", "signature"):
        if key in value and not isinstance(value[key], str):
            raise ValueError(f"confirmation.{key} must be a string")
    if "items" in value:
        items = value["items"]
        if (
            not isinstance(items, list)
            or not items
            or any(not isinstance(item, str) or not item.strip() for item in items)
        ):
            raise ValueError("confirmation.items must be a non-empty list of non-empty strings")
    return ConfirmationText(
        intro=value.get("intro", ""),
        items=tuple(value.get("items", [])),
        signature=value.get("signature", ""),
    )


def _validate_house_fields(data: Dict[str, Any]) -> None:
    """Reject unsupported keys and malformed values without requiring all fields."""
    unknown = set(data) - _HOUSE_KEYS
    if unknown:
        raise ValueError("Unknown house settings: " + ", ".join(sorted(map(str, unknown))))
    for key in ("prepared_by", "disclaimer", "confidential_label"):
        if key in data and not isinstance(data[key], str):
            raise ValueError(f"{key} must be a string")
    if "confirmation" in data:
        _confirmation_text(data["confirmation"])
    if "style" in data:
        validate_style(data["style"])


def load_house(path: Union[str, Path]) -> Dict[str, Any]:
    """Load and validate the fields supplied by a house YAML file.

    Args:
        path: House file to read; callers resolve its source-relative location.

    Returns:
        Validated supplied fields, retaining which optional keys were present.

    Raises:
        ValueError: YAML syntax, mapping shape or a supplied field is invalid.
        OSError: The selected file cannot be read.
    """
    try:
        data = yaml.safe_load(Path(path).read_text(encoding="utf-8-sig"))
    except yaml.YAMLError as exc:
        raise ValueError(f"Invalid house YAML: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("house must be a YAML mapping")
    _validate_house_fields(data)
    if "style" in data:
        # A house logo path is relative to the house file.
        data["style"] = validate_style(data["style"], Path(path).resolve().parent)
    return data


def resolve_term_sheet_texts(
    metadata: DocumentMetadata, house_path: Optional[str] = None
) -> TermSheetTexts:
    """Resolve boilerplate before rendering, preserving explicit empty values.

    Args:
        metadata: Term-sheet metadata containing optional boilerplate overrides.
        house_path: Effective house path after caller/frontmatter precedence.
            Source-relative paths must already be absolute; absent paths use no file.

    Returns:
        Immutable texts for the renderer's current request.

    Raises:
        ValueError: A path is relative, a field is malformed, or required text is absent.
        OSError: The effective house file cannot be read.
    """
    house: Dict[str, Any] = {}
    if house_path is not None:
        if not isinstance(house_path, str) or not house_path.strip():
            raise ValueError("house must be a YAML file path")
        if not Path(house_path).is_absolute():
            raise ValueError("house path must be absolute for input without a source file")
        house = load_house(house_path)
    values = {
        key: metadata.extra[key] if key in metadata.extra else house[key]
        for key in _HOUSE_KEYS - {"style"}
        if key in metadata.extra or key in house
    }
    _validate_house_fields(values)
    for key in ("prepared_by", "disclaimer"):
        if key not in values or not values[key].strip():
            raise ValueError(f"term-sheet requires non-empty {key}")
    # Style settings merge key by key: frontmatter `style` overrides the house file's.
    style = dict(house.get("style", {}))
    if "style" in metadata.extra:
        style.update(validate_style(metadata.extra["style"]))
    logo = style.get("logo")
    if logo and not logo.lower().startswith("data:") and not Path(logo).is_absolute():
        raise ValueError("style.logo must be absolute for input without a source file")
    return TermSheetTexts(
        prepared_by=values["prepared_by"],
        disclaimer=values["disclaimer"],
        confidential_label=values.get("confidential_label", "Strictly Confidential"),
        confirmation=_confirmation_text(values["confirmation"])
        if "confirmation" in values else None,
        style=TermSheetStyle(**style),
    )


# ═══════════════════════════════════════════════════════════════════════════════
# TABLE CELL PRIMITIVES
# ═══════════════════════════════════════════════════════════════════════════════


def set_cell_fill(tc: CT_Tc, hex_color: str) -> None:
    """Give one `w:tc` exactly one solid fill, placed in schema order.

    Args:
        tc: The cell element; a covered vertical-merge cell is a distinct `w:tc`.
        hex_color: Six-digit RGB fill without a leading `#`.
    """
    tc_pr = tc.get_or_add_tcPr()
    for shading in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(shading)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    tc_pr.insert_element_before(shading, *_TCPR_AFTER_SHD)


# ═══════════════════════════════════════════════════════════════════════════════
# LINES AND LIST-LIKE MARKERS
# ═══════════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class MarkerLayout:
    """Hanging indent for a line that starts with a list-like marker (plan §2-3).

    Word receives `left = start + hanging` and `first line = -hanging`, so a
    wrapped line aligns with the first character after the marker.
    """

    start_mm: float
    hanging_mm: float
    note: bool = False  # ※ lines also use the note text size


_MARKER_LAYOUTS = {
    "•": MarkerLayout(0.0, 3.0),
    "-": MarkerLayout(3.0, 3.0),
    "·": MarkerLayout(6.0, 3.0),
    "※": MarkerLayout(0.0, 4.0, note=True),
}
_CIRCLED_LAYOUT = MarkerLayout(0.0, 4.5)
_CIRCLED_NUMBERS = frozenset(chr(code) for code in range(0x2460, 0x2474))  # ①–⑳


def marker_layout(text: str) -> Optional[MarkerLayout]:
    """Return the layout for a line's leading marker; the text is never changed.

    A hyphen counts only when followed by whitespace, so negative amounts and a
    lone dash placeholder keep their ordinary layout.
    """
    stripped = text.lstrip()
    if not stripped:
        return None
    marker = stripped[0]
    if marker in _CIRCLED_NUMBERS:
        return _CIRCLED_LAYOUT
    if marker == "-" and not stripped[1:2].isspace():
        return None
    return _MARKER_LAYOUTS.get(marker)


def apply_marker_layout(paragraph: Any, text: str) -> Optional[Pt]:
    """Give a paragraph its marker hanging indent.

    Args:
        paragraph: python-docx paragraph that will hold exactly this line.
        text: The line's visible text, used only to detect the marker.

    Returns:
        The note size for `※` lines, otherwise None (keep the caller's size).
    """
    layout = marker_layout(text)
    if layout is None:
        return None
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Mm(layout.start_mm + layout.hanging_mm)
    paragraph_format.first_line_indent = Mm(-layout.hanging_mm)
    return STYLE.TS_NOTE_SIZE if layout.note else None


def _is_plain(run: TextRun) -> bool:
    """Semantic runs (equations, footnotes, term values, code, images) are never trimmed or split."""
    return (
        not run.is_latex and run.footnote_id is None and run.term_key is None and not run.code
        and run.image is None
    )


def _strip_edge(line: List[TextRun], leading: bool) -> List[TextRun]:
    """Drop spaces and tabs beside a line break, as HTML does around `<br>`."""
    runs = list(line)
    index = 0 if leading else -1
    while runs and _is_plain(runs[index]):
        text = runs[index].text.lstrip(" \t") if leading else runs[index].text.rstrip(" \t")
        if text:
            runs[index] = replace(runs[index], text=text)
            break
        runs.pop(index)
    return runs


def split_run_lines(runs: Sequence[TextRun]) -> List[List[TextRun]]:
    """Split runs at line breaks into one run list per Word paragraph.

    Formatting, link targets, footnote references and term keys stay on every
    piece, and empty lines are kept as empty lists. Only spaces and tabs next to
    a break are dropped; the source runs are never mutated.

    Args:
        runs: Parsed runs whose text may contain newlines from `<br>` or hard breaks.

    Returns:
        At least one, possibly empty, run list.
    """
    lines: List[List[TextRun]] = [[]]
    for run in runs:
        if not _is_plain(run) or "\n" not in run.text:
            lines[-1].append(run)
            continue
        for index, piece in enumerate(run.text.split("\n")):
            if index:
                lines.append([])
            if piece:
                lines[-1].append(replace(run, text=piece))
    trimmed: List[List[TextRun]] = []
    for index, line in enumerate(lines):
        if index > 0:
            line = _strip_edge(line, leading=True)
        if index < len(lines) - 1:
            line = _strip_edge(line, leading=False)
        trimmed.append(line)
    return trimmed


def line_text(line: Sequence[TextRun]) -> str:
    """Visible text of one split line."""
    return "".join(run.text for run in line)


def _em_width(text: str) -> float:
    """Approximate advance in ems: East Asian wide/full-width 1.0, others 0.55."""
    return sum(1.0 if unicodedata.east_asian_width(char) in ("W", "F") else 0.55 for char in text)


def header_token_width(text: str, font_size: int) -> int:
    """Column width in EMU that keeps a header's longest unbreakable token on one line.

    Tokens split at whitespace (including `<br>` line breaks). The estimate uses
    the same em widths as the row estimator, a 10% allowance for bold glyphs and
    the left and right cell margins.

    Args:
        text: Visible header text of one single-column header cell.
        font_size: Header text size in EMU.

    Returns:
        Required width in EMU, or 0 for an empty header.
    """
    tokens = text.split()
    if not tokens:
        return 0
    widest = max(_em_width(token) for token in tokens)
    return math.ceil(widest * int(font_size) * _BOLD_ALLOWANCE) + 2 * int(Twips(CELL_MARGIN_HORIZONTAL))


def natural_cell_width(texts: Sequence[str], font_size: int, bold: bool = False, markers: bool = False) -> int:
    """Column width in EMU that keeps every line of a cell on one line.

    Uses the row estimator's em widths, the bold allowance for header text,
    marker hanging indents (content cells) and the left and right cell margins.

    Args:
        texts: One entry per cell paragraph, already split at line breaks.
        font_size: Text size in EMU.
        bold: Whether the text is bold (header cells).
        markers: Whether marker lines are indented (content cells).

    Returns:
        Required width in EMU, or 0 for an empty cell.
    """
    widest = 0.0
    for text in texts:
        if not text.strip():
            continue
        size, indent = int(font_size), 0
        layout = marker_layout(text) if markers else None
        if layout is not None:
            indent = int(Mm(layout.start_mm + layout.hanging_mm))
            if layout.note:
                size = int(STYLE.TS_NOTE_SIZE)
        advance = _em_width(text) * size * (_BOLD_ALLOWANCE if bold else 1.0)
        widest = max(widest, advance + indent)
    if not widest:
        return 0
    return math.ceil(widest) + 2 * int(Twips(CELL_MARGIN_HORIZONTAL))


def estimate_cell_lines(texts: Sequence[str], width: int, font_size: int, markers: bool) -> int:
    """Estimate one cell's wrapped line count for row pagination (plan §4 A2).

    Args:
        texts: One entry per cell paragraph, already split at line breaks.
        width: Merged cell width in EMU, including the cell margins.
        font_size: Body text size of the cell in EMU.
        markers: Whether marker lines are indented (content cells).

    Returns:
        Sum over paragraphs of ceil(text advance / inner width), at least one each.
    """
    inner = int(width) - 2 * int(Twips(CELL_MARGIN_HORIZONTAL))
    total = 0
    for text in texts:
        size, indent = int(font_size), 0
        layout = marker_layout(text) if markers else None
        if layout is not None:
            indent = int(Mm(layout.start_mm + layout.hanging_mm))
            if layout.note:
                size = int(STYLE.TS_NOTE_SIZE)
        available = max(inner - indent, 1)
        total += max(1, math.ceil(_em_width(text) * size / available))
    return total


# ═══════════════════════════════════════════════════════════════════════════════
# TABLE LAYOUT
# ═══════════════════════════════════════════════════════════════════════════════


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def label_widths(label_columns: int) -> List[int]:
    """Fixed label column widths in EMU: the label, plus the second label tier."""
    widths = [int(STYLE.TS_LABEL_WIDTH)]
    if label_columns == 2:
        widths.append(int(STYLE.TS_SUBLABEL_WIDTH))
    return widths


def key_value_widths(label_columns: int, available: int) -> Optional[List[int]]:
    """Fixed label grid for key-value tables, so their first vertical line is common.

    The labels keep their fixed widths whenever the content column keeps a
    positive width, however narrow the section is.

    Args:
        label_columns: One or two leading label columns.
        available: Printable width of the table's own section in EMU (portrait
            or landscape); the content column receives the remainder.

    Returns:
        Column widths in EMU, or None when the fixed labels leave no content width
        (the caller reports that explicitly).
    """
    fixed = label_widths(label_columns)
    remaining = int(available) - sum(fixed)
    return fixed + [remaining] if remaining > 0 else None


def _measure(tag: str, twips: int) -> Any:
    element = OxmlElement(tag)
    element.set(qn("w:w"), str(twips))
    element.set(qn("w:type"), "dxa")
    return element


def apply_table_frame(word_table: Any, width: int, open_sides: bool = False) -> None:
    """Apply the term-sheet table frame in schema order.

    The grid is fixed and spans `width`. In compatibility mode 14 Word draws the
    left border at the table indent minus the left cell margin, so an indent equal
    to that margin puts both outer borders on the text margins. The border edges
    are 0.5pt `TS_BORDER_HEX`, except the outer left and right edges of an
    open-sided table (`style.table_sides: open`); cell margins are 70
    (top/bottom) and 100 (left/right) twips.

    Args:
        word_table: python-docx table whose grid widths are already set.
        width: Total grid width in EMU.
        open_sides: Omit the outer left and right borders.
    """
    word_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    word_table.autofit = False
    tbl_pr = word_table._tbl.tblPr
    for tag in ("w:tblInd", "w:tblBorders", "w:tblCellMar"):
        for existing in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(existing)
    table_width = tbl_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.insert_element_before(table_width, "w:jc", "w:tblCellSpacing", "w:tblInd", *_TBLPR_AFTER_IND)
    table_width.set(qn("w:w"), str(Emu(width).twips))
    table_width.set(qn("w:type"), "dxa")
    tbl_pr.insert_element_before(_measure("w:tblInd", CELL_MARGIN_HORIZONTAL), *_TBLPR_AFTER_IND)
    borders = OxmlElement("w:tblBorders")
    for edge in _BORDER_EDGES:
        border = OxmlElement(f"w:{edge}")
        if open_sides and edge in ("left", "right"):
            border.set(qn("w:val"), "nil")
            borders.append(border)
            continue
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), _BORDER_SIZE)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), STYLE.TS_BORDER_HEX)
        borders.append(border)
    tbl_pr.insert_element_before(borders, *_TBLPR_AFTER_IND[1:])
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (
        ("top", CELL_MARGIN_VERTICAL), ("left", CELL_MARGIN_HORIZONTAL),
        ("bottom", CELL_MARGIN_VERTICAL), ("right", CELL_MARGIN_HORIZONTAL),
    ):
        margins.append(_measure(f"w:{edge}", value))
    tbl_pr.insert_element_before(margins, *_TBLPR_AFTER_MARGINS)


def configure_cell_paragraph(paragraph: Any) -> None:
    """Cell paragraph rhythm: no space before, 1.5pt after, 1.05 line spacing."""
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = _CELL_SPACE_AFTER
    paragraph_format.line_spacing = _CELL_LINE_SPACING


def set_row_pagination(word_row: Any, keep_together: bool, repeat_header: bool) -> None:
    """Add `w:cantSplit`/`w:tblHeader` to a row in schema order, only when needed."""
    if not keep_together and not repeat_header:
        return
    tr_pr = word_row._tr.get_or_add_trPr()
    if keep_together:
        tr_pr.insert_element_before(OxmlElement("w:cantSplit"), *_TRPR_AFTER_CANT_SPLIT)
    if repeat_header:
        tr_pr.insert_element_before(OxmlElement("w:tblHeader"), *_TRPR_AFTER_CANT_SPLIT[2:])


def add_table_heading(
    doc: Any, caption: str, unit: str, as_of: str, render_runs: RenderRuns, width: int
) -> None:
    """Caption at left and `(단위 : …, 기준일 : …)` at right, kept with the table.

    Args:
        doc: Document receiving the paragraph before the table.
        caption: Bold accent-coloured caption text; may be empty.
        unit: Unit label; may be empty.
        as_of: Reference date text; may be empty.
        render_runs: Injected run renderer.
        width: Printable width of the current section in EMU (right tab stop).
    """
    details = ", ".join(
        part for part in (f"단위 : {unit}" if unit else "", f"기준일 : {as_of}" if as_of else "") if part
    )
    if not caption and not details:
        return
    paragraph = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
    paragraph.paragraph_format.keep_with_next = True
    if caption:
        runs = [replace(run, bold=True) for run in TextParser.parse_runs(caption)]
        render_runs(paragraph, runs, default_color=STYLE.NAVY, font_size=STYLE.BODY_SIZE)
    if details:
        if caption:
            paragraph.paragraph_format.tab_stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)
            render_runs(paragraph, [TextRun(text="\t")], font_size=STYLE.TS_NOTE_SIZE)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        render_runs(
            paragraph, TextParser.parse_runs(f"({details})"),
            default_color=_rgb(STYLE.TS_MUTED_HEX), font_size=STYLE.TS_NOTE_SIZE,
        )


def add_table_note(doc: Any, note: str, render_runs: RenderRuns) -> None:
    """Right-aligned muted note directly below the table."""
    paragraph = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    render_runs(
        paragraph, TextParser.parse_runs(note),
        default_color=_rgb(STYLE.TS_MUTED_HEX), font_size=STYLE.TS_NOTE_SIZE,
    )


def add_table_source(doc: Any, source: str, render_runs: RenderRuns) -> None:
    """Muted `출처 : …` line below the table (and its note)."""
    paragraph = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
    render_runs(
        paragraph, TextParser.parse_runs("출처 : " + source),
        default_color=_rgb(STYLE.TS_MUTED_HEX), font_size=STYLE.TS_NOTE_SIZE,
    )


def _add_spacer(doc: Any, height: Length, keep_with_next: bool = False) -> None:
    """Empty paragraph of an exact height (no body-size line)."""
    paragraph_format = doc.add_paragraph().paragraph_format
    if keep_with_next:
        paragraph_format.keep_with_next = True
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = height


def add_table_spacer(doc: Any) -> None:
    """Empty paragraph of exactly 4pt after a table instead of a body-size line."""
    _add_spacer(doc, _SPACER_HEIGHT)


# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT STYLES
# ═══════════════════════════════════════════════════════════════════════════════


def _default_properties(doc: Any) -> Tuple[Any, Any]:
    """Return `docDefaults` paragraph and run properties, creating them if absent."""
    root = doc.styles.element
    defaults = root.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        root.insert(0, defaults)
    run_default = defaults.find(qn("w:rPrDefault"))
    if run_default is None:
        run_default = OxmlElement("w:rPrDefault")
        defaults.insert(0, run_default)
    paragraph_default = defaults.find(qn("w:pPrDefault"))
    if paragraph_default is None:
        paragraph_default = OxmlElement("w:pPrDefault")
        run_default.addnext(paragraph_default)
    for parent, tag in ((run_default, "w:rPr"), (paragraph_default, "w:pPr")):
        if parent.find(qn(tag)) is None:
            parent.append(OxmlElement(tag))
    return paragraph_default.find(qn("w:pPr")), run_default.find(qn("w:rPr"))


def _set_child(parent: Any, tag: str, successors: Sequence[str], **attributes: str) -> Any:
    """Replace one optional child, inserted before its schema successors."""
    for existing in parent.findall(qn(tag)):
        parent.remove(existing)
    element = OxmlElement(tag)
    for name, value in attributes.items():
        element.set(qn(f"w:{name}"), value)
    parent.insert_element_before(element, *successors)
    return element


def setup_term_sheet_styles(doc: Any) -> None:
    """Apply term-sheet document defaults and heading accents (term-sheet only).

    Korean text wraps at word (eojeol) boundaries with kinsoku rules: Word breaks
    East Asian text per character when `w:wordWrap` is 0, so it is set to 1
    explicitly. Automatic spacing between East Asian text and Latin letters or
    numbers (`w:autoSpaceDE`/`w:autoSpaceDN`) is off, so `300억원`, `SPC에` and
    `36개월` print tight as in Korean term sheets. The East Asian language is
    `ko-KR`. Unformatted paragraph marks (table cells, spacers) use the body font
    and size, so Word does not size their lines at the 11pt template default.
    Heading 2 gets the accent colour and a 1pt accent rule; Heading 3 the accent.

    Args:
        doc: Document whose request-scoped styles were created already.
    """
    paragraph_defaults, run_defaults = _default_properties(doc)
    _set_child(paragraph_defaults, "w:kinsoku", _PPR_AFTER_KINSOKU, val="1")
    _set_child(paragraph_defaults, "w:wordWrap", _PPR_AFTER_KINSOKU[1:], val="1")
    _set_child(paragraph_defaults, "w:autoSpaceDE", _PPR_AFTER_KINSOKU[4:], val="0")
    _set_child(paragraph_defaults, "w:autoSpaceDN", _PPR_AFTER_KINSOKU[5:], val="0")
    fonts = run_defaults.get_or_add_rFonts()
    for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme"):
        fonts.attrib.pop(qn(f"w:{theme}"), None)
    fonts.set(qn("w:ascii"), STYLE.BODY_FONT)
    fonts.set(qn("w:hAnsi"), STYLE.BODY_FONT)
    fonts.set(qn("w:eastAsia"), STYLE.KOREAN_FONT)
    run_defaults.sz_val = STYLE.BODY_SIZE
    half_points = str(int(round(STYLE.BODY_SIZE.pt * 2)))
    _set_child(run_defaults, "w:szCs", _RPR_AFTER_SZ_CS, val=half_points)
    language = run_defaults.find(qn("w:lang"))
    if language is None:
        language = _set_child(run_defaults, "w:lang", _RPR_AFTER_SZ_CS[11:])
    language.set(qn("w:eastAsia"), "ko-KR")

    heading2 = doc.styles["Heading 2"]
    heading2.font.color.rgb = STYLE.NAVY
    rule = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for name, value in (("val", "single"), ("sz", _HEADING_RULE_SIZE), ("space", "1"), ("color", STYLE.NAVY_HEX)):
        bottom.set(qn(f"w:{name}"), value)
    rule.append(bottom)
    style_paragraph = heading2.element.get_or_add_pPr()
    for existing in style_paragraph.findall(qn("w:pBdr")):
        style_paragraph.remove(existing)
    style_paragraph.insert_element_before(rule, *_PPR_AFTER_PBDR)
    doc.styles["Heading 3"].font.color.rgb = STYLE.NAVY


# ═══════════════════════════════════════════════════════════════════════════════
# OPENING BLOCK, HEADINGS AND BODY PARAGRAPHS
# ═══════════════════════════════════════════════════════════════════════════════


def _display_runs(metadata: DocumentMetadata, key: str, text: str) -> List[TextRun]:
    """Parse-time runs (term substitution) when present, else the parsed text."""
    runs = metadata.display_runs.get(key)
    return list(runs) if runs else TextParser.parse_runs(text)


def _bold(runs: Sequence[TextRun]) -> List[TextRun]:
    return [replace(run, bold=True) for run in runs]


def _opening_paragraph(doc: Any, space_after: Length, keep_with_next: bool = True) -> Any:
    paragraph = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = space_after
    paragraph_format.keep_with_next = keep_with_next
    return paragraph


def _paragraph_rules(paragraph: Any, edges: Sequence[str]) -> None:
    """Add 0.5pt `TS_BORDER_HEX` paragraph rules in schema order."""
    rules = OxmlElement("w:pBdr")
    for edge in edges:
        border = OxmlElement(f"w:{edge}")
        for name, value in (("val", "single"), ("sz", _BORDER_SIZE), ("space", "4"), ("color", STYLE.TS_BORDER_HEX)):
            border.set(qn(f"w:{name}"), value)
        rules.append(border)
    paragraph._p.get_or_add_pPr().insert_element_before(rules, *_PPR_AFTER_PBDR)


def render_term_sheet_opening(
    doc: Any, metadata: DocumentMetadata, texts: TermSheetTexts, render_runs: RenderRuns
) -> bool:
    """Render the first-page block: title, subtitle, date, prepared_by, logo, disclaimer.

    The title uses the non-outline `Title` style, so it stays out of the
    navigation pane and any TOC. Title and subtitle are bold accent-coloured and
    centred; date and preparer are centred meta lines; an optional house logo
    follows; the disclaimer is a justified 7pt muted block between thin rules or
    in a frame (`style.disclaimer`). With `style.cover: page` the block is set
    lower on a cover page of its own and the body starts on the next page. It
    is rendered regardless of the end-disclaimer switch. House and frontmatter
    texts are parsed for inline emphasis; `metadata.display_runs` wins for
    title, subtitle and date.

    Args:
        doc: Document with request-scoped styles already created.
        metadata: Validated term-sheet metadata.
        texts: Resolved boilerplate for this render.
        render_runs: Injected run renderer.

    Returns:
        True, so the caller skips the first body H1 that repeats the title.
    """
    style = texts.style
    cover = style.cover == "page"
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = _COVER_TITLE_SPACE if cover else Pt(0)
    title.paragraph_format.space_after = Pt(2)
    render_runs(
        title, _bold(_display_runs(metadata, "title", metadata.title)),
        default_color=STYLE.NAVY, font_name=STYLE.HEADING_FONT, font_size=STYLE.TS_TITLE_SIZE,
    )
    render_runs(
        _opening_paragraph(doc, Pt(8)), _bold(_display_runs(metadata, "subtitle", metadata.subtitle)),
        default_color=STYLE.NAVY, font_name=STYLE.HEADING_FONT, font_size=STYLE.TS_SUBTITLE_SIZE,
    )
    date = metadata.extra.get("date")
    date_text = str(date).strip() if date is not None else ""
    if metadata.display_runs.get("date") or date_text:
        render_runs(
            _opening_paragraph(doc, Pt(0)), _display_runs(metadata, "date", date_text),
            font_size=STYLE.TS_META_SIZE,
        )
    render_runs(
        _opening_paragraph(doc, Pt(8)), TextParser.parse_runs(texts.prepared_by.strip()),
        font_size=STYLE.TS_META_SIZE,
    )
    if style.logo:
        _add_logo(doc, style)
    edges = ("top", "left", "bottom", "right") if style.disclaimer == "box" else ("top", "bottom")
    lines = split_run_lines(TextParser.parse_runs(texts.disclaimer.strip()))
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_before = _COVER_DISCLAIMER_SPACE if cover and index == 0 else Pt(0)
        paragraph.paragraph_format.space_after = Pt(6) if index == len(lines) - 1 else Pt(0)
        _paragraph_rules(paragraph, edges)
        render_runs(
            paragraph, line, default_color=_rgb(STYLE.TS_MUTED_HEX), font_size=STYLE.TS_DISCLAIMER_SIZE
        )
    if cover:
        doc.add_page_break()
    doc.core_properties.title = metadata.title
    doc.core_properties.subject = metadata.subtitle
    return True


def _add_logo(doc: Any, style: TermSheetStyle) -> None:
    """Centre the house logo; a logo that cannot be loaded is a render diagnostic."""
    paragraph = _opening_paragraph(doc, Pt(8))
    try:
        source: Union[str, BytesIO] = style.logo
        if style.logo.lower().startswith("data:"):
            header, _, payload = style.logo.partition(",")
            if not header.lower().endswith(";base64"):
                raise ValueError("only base64 data: URIs are supported")
            source = BytesIO(base64.b64decode(payload))
        paragraph.add_run().add_picture(source, width=Mm(style.logo_width_mm))
    except Exception as error:  # a missing or unreadable file must not break the opening
        paragraph._p.getparent().remove(paragraph._p)
        errors = getattr(doc.part, "_ib_render_errors", None)
        if errors is not None:
            errors.append(f"House logo could not be rendered: {error}")


def render_term_sheet_heading(doc: Any, heading: Heading, render_runs: RenderRuns) -> None:
    """Accent-coloured heading; parse-time `Heading.runs` are used when present."""
    level = max(1, min(heading.level, 4))
    sizes = {1: STYLE.H1_SIZE, 2: STYLE.H2_SIZE, 3: STYLE.H3_SIZE, 4: STYLE.H4_SIZE}
    paragraph = doc.add_heading(level=level)
    runs = heading.runs or TextParser.parse_runs(heading.text)
    render_runs(
        paragraph, _bold(runs), default_color=STYLE.NAVY if level < 4 else STYLE.DARK_GRAY,
        font_name=STYLE.HEADING_FONT, font_size=sizes[level],
    )


def render_term_sheet_paragraph(doc: Any, paragraph: Paragraph, render_runs: RenderRuns) -> None:
    """Render each line of a body paragraph as its own Word paragraph (plan §2-8).

    Separate paragraphs let every line take its own marker hanging indent;
    `※` lines use the note size. The text itself is never changed.
    """
    runs = paragraph.runs or TextParser.parse_runs(paragraph.text)
    for line in split_run_lines(runs):
        target = doc.add_paragraph(style=STYLE.STYLE_IB_BODY)
        render_runs(target, line, font_size=apply_marker_layout(target, line_text(line)))


# ═══════════════════════════════════════════════════════════════════════════════
# CONFIRMATION BOX
# ═══════════════════════════════════════════════════════════════════════════════


def confirmation_has_text(confirmation: ConfirmationText) -> bool:
    """True when any confirmation field has visible text to render."""
    return any(text.strip() for text in (confirmation.intro, *confirmation.items, confirmation.signature))


def printable_width(section: Any) -> int:
    """Actual printable width of a section in EMU, without any minimum floor.

    Raises:
        ValueError: The section lacks explicit page width or margins.
    """
    width, left, right = section.page_width, section.left_margin, section.right_margin
    if width is None or left is None or right is None:
        raise ValueError("Section dimensions are required for term-sheet layout")
    return int(width) - int(left) - int(right)


def positive_printable_width(section: Any, what: str) -> int:
    """Printable width for content that must fit its section; diagnose impossible geometry.

    Raises:
        ValueError: The section has no positive printable width for `what`.
    """
    width = printable_width(section)
    if width <= 0:
        raise ValueError(f"{what} has no printable width in its section")
    return width


def _keep_previous_with_next(doc: Any) -> None:
    """Keep the paragraph that precedes the next body block on the same page."""
    body = doc.element.body
    previous = body[-1] if len(body) else None
    if previous is not None and previous.tag == qn("w:sectPr"):
        previous = previous.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        previous.get_or_add_pPr().keepNext_val = True


class _CellWriter:
    """Hand out a cell's paragraphs in order, reusing its initial empty one."""

    def __init__(self, cell: Any) -> None:
        self.cell = cell
        self.count = 0

    def paragraph(self) -> Any:
        paragraph = self.cell.paragraphs[0] if self.count == 0 else self.cell.add_paragraph()
        self.count += 1
        configure_cell_paragraph(paragraph)
        return paragraph


def _item_layout(text: str) -> Optional[MarkerLayout]:
    if text.lstrip().startswith(_CHECKBOX):
        return MarkerLayout(0.0, _CHECKBOX_LAYOUT_MM)
    return marker_layout(text)


def render_confirmation(doc: Any, confirmation: ConfirmationText, render_runs: RenderRuns) -> None:
    """Render the customer confirmation box as a one-column, two-row table.

    Row 1 holds the intro (8pt muted) and the check items (10pt, `□` hanging
    indent; later lines of an item align with its text). Row 2 holds the
    signature, centred bold 10pt on `TS_SIGNATURE_BG_HEX`. Both rows are
    unsplittable, every row-1 paragraph is kept with the next and so are the
    paragraph before the box and a 6pt spacer that separates the box from it, so
    the whole box stays on one page. A row without text is omitted. The frame is
    0.5pt `TS_BORDER_HEX`.

    Args:
        doc: Document receiving the box at the current position.
        confirmation: Wording with at least one non-blank field.
        render_runs: Injected run renderer.

    Raises:
        ValueError: The section has no positive printable width (nothing is emitted).
    """
    width = positive_printable_width(doc.sections[-1], "Confirmation box")
    _keep_previous_with_next(doc)
    _add_spacer(doc, _CONFIRMATION_GAP, keep_with_next=True)
    intro = confirmation.intro.strip()
    items = [item.strip() for item in confirmation.items if item.strip()]
    signature = confirmation.signature.strip()
    kinds = (["body"] if intro or items else []) + (["signature"] if signature else [])
    table = doc.add_table(rows=len(kinds), cols=1)
    table.style = STYLE.STYLE_TABLE_GRID
    table.columns[0].width = Emu(width)
    for cell in table.columns[0].cells:
        cell.width = Emu(width)
    apply_table_frame(table, width)
    muted = _rgb(STYLE.TS_MUTED_HEX)
    for row_index, kind in enumerate(kinds):
        cell = table.rows[row_index].cells[0]
        if kind == "signature":
            set_cell_fill(cell._tc, STYLE.TS_SIGNATURE_BG_HEX)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        writer = _CellWriter(cell)
        if kind == "body":
            for line in split_run_lines(TextParser.parse_runs(intro)) if intro else []:
                render_runs(writer.paragraph(), line, default_color=muted, font_size=STYLE.TS_NOTE_SIZE)
            for item in items:
                layout: Optional[MarkerLayout] = None
                for line_index, line in enumerate(split_run_lines(TextParser.parse_runs(item))):
                    paragraph = writer.paragraph()
                    if line_index == 0:
                        layout = _item_layout(line_text(line))
                        if layout is not None:
                            paragraph.paragraph_format.first_line_indent = Mm(-layout.hanging_mm)
                    if layout is not None:
                        paragraph.paragraph_format.left_indent = Mm(layout.start_mm + layout.hanging_mm)
                    render_runs(paragraph, line, font_size=STYLE.TS_META_SIZE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        else:
            for line in split_run_lines(TextParser.parse_runs(signature)):
                paragraph = writer.paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                render_runs(paragraph, _bold(line), font_size=STYLE.TS_META_SIZE)
        set_row_pagination(table.rows[row_index], keep_together=True, repeat_header=False)
    add_table_spacer(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# HEADER AND FOOTER
# ═══════════════════════════════════════════════════════════════════════════════


def _reset_story(story: Any) -> Any:
    """Unlink a header/footer and return its single, emptied paragraph."""
    story.is_linked_to_previous = False
    paragraphs = story.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    paragraph = paragraphs[0] if paragraphs else story.add_paragraph()
    paragraph.clear()
    return paragraph


def _accent_rule(paragraph: Any, edge: str) -> None:
    """A 1.5pt accent rule on one edge of a header or footer paragraph."""
    rules = OxmlElement("w:pBdr")
    border = OxmlElement(f"w:{edge}")
    for name, value in (("val", "single"), ("sz", "12"), ("space", "4"), ("color", STYLE.NAVY_HEX)):
        border.set(qn(f"w:{name}"), value)
    rules.append(border)
    paragraph._p.get_or_add_pPr().insert_element_before(rules, *_PPR_AFTER_PBDR)


def _add_field(paragraph: Any, instruction: str, color: RGBColor, size: Length) -> None:
    """Append a simple field run (e.g. PAGE) styled like the surrounding text."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        element = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            element.set(qn("w:fldCharType"), kind)
        else:
            element.text = text
        run._r.append(element)
    run.font.name = STYLE.BODY_FONT
    run.font.size = size
    run.font.color.rgb = color
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), STYLE.KOREAN_FONT)


def setup_term_sheet_header_footer(
    doc: Any,
    metadata: DocumentMetadata,
    texts: Optional[TermSheetTexts],
    confidential: bool,
    render_runs: RenderRuns,
) -> None:
    """Set the header and footer of every section once all sections exist.

    Header: right-aligned italic 7.5pt `TS_CONFIDENTIAL_HEX` confidentiality
    label (or `style.label_color`); empty when `confidential` is off or the
    label is blank (plan §2-10). Footer: `"{subtitle} {version}"` at left (empty
    without a version) and the page text (`style.page_number`, default
    `{page} / {pages}`) at a right or centre tab on each section's own printable
    width, so portrait and landscape sections both align. `style.header_rule`
    and `style.footer_rule` add accent rules under the header and over the footer.

    Args:
        doc: Document whose sections are final.
        metadata: Term-sheet metadata supplying subtitle and `version`.
        texts: Resolved boilerplate carrying the confidentiality label.
        confidential: Resolved confidential option for this render.
        render_runs: Injected run renderer.
    """
    label = texts.confidential_label if texts is not None and confidential else ""
    style = texts.style if texts is not None else TermSheetStyle()
    version = str(metadata.extra.get("version") or "").strip()
    left = f"{metadata.subtitle} {version}" if version else ""
    grey = _rgb(STYLE.TS_CONFIDENTIAL_HEX)
    label_color = _rgb(style.label_color[1:]) if style.label_color else grey
    size = STYLE.TS_HEADER_FOOTER_SIZE
    for section in doc.sections:
        width = printable_width(section)
        header = _reset_story(section.header)
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if label.strip():
            render_runs(header, [TextRun(text=label, italic=True)], default_color=label_color, font_size=size)
        if style.header_rule:
            _accent_rule(header, "bottom")
        footer = _reset_story(section.footer)
        if footer.style is not None:
            footer.style.paragraph_format.tab_stops.clear_all()
        if width > 0:  # a section without printable width is reported by the structural audit
            if style.page_number_align == "center":
                footer.paragraph_format.tab_stops.add_tab_stop(Emu(width // 2), WD_TAB_ALIGNMENT.CENTER)
            else:
                footer.paragraph_format.tab_stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)
        if style.footer_rule:
            _accent_rule(footer, "top")
        runs = ([TextRun(text=left)] if left else []) + [TextRun(text="\t")]
        render_runs(footer, runs, default_color=grey, font_size=size)
        for part in _PAGE_FIELD_RE.split(style.page_number):
            if part == "{page}":
                _add_field(footer, "PAGE", grey, size)
            elif part == "{pages}":
                _add_field(footer, "NUMPAGES", grey, size)
            elif part:
                render_runs(footer, [TextRun(text=part)], default_color=grey, font_size=size)
