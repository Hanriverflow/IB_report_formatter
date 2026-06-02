#!/usr/bin/env python3
"""IB style audit tool.

Renders a Markdown input under both style profiles (``classic`` and
``ib-pro``), extracts structural attributes from the resulting ``.docx`` files
via python-docx, evaluates them against an IB-grade rubric, and prints a
comparison report.

This is the measurement backbone of the IB style workflow: it makes "is the
output IB-grade?" answerable with concrete numbers and makes ``classic`` vs
``ib-pro`` differences explicit. It is also importable from tests as a
regression gate.

Usage:
    uv run python tools/ib_style_audit.py output/p1_input.md
    uv run python tools/ib_style_audit.py tests/some_report.md --json
    uv run python tools/ib_style_audit.py --docx a.docx b.docx   # compare two existing files
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# Make project-root modules importable when run as tools/ib_style_audit.py
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

#: Styles we care about for the rubric (others are ignored to keep noise down).
_TRACKED_STYLES = (
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "Normal",
    "Title",
    "IB Body",
    "IB Bullet",
)


# ═══════════════════════════════════════════════════════════════════════════════
# FEATURE EXTRACTION
# ═══════════════════════════════════════════════════════════════════════════════


def _color_hex(font) -> Optional[str]:
    """Return a run/style font color as an uppercase hex string, or None."""
    try:
        color = font.color
        if color is not None and color.type is not None and color.rgb is not None:
            return str(color.rgb).upper()
    except (AttributeError, ValueError):
        pass
    return None


def _style_features(doc) -> Dict[str, Dict[str, Any]]:
    """Extract font/size/color for the tracked named styles."""
    out: Dict[str, Dict[str, Any]] = {}
    for style in doc.styles:
        if style.name not in _TRACKED_STYLES:
            continue
        try:
            font = style.font
            out[style.name] = {
                "font": font.name,
                "size_pt": font.size.pt if font.size is not None else None,
                "bold": bool(font.bold) if font.bold is not None else None,
                "color": _color_hex(font),
            }
        except AttributeError:
            continue
    return out


def _cell_shading(cell) -> Optional[str]:
    """Return the fill color (hex) of a table cell, or None if unshaded."""
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    fill = shd.get(qn("w:fill"))
    if fill in (None, "auto"):
        return None
    return fill.upper()


def _table_borders(table) -> Dict[str, Dict[str, Optional[str]]]:
    """Extract table-level border definitions by edge."""
    result: Dict[str, Dict[str, Optional[str]]] = {}
    tblPr = table._tbl.tblPr
    if tblPr is None:
        return result
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        return result
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is not None:
            result[edge] = {
                "val": el.get(qn("w:val")),
                "sz": el.get(qn("w:sz")),
                "color": el.get(qn("w:color")),
            }
    return result


def _border_style_label(borders: Dict[str, Dict[str, Optional[str]]]) -> str:
    """Classify a table's border treatment for the rubric."""
    if not borders:
        return "none"
    inside = any(
        edge in borders and (borders[edge].get("val") not in (None, "none", "nil"))
        for edge in ("insideH", "insideV")
    )
    return "grid" if inside else "horizontal"


def _table_features(doc) -> List[Dict[str, Any]]:
    """Extract per-table structure: dimensions, header shading, borders."""
    tables = []
    for table in doc.tables:
        rows = len(table.rows)
        cols = len(table.columns) if table.rows else 0
        header_shading = None
        if rows:
            shades = [_cell_shading(c) for c in table.rows[0].cells]
            header_shading = next((s for s in shades if s), None)
        borders = _table_borders(table)
        tables.append(
            {
                "rows": rows,
                "cols": cols,
                "header_shading": header_shading,
                "border_style": _border_style_label(borders),
            }
        )
    return tables


def _alignment_distribution(doc) -> Dict[str, int]:
    """Count paragraph alignments (LEFT/CENTER/RIGHT/JUSTIFY/None)."""
    dist: Dict[str, int] = {}
    for p in doc.paragraphs:
        name = p.alignment.name if p.alignment is not None else "INHERIT"
        dist[name] = dist.get(name, 0) + 1
    return dist


def _color_inventory(doc, styles: Dict[str, Dict[str, Any]]) -> List[str]:
    """Collect the distinct colors used across tracked styles and runs."""
    colors = {info["color"] for info in styles.values() if info.get("color")}
    for p in doc.paragraphs:
        for run in p.runs:
            c = _color_hex(run.font)
            if c:
                colors.add(c)
    return sorted(colors)


def extract_features(docx_path: Path) -> Dict[str, Any]:
    """Extract the full structural feature set from a docx file."""
    doc = Document(str(docx_path))
    section = doc.sections[0]
    styles = _style_features(doc)
    return {
        "margins_in": {
            "top": round(section.top_margin.inches, 3),
            "bottom": round(section.bottom_margin.inches, 3),
            "left": round(section.left_margin.inches, 3),
            "right": round(section.right_margin.inches, 3),
        },
        "styles": styles,
        "alignment": _alignment_distribution(doc),
        "tables": _table_features(doc),
        "colors": _color_inventory(doc, styles),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# RUBRIC EVALUATION
# ═══════════════════════════════════════════════════════════════════════════════


def evaluate_rubric(features: Dict[str, Any]) -> List[Dict[str, str]]:
    """Evaluate extracted features against the IB-grade rubric.

    Returns a list of {item, verdict, detail} where verdict is pass/warn/fail.
    """
    results: List[Dict[str, str]] = []
    styles = features.get("styles", {})

    def size(name: str) -> Optional[float]:
        info = styles.get(name)
        return info.get("size_pt") if info else None

    # 1. Type hierarchy: H1 > H2 > H3 in size
    h1, h2, h3 = size("Heading 1"), size("Heading 2"), size("Heading 3")
    if None in (h1, h2, h3):
        results.append({"item": "type_hierarchy", "verdict": "warn", "detail": "missing heading sizes"})
    elif h1 > h2 > h3:
        results.append({"item": "type_hierarchy", "verdict": "pass", "detail": f"H1={h1} > H2={h2} > H3={h3}"})
    else:
        results.append({"item": "type_hierarchy", "verdict": "fail", "detail": f"H1={h1}, H2={h2}, H3={h3}"})

    # 2. Color restraint: a focused palette reads more institutional
    n_colors = len(features.get("colors", []))
    if n_colors <= 6:
        results.append({"item": "color_restraint", "verdict": "pass", "detail": f"{n_colors} colors"})
    elif n_colors <= 9:
        results.append({"item": "color_restraint", "verdict": "warn", "detail": f"{n_colors} colors"})
    else:
        results.append({"item": "color_restraint", "verdict": "fail", "detail": f"{n_colors} colors"})

    # 3. Table treatment: report border style and header emphasis per table
    tables = features.get("tables", [])
    if not tables:
        results.append({"item": "tables", "verdict": "warn", "detail": "no tables in sample"})
    else:
        grid = sum(1 for t in tables if t["border_style"] == "grid")
        shaded = sum(1 for t in tables if t["header_shading"])
        results.append(
            {
                "item": "tables",
                "verdict": "pass",
                "detail": f"{len(tables)} tables; {grid} full-grid; {shaded} with header shading",
            }
        )

    # 4. Body alignment: Korean reads best left-aligned (justify is a warn)
    align = features.get("alignment", {})
    justified = align.get("JUSTIFY", 0)
    if justified:
        results.append({"item": "body_alignment", "verdict": "warn", "detail": f"{justified} justified paragraphs"})
    else:
        results.append({"item": "body_alignment", "verdict": "pass", "detail": "no justified body text"})

    return results


# ═══════════════════════════════════════════════════════════════════════════════
# DIFF
# ═══════════════════════════════════════════════════════════════════════════════


def diff_features(a: Dict[str, Any], b: Dict[str, Any], path: str = "") -> List[str]:
    """Return a flat list of human-readable differences between two feature sets."""
    diffs: List[str] = []
    if isinstance(a, dict) and isinstance(b, dict):
        for key in sorted(set(a) | set(b)):
            sub = f"{path}.{key}" if path else key
            if key not in a:
                diffs.append(f"+ {sub} = {b[key]!r} (only in ib-pro)")
            elif key not in b:
                diffs.append(f"- {sub} = {a[key]!r} (only in classic)")
            else:
                diffs.extend(diff_features(a[key], b[key], sub))
    elif a != b:
        diffs.append(f"~ {path}: classic={a!r} -> ib-pro={b!r}")
    return diffs


# ═══════════════════════════════════════════════════════════════════════════════
# ORCHESTRATION
# ═══════════════════════════════════════════════════════════════════════════════


def render_profile(md_path: Path, profile: str, out_dir: Path) -> Path:
    """Convert a markdown file to docx under the given style profile."""
    import md_to_word

    out_path = out_dir / f"{md_path.stem}__{profile}.docx"
    args = argparse.Namespace(
        no_cover=False,
        no_toc=False,
        no_disclaimer=False,
        separator_mode="auto",
        format=False,
        deepresearch_cleaner="off",
        cite_mode="footnote",
        drop_unknown_markers=False,
        cleaner_report=False,
        verbose=False,
        output_file=str(out_path),
        style=profile,
    )
    exit_code = md_to_word.run_conversion(md_path, args)
    if exit_code != 0:
        raise RuntimeError(f"Conversion failed for profile {profile!r}")
    return out_path


def audit_markdown(md_path: Path) -> Dict[str, Any]:
    """Render md under both profiles and return features, rubric, and diff."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        classic_docx = render_profile(md_path, "classic", tmp_dir)
        ibpro_docx = render_profile(md_path, "ib-pro", tmp_dir)
        classic = extract_features(classic_docx)
        ibpro = extract_features(ibpro_docx)
    return {
        "input": str(md_path),
        "classic": {"features": classic, "rubric": evaluate_rubric(classic)},
        "ib_pro": {"features": ibpro, "rubric": evaluate_rubric(ibpro)},
        "diff": diff_features(classic, ibpro),
    }


def audit_docx_pair(classic_docx: Path, ibpro_docx: Path) -> Dict[str, Any]:
    """Audit two already-rendered docx files."""
    classic = extract_features(classic_docx)
    ibpro = extract_features(ibpro_docx)
    return {
        "input": f"{classic_docx} vs {ibpro_docx}",
        "classic": {"features": classic, "rubric": evaluate_rubric(classic)},
        "ib_pro": {"features": ibpro, "rubric": evaluate_rubric(ibpro)},
        "diff": diff_features(classic, ibpro),
    }


def _format_report(result: Dict[str, Any]) -> str:
    """Render the audit result as a readable markdown report."""
    lines = [f"# IB Style Audit — {result['input']}", ""]
    for profile_key, label in (("classic", "classic"), ("ib_pro", "ib-pro")):
        lines.append(f"## {label} rubric")
        for r in result[profile_key]["rubric"]:
            mark = {"pass": "[PASS]", "warn": "[WARN]", "fail": "[FAIL]"}[r["verdict"]]
            lines.append(f"  {mark} {r['item']}: {r['detail']}")
        lines.append("")
    lines.append("## classic -> ib-pro diff")
    if result["diff"]:
        lines.extend(f"  {d}" for d in result["diff"])
    else:
        lines.append("  (identical — no structural differences)")
    lines.append("")
    return "\n".join(lines)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Audit IB Word output style (classic vs ib-pro)")
    parser.add_argument("input", nargs="?", help="Markdown input file to render under both profiles")
    parser.add_argument(
        "--docx",
        nargs=2,
        metavar=("CLASSIC", "IBPRO"),
        help="Compare two already-rendered docx files instead of rendering",
    )
    parser.add_argument("--json", action="store_true", help="Emit JSON instead of a text report")
    return parser


def main() -> int:
    # Windows consoles default to cp949 (Korean), which cannot encode em dashes
    # or Korean paths; match the rest of the project and force UTF-8 stdout.
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except (ValueError, OSError):
            pass

    args = build_parser().parse_args()
    if args.docx:
        result = audit_docx_pair(Path(args.docx[0]), Path(args.docx[1]))
    elif args.input:
        result = audit_markdown(Path(args.input))
    else:
        build_parser().print_help()
        return 1

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(_format_report(result))
    return 0


if __name__ == "__main__":
    sys.exit(main())
