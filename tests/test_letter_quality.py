"""Regression coverage from office-letter production failures (synthetic data)."""

import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from document_profiles import RenderOptions
from docx_audit import inspect_document
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, TextParser


def render(source, profile="plain"):
    """Exercise parser, renderer and saved Word package together."""
    model = MarkdownParser(profile=profile).parse(source)
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return Document(buffer)


@pytest.mark.parametrize("tag", ["<br>", "<br/>", "<BR />"])
@pytest.mark.parametrize("context", ["paragraph", "bold", "heading", "list", "table"])
def test_inline_break_reaches_word_in_every_text_context(tag, context):
    source = {
        "paragraph": "앞" + tag + "뒤",
        "bold": "**앞" + tag + "뒤**",
        "heading": "## 앞" + tag + "뒤",
        "list": "1. 앞" + tag + "뒤",
        "table": "| 항목 | 값 |\n|---|---|\n| 앞" + tag + "뒤 | 001234 |",
    }[context]
    doc = render(source)
    assert len(doc.element.xpath('.//w:br[not(@w:type) or @w:type="textWrapping"]')) == 1
    assert not any("<br" in t.lower() for t in doc.element.xpath(".//w:t/text()"))
    if context == "bold":
        assert doc.element.xpath(".//w:r[w:rPr/w:b]/w:br")
    if context == "table":
        assert doc.tables[0].cell(1, 1).text == "001234"


def test_breaks_do_not_add_an_extra_space_or_duplicate_source_newline():
    doc = render("앞<br>\n뒤<br><br>끝")
    assert doc.paragraphs[0].text == "앞\n뒤\n\n끝"


@pytest.mark.parametrize("source, expected", [(r"앞\<br>뒤", "앞<br>뒤"), ("`<br>`", "`<br>`"), ("`**<br>**`", "`**<br>**`"), ("`$5<br>$`", "`$5<br>$`")])
def test_literal_break_syntax_stays_literal(source, expected):
    doc = render(source)
    assert doc.paragraphs[0].text == expected
    assert not doc.element.xpath(".//w:br")


def test_br_does_not_change_link_destination_or_currency():
    runs = TextParser.parse_runs("[첫<br>둘](https://example.com/<br>)")
    assert "".join(r.text for r in runs) == "첫\n둘"
    assert runs[0].hyperlink == "https://example.com/<br>"
    assert "".join(r.text for r in TextParser.parse_runs_plain("$123<br>001234")) == "$123\n001234"


def test_literal_code_keeps_surrounding_emphasis_and_table_context():
    runs = TextParser.parse_runs("**앞 `*<br>*` 뒤**")
    assert "".join(r.text for r in runs) == "앞 `*<br>*` 뒤"
    assert all(r.bold for r in runs)
    doc = render("| 항목 |\n|---|\n| `$5<br>$` |")
    assert doc.tables[0].cell(1, 0).text == "`$5<br>$`"
    assert not doc.element.xpath(".//w:br")


LETTER = """---
profile: office-letter
title: 자료 확인 요청
document_no: 예시-001
date: "2026-09-15"
sender:
  organization: 가상회사
  signatory: 대표자 [성명 기재]
  contact: office@example.com
recipients: [가상 수신기관]
cc: [가상 담당부서]
attachments: [상세 내역 1부, 확인서 1부]
letter:
  appendix_heading: 상세 내역
  appendix_label: 붙임 1
---
# 자료 확인 요청

1. 입력 내용을 확인해 주시기 바랍니다.
2. 처리 결과를 담당자에게 알려주시기 바랍니다.

# 상세 내역

| 구분 | 내용 |
|---|---|
| 코드 | 001234 |
| 연락처 | 가상 담당자<br>office@example.com |
"""


def test_letter_closes_before_explicit_appendix_and_uses_native_attachments():
    doc = render(LETTER, "office-letter")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.index("대표자 [성명 기재]") < text.rindex("상세 내역")
    assert text.count("끝.") == 1
    assert len(doc.element.xpath('.//w:br[@w:type="page"]')) == 1
    assert len(doc.element.xpath(".//w:pPr/w:numPr")) == 4
    assert not any(p.text == "끝." for p in doc.paragraphs)
    assert "수신\t가상 수신기관" in text
    assert "제목\t자료 확인 요청" in text
    assert "office@example.com" in text
    assert not doc.sections[0].header.paragraphs[0].text
    assert doc.tables[0].cell(1, 1).text == "001234"


@pytest.mark.parametrize("replacement", ["# 다른 내역", "# 상세 내역\n\n# 상세 내역"])
def test_invalid_appendix_boundary_is_rejected(replacement):
    with pytest.raises(ValueError, match="appendix_heading"):
        render(LETTER.replace("# 상세 내역", replacement), "office-letter")


@pytest.mark.parametrize("metadata", ["letter: wrong", "letter: {unknown: true}", "letter: {appendix_label: 붙임}", "letter: {appendix_heading: 123}"])
def test_letter_options_fail_closed(metadata):
    source = LETTER.replace("letter:\n  appendix_heading: 상세 내역\n  appendix_label: 붙임 1", metadata)
    with pytest.raises(ValueError, match="letter"):
        render(source, "office-letter")


def test_audit_distinguishes_pagination_marks_from_bullets_and_visual_review():
    doc = Document()
    doc.styles["Normal"].paragraph_format.keep_together = True
    doc.add_paragraph("A normal paragraph, not a bullet")
    report = inspect_document(doc)
    assert report.pagination_marked_paragraphs == 1
    assert report.numbered_paragraphs == 0
    assert report.warnings and "Normal" in report.warnings[0]
    assert report.visual_review == "not_performed"
    assert not report.issues


def test_local_false_overrides_inherited_pagination_marker():
    doc = Document()
    doc.styles["Normal"].paragraph_format.keep_together = True
    doc.add_paragraph("Override").paragraph_format.keep_together = False
    assert inspect_document(doc).pagination_marked_paragraphs == 0


def test_office_body_does_not_inherit_blanket_pagination_flags():
    doc = render(LETTER, "office-letter")
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("입력 내용을"):
            assert not paragraph.paragraph_format.keep_together
            assert not paragraph.style.paragraph_format.keep_together
    assert not doc.styles["Normal"].paragraph_format.keep_together
    assert not doc.styles["Normal"].paragraph_format.keep_with_next
    # Proper headings retain meaningful pagination control; do not strip it globally.
    assert doc.styles["Heading 1"].paragraph_format.keep_with_next
    assert doc.element.xpath(".//w:trPr/w:cantSplit")


@pytest.mark.parametrize("source", [r"앞\<br>", "```html\n앞<br>뒤\n```"])
def test_literal_break_at_end_of_line_and_in_code_block(source):
    doc = render(source)
    assert "<br>" in "".join(doc.element.xpath(".//w:t/text()"))
    assert not doc.element.xpath(".//w:br")


def test_office_letter_with_explicit_cover_has_closing_styles():
    model = MarkdownParser().parse(LETTER)
    doc = IBDocumentRenderer(options=RenderOptions(include_cover=True, strict=True)).render(model)
    assert "Office Signatory" in doc.styles
    assert "대표자 [성명 기재]" in "\n".join(p.text for p in doc.paragraphs)


def test_existing_closing_marker_is_not_duplicated_without_attachments():
    source = LETTER.split("# 상세 내역")[0].replace(
        "attachments: [상세 내역 1부, 확인서 1부]\n", ""
    ).replace("letter:\n  appendix_heading: 상세 내역\n  appendix_label: 붙임 1\n", "")
    doc = render(source + "\n본문을 마칩니다.  끝.", "office-letter")
    assert sum(p.text.count("끝.") for p in doc.paragraphs) == 1


def test_office_letter_cli_api_registry_share_appendix_composition(tmp_path, monkeypatch):
    from converters import get_default_registry
    from md_to_word import IBReportConverter, main

    source = tmp_path / "가상 공문.md"
    source.write_text(LETTER, encoding="utf-8")
    paths = [tmp_path / (name + ".docx") for name in ("cli", "api", "registry")]
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(paths[0]), "--strict"])
    with pytest.raises(SystemExit) as result:
        main()
    assert result.value.code == 0
    IBReportConverter(str(source), str(paths[1]), render_options=RenderOptions(strict=True)).convert()
    registry = get_default_registry()
    model = registry.convert(str(source))
    registry.convert(model, output_format="docx", output_path=str(paths[2]), strict=True)
    documents = [Document(path) for path in paths]
    assert documents[0].element.xml == documents[1].element.xml == documents[2].element.xml


def test_invalid_appendix_cli_does_not_save(tmp_path, monkeypatch):
    from md_to_word import main

    source, output = tmp_path / "bad.md", tmp_path / "bad.docx"
    source.write_text(LETTER.replace("# 상세 내역", "# 경계 없음"), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output), "--strict"])
    with pytest.raises(SystemExit) as result:
        main()
    assert result.value.code != 0
    assert not output.exists()


def test_default_letter_has_no_forced_appendix_break():
    source = Path(__file__).resolve().parents[1] / "samples" / "profiles" / "office-letter.md"
    doc = render(source.read_text(encoding="utf-8"), "office-letter")
    assert not doc.element.xpath('.//w:br[@w:type="page"]')
    assert sum(p.text.count("끝.") for p in doc.paragraphs) == 1


def test_wrapped_letter_metadata_aligns_below_value_not_label():
    doc = render(LETTER, "office-letter")
    for name in ("Office Metadata", "Office Subject", "Office Contact"):
        fmt = doc.styles[name].paragraph_format
        assert fmt.left_indent.pt == 42
        assert fmt.first_line_indent.pt == -42
        assert fmt.tab_stops[0].position.pt == 42


def test_appendix_requires_preceding_letter_body():
    source = LETTER.replace("1. 입력 내용을 확인해 주시기 바랍니다.\n2. 처리 결과를 담당자에게 알려주시기 바랍니다.", "")
    with pytest.raises(ValueError, match="appendix_heading must follow"):
        render(source, "office-letter")
