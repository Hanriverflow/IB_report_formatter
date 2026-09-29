"""Schema order, forced-TOC title placement and UTF-8 audit output for existing profiles."""

import glob
import json
import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, parse_markdown_file
from ooxml_order import SEQUENCES, insert_ordered

ROOT = Path(__file__).resolve().parents[1]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
LEGACY_SAMPLES = sorted(
    path for path in glob.glob(str(ROOT / "samples" / "profiles" / "*.md"))
    + glob.glob(str(ROOT / "samples" / "qa" / "*.md"))
    if "term-sheet" not in Path(path).name
)


def _reopen(doc):
    payload = BytesIO()
    doc.save(payload)
    return Document(payload)


def _order_violations(root) -> list:
    problems = []
    for kind, sequence in SEQUENCES.items():
        for element in root.iter(W + kind):
            names = [etree.QName(c).localname for c in element if isinstance(c.tag, str)]
            known = [name for name in names if name in sequence]
            ranks = [sequence.index(name) for name in known]
            if ranks != sorted(ranks) or len(set(known)) != len(known):
                problems.append((kind, names))
    return problems


# ═══════════════════════════════════════════════════════════════════════════════
# SCHEMA ORDER
# ═══════════════════════════════════════════════════════════════════════════════


def test_insert_ordered_places_child_before_its_successors_and_replaces_duplicates():
    tc_pr = OxmlElement("w:tcPr")
    for tag in ("w:tcW", "w:vAlign"):
        tc_pr.append(OxmlElement(tag))
    first = OxmlElement("w:shd")
    first.set(qn("w:fill"), "111111")
    insert_ordered(tc_pr, first)
    second = OxmlElement("w:shd")
    second.set(qn("w:fill"), "222222")
    insert_ordered(tc_pr, second)
    assert [etree.QName(c).localname for c in tc_pr] == ["tcW", "shd", "vAlign"]
    assert tc_pr.find(qn("w:shd")).get(qn("w:fill")) == "222222"


def test_insert_ordered_rejects_unknown_containers():
    with pytest.raises(KeyError):
        insert_ordered(OxmlElement("w:sectPr"), OxmlElement("w:pgSz"))


@pytest.mark.parametrize("toc", [None, True])
@pytest.mark.parametrize("path", LEGACY_SAMPLES, ids=lambda p: Path(p).name)
def test_legacy_samples_follow_the_schema_order(path: str, toc) -> None:
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(include_toc=toc)).render(
        parse_markdown_file(path)
    ))
    roots = [doc.element.body, doc.styles.element]
    for section in doc.sections:
        roots += [section.header._element, section.footer._element]
    assert [problem for root in roots for problem in _order_violations(root)] == []


# ═══════════════════════════════════════════════════════════════════════════════
# FORCED TOC WITH AN OPENING TITLE
# ═══════════════════════════════════════════════════════════════════════════════

OFFICE_FRONTMATTER = (
    "sender:\n  organization: 가상회사\n  department: 기획팀\nrecipients:\n  - 협력사 담당부서\n"
)


@pytest.mark.parametrize("profile,toc_title", [
    ("business-report", "목차"),
    ("meeting-minutes", "목차"),
    ("office-letter", "목차"),
    ("ib-memo", "TABLE OF CONTENTS"),
])
def test_forced_toc_follows_the_opening_title_and_omits_it(profile: str, toc_title: str) -> None:
    extra = OFFICE_FRONTMATTER if profile == "office-letter" else ""
    source = (
        f"---\nprofile: {profile}\ntitle: 가상 보고서\n{extra}---\n"
        "# 가상 보고서\n\n## 1. 개요\n\n본문.\n"
    )
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(include_toc=True, strict=True)).render(
        MarkdownParser().parse(source)
    ))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    title_lines = [index for index, text in enumerate(texts) if text.endswith("가상 보고서")]
    assert len(title_lines) == 1, texts
    assert title_lines[0] < texts.index(toc_title)
    # The TOC preview lists body headings only; the body keeps its section heading.
    assert texts.count("1. 개요") == 2


# ═══════════════════════════════════════════════════════════════════════════════
# DOCX-AUDIT OUTPUT ENCODING
# ═══════════════════════════════════════════════════════════════════════════════


def test_docx_audit_writes_utf8_json_under_a_legacy_code_page(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("[Image: 가상 도표]")
    path = tmp_path / "audit.docx"
    document.save(str(path))
    env = dict(os.environ, PYTHONIOENCODING="cp949")
    completed = subprocess.run(
        [sys.executable, "-B", "-m", "docx_audit", str(path)],
        cwd=str(ROOT), env=env, capture_output=True, check=False,
    )
    assert completed.returncode == 0, completed.stderr
    payload = json.loads((completed.stdout + completed.stderr).decode("utf-8"))
    assert any("가상 도표" in warning for warning in payload["warnings"])
