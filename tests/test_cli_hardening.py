"""End-to-end regressions for saving, batch conversion, and CLI boundaries."""

import logging
import os
import sys
import zipfile
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from io import BytesIO
from pathlib import Path
from threading import Barrier, Event

import pytest
import yaml
from docx import Document
from docx.document import Document as WordDocument

import cli_utils
import converters
import md_to_word
from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser
from md_to_word import IBReportConverter, build_parser, generate_output_path, run_batch_conversion


def _render_document(text):
    model = MarkdownParser(profile="plain").parse(text)
    return IBDocumentRenderer(options=RenderOptions(profile="plain", strict=True)).render(model)


def _convert(source, output, entry_point):
    if entry_point == "cli":
        return Path(IBReportConverter(
            str(source), str(output), RenderOptions(profile="plain", strict=True)
        ).convert())
    registry = converters.get_default_registry()
    model = registry.convert(source, profile="plain")
    return Path(registry.convert(
        model, output_format="docx", output_path=output, profile="plain", strict=True
    ))


@pytest.mark.parametrize("entry_point", ["cli", "registry"])
def test_save_serialization_failure_preserves_existing_report(tmp_path, monkeypatch, entry_point):
    source, output = tmp_path / "report.md", tmp_path / "report.docx"
    source.write_text("Replacement report", encoding="utf-8")
    _render_document("Original report").save(output)
    original_bytes = output.read_bytes()
    original_save = WordDocument.save

    def failing_save(self, path):
        original_save(self, path)
        Path(path).write_bytes(b"incomplete serialization")
        raise OSError("injected serialization failure")

    monkeypatch.setattr(WordDocument, "save", failing_save)
    with pytest.raises(OSError, match="injected serialization failure"):
        _convert(source, output, entry_point)

    assert output.read_bytes() == original_bytes
    assert set(tmp_path.iterdir()) == {source, output}
    assert "Original report" in [p.text for p in Document(output).paragraphs]


@pytest.mark.parametrize("error_type", [OSError, PermissionError])
def test_save_replace_failure_cleans_temp_and_handles_lock(tmp_path, monkeypatch, error_type):
    source, output = tmp_path / "report.md", tmp_path / "report.docx"
    source.write_text("Replacement report", encoding="utf-8")
    _render_document("Original report").save(output)
    original_bytes = output.read_bytes()
    original_replace = os.replace

    def failing_replace(src, dst):
        if Path(dst) == output:
            raise error_type("injected replace failure")
        return original_replace(src, dst)

    monkeypatch.setattr(os, "replace", failing_replace)
    expected_paths = {source, output}
    if error_type is PermissionError:
        saved = _convert(source, output, "cli")
        assert saved != output
        assert "Replacement report" in [p.text for p in Document(saved).paragraphs]
        expected_paths.add(saved)
    else:
        with pytest.raises(OSError, match="injected replace failure"):
            _convert(source, output, "cli")
    assert output.read_bytes() == original_bytes
    assert set(tmp_path.iterdir()) == expected_paths


@pytest.mark.skipif(os.name == "nt", reason="POSIX permission bits")
def test_atomic_save_keeps_direct_save_permissions(tmp_path):
    source = tmp_path / "report.md"
    source.write_text("Report", encoding="utf-8")
    new_output, existing_output = tmp_path / "new.docx", tmp_path / "existing.docx"
    existing_output.write_bytes(b"old")
    existing_output.chmod(0o640)
    umask = os.umask(0o022)
    try:
        _convert(source, new_output, "cli")
        _convert(source, existing_output, "cli")
    finally:
        os.umask(umask)

    assert new_output.stat().st_mode & 0o777 == 0o644
    assert existing_output.stat().st_mode & 0o777 == 0o640


@pytest.mark.parametrize("concurrent", [False, True])
def test_locked_fallback_names_are_exclusive_with_frozen_clock(tmp_path, monkeypatch, concurrent):
    output = tmp_path / "report.docx"
    _render_document("Locked original").save(output)
    original_bytes = output.read_bytes()
    prior_fallback = tmp_path / "report_1700000000.docx"
    prior_fallback.write_bytes(b"existing fallback must survive")
    sources = [tmp_path / "first.md", tmp_path / "second.md"]
    for index, source in enumerate(sources):
        source.write_text(f"Report {index}", encoding="utf-8")
    original_replace = os.replace
    ready = Barrier(2)

    def locked_replace(src, dst):
        if Path(dst) == output:
            if concurrent:
                ready.wait(timeout=10)
            raise PermissionError("report is open in Word")
        return original_replace(src, dst)

    monkeypatch.setattr(os, "replace", locked_replace)
    monkeypatch.setattr(cli_utils.time, "time", lambda: 1700000000)
    if concurrent:
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(_convert, source, output, "cli") for source in sources]
            saved = [future.result(timeout=15) for future in futures]
    else:
        saved = [_convert(source, output, "cli") for source in sources]

    assert prior_fallback.read_bytes() == b"existing fallback must survive"
    assert len(set(saved + [output, prior_fallback])) == 4
    assert output.read_bytes() == original_bytes
    for index, path in enumerate(saved):
        assert f"Report {index}" in [p.text for p in Document(path).paragraphs]
    assert set(tmp_path.iterdir()) == set(sources + saved + [output, prior_fallback])


def test_locked_fallback_failure_removes_reserved_path(tmp_path, monkeypatch):
    source, output = tmp_path / "report.md", tmp_path / "report.docx"
    source.write_text("Replacement report", encoding="utf-8")
    _render_document("Original report").save(output)
    original_bytes = output.read_bytes()

    def failing_replace(src, dst):
        if Path(dst) == output:
            raise PermissionError("report is open in Word")
        raise OSError("fallback failed")

    monkeypatch.setattr(os, "replace", failing_replace)
    with pytest.raises(OSError, match="fallback failed"):
        _convert(source, output, "cli")
    assert output.read_bytes() == original_bytes
    assert set(tmp_path.iterdir()) == {source, output}


@pytest.mark.parametrize("separate_output", [False, True])
@pytest.mark.parametrize("names", [
    ("가상 보고.md", "가상_보고.md"),
    pytest.param(
        ("가상 보고A.md", "가상_보고a.md"),
        marks=pytest.mark.skipif(os.name != "nt", reason="Windows case-insensitive paths"),
    ),
])
def test_batch_preflight_rejects_collisions_and_converts_other_inputs(
    tmp_path, caplog, separate_output, names
):
    input_dir = tmp_path / "보고 입력"
    input_dir.mkdir()
    output_dir = tmp_path / "보고 출력" if separate_output else input_dir
    output_dir.mkdir(exist_ok=True)
    sources = [input_dir / name for name in names]
    planned = [output_dir / generate_output_path(source).name for source in sources]
    healthy = input_dir / "정상 보고.md"
    for index, source in enumerate(sources + [healthy]):
        source.write_text(f"Report {index}", encoding="utf-8")
    planned[0].write_bytes(b"Existing report")
    args = build_parser().parse_args([
        str(input_dir), *([str(output_dir)] if separate_output else []),
        "--batch", "--profile", "plain", "--strict",
    ])
    caplog.set_level(logging.INFO)

    assert run_batch_conversion(input_dir, args) == 1

    assert planned[0].read_bytes() == b"Existing report"
    healthy_output = output_dir / generate_output_path(healthy).name
    assert "Report 2" in [p.text for p in Document(healthy_output).paragraphs]
    assert set(output_dir.glob("*.docx")) == {planned[0], healthy_output}
    errors = [r.getMessage() for r in caplog.records if r.levelno == logging.ERROR]
    assert any("collision" in message.lower() and all(
        str(source) in message for source in sources
    ) for message in errors)
    assert "1 succeeded, 2 failed" in caplog.text


@pytest.mark.parametrize("invalid_markdown", ["", "![Missing](missing.png)"])
def test_batch_isolates_invalid_and_strict_failures(tmp_path, invalid_markdown):
    input_dir, output_dir = tmp_path / "보고 입력", tmp_path / "보고 출력"
    input_dir.mkdir()
    sources = [input_dir / name for name in ("1 정상.md", "2 오류.md", "3 정상.md")]
    for source, text in zip(sources, ["First report", invalid_markdown, "Last report"]):
        source.write_text(text, encoding="utf-8")
    args = build_parser().parse_args([
        str(input_dir), str(output_dir), "--batch", "--profile", "plain", "--strict",
    ])

    assert run_batch_conversion(input_dir, args) == 1

    assert not (output_dir / generate_output_path(sources[1]).name).exists()
    for source, text in [(sources[0], "First report"), (sources[2], "Last report")]:
        output = output_dir / generate_output_path(source).name
        assert text in [p.text for p in Document(output).paragraphs]
    assert len(list(output_dir.iterdir())) == 2


@pytest.mark.parametrize("relative_input", [str(Path("reports") / "q4" / "report.md"), "./report.md"])
def test_cli_missing_qualified_input_does_not_convert_parent_basename(
    tmp_path, monkeypatch, caplog, relative_input
):
    parent_dir, cwd = tmp_path / "parent", tmp_path / "workspace"
    parent_dir.mkdir()
    cwd.mkdir()
    wrong_source = parent_dir / "report.md"
    wrong_source.write_text("Wrong document", encoding="utf-8")
    monkeypatch.chdir(cwd)
    monkeypatch.setattr(md_to_word, "PARENT_DIR", parent_dir)
    monkeypatch.setattr(md_to_word, "configure_logging", lambda verbose: None)
    monkeypatch.setattr(sys, "argv", ["md-to-word", relative_input, "--profile", "plain", "--strict"])

    with pytest.raises(SystemExit) as result:
        md_to_word.main()

    assert result.value.code == 1
    assert "File not found" in caplog.text
    assert str(cwd / relative_input) in caplog.text
    assert not list(tmp_path.rglob("*.docx"))


def test_cli_bare_filename_keeps_parent_fallback(tmp_path, monkeypatch):
    parent_dir, cwd = tmp_path / "parent", tmp_path / "workspace"
    parent_dir.mkdir()
    cwd.mkdir()
    source = parent_dir / "보고 파일.md"
    source.write_text("Intended document", encoding="utf-8")
    monkeypatch.chdir(cwd)
    monkeypatch.setattr(md_to_word, "PARENT_DIR", parent_dir)
    monkeypatch.setattr(md_to_word, "configure_logging", lambda verbose: None)
    monkeypatch.setattr(sys, "argv", ["md-to-word", source.name, "--profile", "plain", "--strict"])

    with pytest.raises(SystemExit) as result:
        md_to_word.main()

    assert result.value.code == 0
    assert "Intended document" in [p.text for p in Document(generate_output_path(source)).paragraphs]


def test_concurrent_first_registry_calls_wait_for_builtin_registration(tmp_path, monkeypatch):
    source = tmp_path / "report.md"
    source.write_text("Threaded report", encoding="utf-8")
    registration_started, release_registration = Event(), Event()
    contenders_ready = Barrier(4)
    original_register = converters._register_builtin_converters
    registrations = []

    def slow_register(registry):
        registrations.append(registry)
        registration_started.set()
        assert release_registration.wait(timeout=10)
        original_register(registry)

    def convert_in_thread(index):
        if index:
            contenders_ready.wait(timeout=10)
        registry = converters.get_default_registry()
        model = registry.convert(source, profile="plain")
        output = registry.convert(
            model, output_format="docx", output_path=tmp_path / f"report-{index}.docx",
            profile="plain", strict=True,
        )
        assert "Threaded report" in [p.text for p in Document(output).paragraphs]
        return registry

    monkeypatch.setattr(converters, "_default_registry", None)
    monkeypatch.setattr(converters, "_register_builtin_converters", slow_register)
    with ThreadPoolExecutor(max_workers=4) as pool:
        first = pool.submit(convert_in_thread, 0)
        try:
            assert registration_started.wait(timeout=10)
            contenders = [pool.submit(convert_in_thread, index) for index in range(1, 4)]
            contenders_ready.wait(timeout=10)
            # Give contenders time to observe an early-published registry, if present.
            wait(contenders, timeout=0.5, return_when=FIRST_COMPLETED)
        finally:
            release_registration.set()
        registries = [future.result(timeout=15) for future in [first] + contenders]

    assert len(registrations) == 1
    assert all(registry is registries[0] for registry in registries)
    assert [converter.name for converter in registries[0].converters] == [
        "markdown-input", "docx-output",
    ]


def _ci_workflow():
    root = Path(__file__).resolve().parents[1]
    return yaml.safe_load((root / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8"))


def test_ci_workflow_config_and_python312_gate():
    workflow = _ci_workflow()
    assert set(workflow["on"]) == {"push", "pull_request"}
    job = workflow["jobs"]["quality"]
    assert job["strategy"]["fail-fast"] is False
    assert job["strategy"]["matrix"] == {
        "os": ["ubuntu-latest", "windows-latest"], "python-version": ["3.12", "3.13"],
    }
    assert job["runs-on"] == "${{ matrix.os }}"
    steps = job["steps"]
    assert steps[0]["uses"] == "actions/checkout@v4"
    assert steps[1]["uses"] == "astral-sh/setup-uv@v6"
    assert steps[1]["with"]["python-version"] == "${{ matrix.python-version }}"
    commands = [step["run"] for step in steps if "run" in step]
    assert commands[:5] == [
        "uv sync --locked --extra dev", "uv run ruff check .", "uv run mypy .",
        "uv run pytest tests/ -q", "uv build",
    ]
    gate = next(step for step in steps if step.get("id") == "python312")
    assert gate["name"] == "Verify Python 3.12 syntax"
    assert gate["shell"] == "uv run python {0}"
    assert "feature_version=(3, 12)" in gate["run"]
    exec(compile(gate["run"], "ci-python312", "exec"), {})


@pytest.mark.parametrize("source, supported", [
    pytest.param("type Alias[T] = list[T]\n", True, id="python312-type-alias"),
    pytest.param("type Alias[T = int] = list[T]\n", False, id="python313-type-default"),
])
def test_ci_syntax_gate_enforces_python312(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source: str, supported: bool
) -> None:
    gate = next(step for step in _ci_workflow()["jobs"]["quality"]["steps"]
                if step.get("name", "").startswith("Verify Python "))
    (tmp_path / "engine.py").write_text(source, encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    if supported:
        exec(compile(gate["run"], "ci-python312", "exec"), {})
    else:
        with pytest.raises(SyntaxError):
            exec(compile(gate["run"], "ci-python312", "exec"), {})


@pytest.mark.parametrize("contents", ["valid", "missing", "report", "extra_module"])
def test_ci_wheel_gate_checks_exact_engine_payload(tmp_path, monkeypatch, contents):
    gate = next(step for step in _ci_workflow()["jobs"]["quality"]["steps"]
                if step.get("id") == "wheel_contents")
    assert gate["shell"] == "uv run python {0}"
    (tmp_path / "pyproject.toml").write_text(
        '[tool.hatch.build.targets.wheel]\nonly-include = ["engine.py"]\n', encoding="utf-8"
    )
    dist = tmp_path / "dist"
    dist.mkdir()
    report = BytesIO()
    _render_document("Private report must never be packaged").save(report)
    with zipfile.ZipFile(dist / "engine.whl", "w") as wheel:
        wheel.writestr("engine-1.0.dist-info/METADATA", "Name: engine")
        if contents != "missing":
            wheel.writestr("engine.py", "# engine module\n")
        if contents == "report":
            wheel.writestr("reports/private.docx", report.getvalue())
        if contents == "extra_module":
            wheel.writestr("unexpected.py", "# unexpected module\n")
    monkeypatch.chdir(tmp_path)
    if contents == "valid":
        exec(compile(gate["run"], "ci-wheel", "exec"), {})
    else:
        with pytest.raises(AssertionError, match="Wheel payload"):
            exec(compile(gate["run"], "ci-wheel", "exec"), {})
