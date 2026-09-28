"""Chart regressions through the public parser, renderer, converter and CLI."""

import sys
import warnings
from io import BytesIO
from uuid import UUID

import pytest
from docx import Document

import md_to_word
from converters import get_default_registry
from document_model import ElementType
from document_profiles import PROFILES, RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser

SPEC = """chart_type: bar
title: 가상 매출
labels: [상반기, 하반기]
series:
  - name: 금액
    values: [100, -30]
y_label: 백만원
source: 가상 자료
"""
CHART = "```chart\n" + SPEC + "```\n"


def render(markdown=CHART, **kwargs):
    """Exercise actual parsing and round-trip DOCX serialization."""
    profile = kwargs.pop("profile", "plain")
    model = MarkdownParser(profile=profile).parse(markdown)
    renderer = IBDocumentRenderer(options=RenderOptions(profile=profile, **kwargs))
    doc = renderer.render(model)
    buffer = BytesIO()
    doc.save(buffer)
    return model, renderer, Document(buffer)


@pytest.mark.parametrize("profile", PROFILES)
def test_chart_off_matches_legacy_code_panel(profile):
    prefix = "---\nsender: {organization: Example}\nrecipients: [Team]\n---\n"
    _, renderer, doc = render(prefix + CHART, profile=profile)
    assert not doc.inline_shapes
    assert any(SPEC.rstrip() == table.cell(0, 0).text for table in doc.tables)
    assert not renderer.errors


@pytest.mark.parametrize("kind", ["bar", "line", "waterfall"])
@pytest.mark.parametrize("type_key", ["chart_type", "type"])
def test_chart_is_model_element_and_saved_image(kind, type_key):
    from matplotlib.font_manager import FontProperties, findfont

    from ib_renderer import FontPolicy

    try:
        findfont(FontProperties(family=FontPolicy.resolve_korean_font()), fallback_to_default=False)
        korean_font_available = True
    except ValueError:
        korean_font_available = False
    markdown = CHART.replace("chart_type: bar", f"{type_key}: {kind}")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        model, renderer, doc = render(markdown, charts=True, strict=True)
    assert model.elements[0].element_type == ElementType.CHART
    assert model.elements[0].content.spec.chart_type == kind
    assert len(doc.inline_shapes) == 1
    assert not doc.tables
    assert not renderer.errors
    if korean_font_available:
        assert not [w for w in caught if "Glyph" in str(w.message) and "missing" in str(w.message)]


@pytest.mark.parametrize("frontmatter, override, count", [
    ("", None, 0), ("charts: true", None, 1), ("charts: false", True, 1),
    ("charts: true", False, 0), ("charts: false", None, 0),
])
def test_charts_precedence(frontmatter, override, count):
    _, _, doc = render("---\n" + frontmatter + "\n---\n" + CHART, charts=override)
    assert len(doc.inline_shapes) == count


@pytest.mark.parametrize("spec, diagnostic", [
    (SPEC.replace("bar", "pie"), "chart_type"),
    (SPEC.replace("[100, -30]", "[100]"), "length"),
    (SPEC.replace("[100, -30]", "[100, nope]"), "numeric"),
    (SPEC.replace("[100, -30]", "[100, true]"), "numeric"),
    (SPEC.replace("[100, -30]", "[100, .nan]"), "finite"),
    (SPEC.replace("[100, -30]", "[100, .inf]"), "finite"),
    (SPEC + "surprise: 1\n", "surprise"),
    (SPEC + "unit: []\n", "unit"),
    (SPEC + "number_format: unknown\n", "number_format"),
    (SPEC + "type: line\n", "conflict"),
    ("title: Broken\nlabels: [\n", "YAML"),
])
def test_chart_invalid_fallback_and_strict_save_safety(spec, diagnostic, tmp_path):
    markdown = "```chart\n" + spec + "```\n"
    _, renderer, doc = render(markdown, charts=True)
    assert renderer.errors and diagnostic.lower() in ";".join(renderer.errors).lower()
    assert "Chart" in renderer.errors[0]
    if "가상 매출" in spec:
        assert "가상 매출" in renderer.errors[0]
    assert spec.rstrip() == doc.tables[0].cell(0, 0).text
    assert not doc.inline_shapes
    source, output = tmp_path / "bad.md", tmp_path / "result.docx"
    source.write_text(markdown, encoding="utf-8")
    for existing in (False, True):
        if existing:
            output.write_bytes(b"preserve existing")
        with pytest.raises(RuntimeError, match="Chart"):
            md_to_word.IBReportConverter(str(source), str(output), RenderOptions(
                profile="plain", charts=True, strict=True,
            )).convert()
        assert output.read_bytes() == b"preserve existing" if existing else not output.exists()
    _, disabled, code_doc = render(markdown, charts=False, strict=True)
    assert not disabled.errors and len(code_doc.tables) == 1


@pytest.mark.parametrize("failure", ["png", "insertion"])
def test_chart_render_failure_has_visible_code_and_diagnostic(monkeypatch, failure):
    import chart_renderer

    def fail(*args, **kwargs):
        raise OSError("injected failure")

    if failure == "png":
        monkeypatch.setattr(chart_renderer, "render_chart_png", fail)
    else:
        monkeypatch.setattr("docx.text.run.Run.add_picture", fail)
    _, renderer, doc = render(charts=True)
    assert "injected failure" in renderer.errors[0]
    assert SPEC.rstrip() == doc.tables[0].cell(0, 0).text
    assert not doc.inline_shapes
    with pytest.raises(ValueError, match="Chart.*injected failure"):
        render(charts=True, strict=True)


def test_waterfall_arithmetic_and_formats_do_not_rescale():
    from chart_renderer import format_value, parse_chart_spec, waterfall_positions

    assert waterfall_positions([100, -30, 20]) == ([0, 70, 70], [100, 100, 90], 90)
    assert waterfall_positions([-100, 30, -20]) == ([-100, -100, -90], [0, -70, -70], -90)
    assert waterfall_positions([0, -3.5, 3.5]) == ([0, -3.5, -3.5], [0, 0, 0], 0)
    for number_format, expected in [("percent", "12.5%"), ("bps", "12.5 bps"),
                                    ("multiple", "12.5x"), (",.2f", "12.50")]:
        spec = parse_chart_spec(SPEC + f"number_format: '{number_format}'\nunit: 금액\n")
        assert spec.series[0].values == [100, -30]
        assert format_value(12.5, spec.number_format) == expected


def test_chart_registry_and_cli_use_same_options(tmp_path, monkeypatch):
    source = tmp_path / "chart.md"
    source.write_text(CHART, encoding="utf-8")
    registry = get_default_registry()
    model = registry.convert(source, profile="plain")
    output = tmp_path / "registry.docx"
    registry.convert(model, output_format="docx", output_path=output, charts=True, strict=True)
    assert len(Document(output).inline_shapes) == 1
    for flag, count in [("--charts", 1), ("--no-charts", 0)]:
        source.write_text("---\ncharts: " + ("false" if count else "true") + "\n---\n" + CHART,
                          encoding="utf-8")
        output = tmp_path / (flag[2:] + ".docx")
        monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output),
                                         "--profile", "plain", flag, "--strict"])
        with pytest.raises(SystemExit) as exc:
            md_to_word.main()
        assert exc.value.code == 0
        assert len(Document(output).inline_shapes) == count


@pytest.mark.parametrize("value", ["yesplease", "1", "[]", "null"])
def test_chart_frontmatter_requires_boolean(value):
    with pytest.raises(ValueError, match="charts.*true or false"):
        render("---\ncharts: " + value + "\n---\n" + CHART)


@pytest.mark.parametrize("profile", PROFILES)
def test_disabled_chart_xml_is_identical_to_code_block(profile, monkeypatch):
    from document_model import CodeBlock

    monkeypatch.setattr("ib_renderer.uuid4", lambda: UUID(int=0))
    model = MarkdownParser(profile=profile).parse(
        "---\nsender: {organization: Example}\nrecipients: [Team]\n---\n" + CHART
    )
    chart_doc = IBDocumentRenderer().render(model)
    element = model.elements[0]
    element.element_type = ElementType.CODE_BLOCK
    element.content = CodeBlock(SPEC.rstrip(), "chart")
    code_doc = IBDocumentRenderer().render(model)
    assert chart_doc.element.xml == code_doc.element.xml
    assert chart_doc.styles.element.xml == code_doc.styles.element.xml
    assert chart_doc.part.numbering_part.element.xml == code_doc.part.numbering_part.element.xml


@pytest.mark.parametrize("profile, positive, negative", [
    ("ib-memo", "#003366", "#c00000"), ("plain", "#202020", "#404040"),
])
def test_waterfall_geometry_palette_and_unscaled_labels(profile, positive, negative, monkeypatch):
    from matplotlib.colors import to_hex
    from matplotlib.figure import Figure

    saved = Figure.savefig
    observations = []

    def inspect(self, *args, **kwargs):
        ax = self.axes[0]
        observations.append((
            [p.get_y() for p in ax.patches], [p.get_height() for p in ax.patches],
            [to_hex(p.get_facecolor()) for p in ax.patches], [t.get_text() for t in ax.texts],
        ))
        return saved(self, *args, **kwargs)

    monkeypatch.setattr(Figure, "savefig", inspect)
    _, _, doc = render(CHART.replace("bar", "waterfall").replace(
        "source: 가상 자료", "number_format: percent\nsource: 가상 자료"),
        profile=profile, charts=True, strict=True)
    assert len(doc.inline_shapes) == 1
    bottoms, heights, colors, labels = observations[0]
    assert bottoms == [0, 70, 0]
    assert heights == [100, 30, 70]
    assert colors[:2] == [positive, negative]
    assert labels == ["100%", "(30%)", "70%"]


def test_chart_failure_preserves_matplotlib_and_releases_figure(monkeypatch, tmp_path):
    import matplotlib
    import matplotlib.pyplot as plt
    from matplotlib.figure import Figure

    figures = plt.get_fignums()
    before = dict(matplotlib.rcParams)
    created = []

    def fail(self, destination, **kwargs):
        assert isinstance(destination, BytesIO)
        created.append(self)
        destination.write(b"partial image")
        raise OSError("injected PNG failure")

    monkeypatch.setattr(Figure, "savefig", fail)
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))
    _, renderer, doc = render(charts=True)
    assert renderer.errors and doc.tables
    assert dict(matplotlib.rcParams) == before
    assert plt.get_fignums() == figures
    assert not created[0].axes
    assert not list(tmp_path.iterdir())
