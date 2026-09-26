"""Preserve PR #2's blank-cell contract through parsing and DOCX saving."""

from io import BytesIO

import pytest
from docx import Document

from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, TableParser


@pytest.mark.parametrize(
    ("lines", "expected"),
    [
        (
            [
                "|| Amount | Share |",
                "|---|---|---|",
                "| Tranche A | 7,000 | 63.6% |",
            ],
            [["", "Amount", "Share"], ["Tranche A", "7,000", "63.6%"]],
        ),
        (
            [
                "| Item | Amount | Notes |",
                "|---|---|---|",
                "| Tranche B | | Pending |",
            ],
            [["Item", "Amount", "Notes"], ["Tranche B", "", "Pending"]],
        ),
        (
            [
                "| Item | Amount | Notes |",
                "|---|---|---|",
                "| Tranche B | 2,000 | |",
            ],
            [["Item", "Amount", "Notes"], ["Tranche B", "2,000", ""]],
        ),
    ],
    ids=["blank-leading-header", "blank-interior-cell", "blank-trailing-cell"],
)
def test_blank_cells_keep_their_columns_after_docx_save(lines, expected):
    table = TableParser.parse(lines)
    assert table.col_count == 3
    assert [[cell.content for cell in row.cells] for row in table.rows] == expected

    model = MarkdownParser(profile="plain").parse("\n".join(lines))
    document = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    reopened = Document(output)
    assert len(reopened.tables) == 1
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == expected
