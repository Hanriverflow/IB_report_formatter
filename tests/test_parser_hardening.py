"""Parser regressions verified through the production DOCX render path."""

import base64
from pathlib import Path
from typing import List, Optional, Tuple

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn

from document_model import CodeBlock, DocumentMetadata, DocumentModel, Element, ElementType, Heading
from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser


def _render(
    markdown: str, profile: str = "plain", strict: bool = True,
    cover: Optional[bool] = None,
) -> Tuple[DocumentModel, DocxDocument]:
    """Parse and render with matching profiles and explicit strictness."""
    model = MarkdownParser(profile=profile).parse(markdown)
    doc = IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=strict, include_cover=cover,
        include_toc=False, include_disclaimer=False,
    )).render(model)
    return model, doc


def _texts(doc: DocxDocument) -> List[str]:
    """Include hyperlink text and text inside tables in source order."""
    return [
        "".join("\n" if node.tag == qn("w:br") else (node.text or "")
                for node in p.iter() if node.tag in {qn("w:t"), qn("w:br")})
        for p in doc.element.body.iter(qn("w:p"))
    ]


@pytest.mark.parametrize("template", ["{}", "- {}", "## {}", "| A |\n|---|\n| {} |"])
def test_p1a_escaped_dollars_survive_all_inline_contexts(template: str) -> None:
    _, doc = _render(template.format(r"Pay \$5 and \$10 today."))
    assert "Pay $5 and $10 today." in _texts(doc)
    assert not doc.inline_shapes


@pytest.mark.parametrize("template", ["{}", "- {}", "| A |\n|---|\n| {} |"])
def test_p1b_escaped_emphasis_stays_literal(template: str) -> None:
    _, doc = _render(template.format(r"Show \*\*literal\*\* stars."))
    assert "Show **literal** stars." in _texts(doc)


def test_p1c_references_end_at_first_ordinary_block() -> None:
    model, doc = _render(
        "# T\n\nIntro.\n\n## References\n\n1. Source A\n2. Source B\n\n"
        "IMPORTANT CAVEAT paragraph.\n\n| AAA | BBB |\n|---|---|\n| xq | yq |\n",
        profile="ib-report",
    )
    assert "IMPORTANT CAVEAT paragraph." in _texts(doc)
    assert [c.text for c in doc.tables[-1].rows[-1].cells] == ["xq", "yq"]
    assert model.footnotes == {1: "Source A", 2: "Source B"}
    assert not model.warnings


def test_p1d_unknown_header_label_remains_in_body() -> None:
    model, doc = _render("# Report\n\n**Risk:** Total loss is possible.\n\nBody.\n", "ib-report")
    assert "Risk: Total loss is possible." in _texts(doc)
    assert "Risk" not in model.metadata.extra


def test_p1d_recognized_metadata_is_rendered_once() -> None:
    model, doc = _render("# Report\n\n**작성일:** 2026-09-29\n\nBody.\n", "ib-report")
    assert _texts(doc).count("2026-09-29") == 1
    assert not any(text.startswith("작성일:") for text in _texts(doc))
    assert model.metadata.extra["date"] == "2026-09-29"


@pytest.mark.parametrize("cover", [False, True])
def test_p1d_header_h2_survives_cover_choice(cover: bool) -> None:
    _, doc = _render("# Report\n\n## Critical scope\n\nBody.\n", "ib-report", cover=cover)
    assert _texts(doc).count("Critical scope") == 1


def test_p1e_dash_only_body_row_is_preserved(tmp_path: Path) -> None:
    _, doc = _render("| A | B |\n|---|---|\n| x | 1 |\n| - | - |\n| z | 3 |")
    target = tmp_path / "dashes.docx"
    doc.save(str(target))
    reopened = Document(str(target))
    assert [[c.text for c in r.cells] for r in reopened.tables[0].rows] == [
        ["A", "B"], ["x", "1"], ["-", "-"], ["z", "3"],
    ]


@pytest.mark.parametrize("source,expected", [
    ("Use `[^7]` as an example.", "Use `[^7]` as an example."),
    (r"Use \[^7] as an example.", "Use [^7] as an example."),
])
def test_p2a_literal_footnote_has_no_warning(source: str, expected: str) -> None:
    model, doc = _render(source)
    assert not model.warnings
    assert expected in _texts(doc)
    assert not doc.element.body.xpath(".//w:footnoteReference")


@pytest.mark.parametrize("source,expected", [
    ("Prose\n```python\nx = 1\n```\n\nAfter.", "x = 1"),
    ("~~~python\nx = 1\n~~~\n\nAfter.", "x = 1"),
    ("````python\nx = 1\n```\n~~~~\n````info\n`````\n\nAfter.",
     "x = 1\n```\n~~~~\n````info"),
    ("~~~~\n[^7]: literal\n~~~\n```\n~~~~~\n\nAfter.", "[^7]: literal\n~~~\n```"),
])
def test_p2b_shared_fence_boundaries(source: str, expected: str) -> None:
    model, doc = _render(source)
    blocks = [e.content for e in model.elements if isinstance(e.content, CodeBlock)]
    assert len(blocks) == 1
    assert blocks[0].code == expected
    assert expected in _texts(doc) or expected.splitlines()[0] in _texts(doc)
    assert "After." in _texts(doc)
    assert not model.footnotes


@pytest.mark.parametrize("newline", ["\r\n", "\r"])
def test_p2c_windows_and_old_mac_newlines(newline: str) -> None:
    _, doc = _render(newline.join(["- First", "- Second", ""]))
    assert _texts(doc) == ["First", "Second"]


@pytest.mark.parametrize("continuation", ["    continued explanation", "continued explanation"])
def test_p2d_list_continuation_keeps_numbering(continuation: str) -> None:
    _, doc = _render("1. First line\n" + continuation + "\n2. Second item")
    assert _texts(doc) == ["First line continued explanation", "Second item"]
    assert [p._p.pPr.numPr.numId.val for p in doc.paragraphs] == [
        doc.paragraphs[0]._p.pPr.numPr.numId.val,
        doc.paragraphs[0]._p.pPr.numPr.numId.val,
    ]


@pytest.mark.parametrize("angled", [False, True])
def test_p2e_image_destination_excludes_title(tmp_path: Path, angled: bool) -> None:
    image = tmp_path / ("chart with spaces.png" if angled else "chart.png")
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aWZkAAAAASUVORK5CYII="
    ))
    destination = f"<{image.as_posix()}>" if angled else image.as_posix()
    _, doc = _render(f'![Chart]({destination} "Revenue")')
    target = tmp_path / "image.docx"
    doc.save(str(target))
    assert len(Document(str(target)).inline_shapes) == 1


def test_p2f_setext_equals_heading_preserves_dash_separator() -> None:
    model, doc = _render("Title\n===\n\nBody.\n---\nAfter.")
    assert doc.paragraphs[0].text == "Title"
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert sum(e.element_type == ElementType.SEPARATOR for e in model.elements) == 1


def test_p2f_comments_are_hidden_outside_code() -> None:
    model, doc = _render(
        "Before <!-- hidden -->after.\n\n<!-- multi\nline -->\n\n"
        "`<!-- literal -->`\n\n```\n<!-- code -->\n```\n\nEnd."
    )
    assert [text for text in _texts(doc) if text] == [
        "Before after.", "`<!-- literal -->`", "<!-- code -->", "End.",
    ]
    assert not model.warnings


def test_p2f_reference_links_resolve_and_definitions_disappear(tmp_path: Path) -> None:
    _, doc = _render(
        "Use [docs][1], [docs][], [docs], [missing][no] and ` [docs][1] `.\n\n"
        '[1]: https://x.test "opt title"\n[docs]: https://docs.test\n'
    )
    target = tmp_path / "links.docx"
    doc.save(str(target))
    reopened = Document(str(target))
    assert _texts(reopened) == ["Use docs, docs, docs, [missing][no] and ` [docs][1] `."]
    links = reopened.element.body.xpath(".//w:hyperlink")
    assert len(links) == 3
    assert [reopened.part.rels[link.get(qn("r:id"))].target_ref for link in links] == [
        "https://x.test", "https://docs.test", "https://docs.test",
    ]


def test_p2a_multiline_code_span_does_not_define_or_reference_footnotes() -> None:
    model, doc = _render("Use `literal\n[^7]\n[^8]: example\ncode` here.")
    assert not model.warnings
    assert not model.footnotes
    assert "Use `literal [^7] [^8]: example code` here." in _texts(doc)


def test_p2d_list_continuation_preserves_inline_code_and_hard_breaks() -> None:
    _, doc = _render("- First `  a  b  `\n    continued<br>\n    next\n    - Child\n- Last")
    assert _texts(doc) == ["First `  a  b  ` continued\nnext", "Child", "Last"]
    assert doc.paragraphs[1]._p.pPr.numPr.ilvl.val == 1


def test_p1b_escaped_emphasis_does_not_rewrite_link_destination() -> None:
    destination = r"https://x.test/\_literal\*"
    _, doc = _render(r"Show \*literal\* [docs](" + destination + ").")
    assert _texts(doc) == ["Show *literal* docs."]
    link = doc.element.body.xpath(".//w:hyperlink")[0]
    assert doc.part.rels[link.get(qn("r:id"))].target_ref == destination


@pytest.mark.parametrize("profile", ["ib-report", "ib-memo"])
@pytest.mark.parametrize("cover", [True, False])
@pytest.mark.parametrize("toc", [True, False])
def test_d1_inferred_subtitle_follows_cover_and_toc(
    tmp_path: Path, profile: str, cover: bool, toc: bool,
) -> None:
    model = MarkdownParser(profile=profile).parse(
        "# 예시기업 ABCP 구조\n## 투자 검토 요약\n\n본문 문단.\n\n## 본문 범위\n\n검토 사항.\n"
    )
    renderer = IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=True, include_cover=cover,
        include_toc=toc, include_disclaimer=False,
    ))
    target = tmp_path / "subtitle.docx"
    renderer.render(model).save(str(target))
    doc = Document(str(target))
    matches = [p for p in doc.paragraphs if p.text == "투자 검토 요약"]
    assert [p.text for p in matches if p.style.name == "Heading 2"] == (
        [] if cover else ["투자 검토 요약"]
    )
    assert len(matches) == (1 if cover or not toc else 2)
    assert model.metadata.subtitle == "투자 검토 요약"
    assert sum(
        isinstance(e.content, Heading) and e.content.text == "투자 검토 요약"
        for e in model.elements
    ) == 1
    assert _texts(doc).count("본문 범위") == (2 if toc else 1)
    assert not any("**" in text for text in _texts(doc))
    # Rendering the same model with the other cover setting must retain the H2.
    other = IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=True, include_cover=not cover,
        include_toc=False, include_disclaimer=False,
    )).render(model)
    assert sum(p.text == "투자 검토 요약" and p.style.name == "Heading 2"
               for p in other.paragraphs) == int(cover)


@pytest.mark.parametrize("profile", ["ib-report", "ib-memo"])
@pytest.mark.parametrize("cover", [True, False])
def test_d1_explicit_yaml_subtitle_keeps_h2_in_body(profile: str, cover: bool) -> None:
    model, doc = _render(
        '---\nsubtitle: "명시적 부제"\n---\n# 보고서\n## 투자 검토 요약\n\n본문.\n',
        profile, cover=cover,
    )
    assert model.metadata.subtitle == "명시적 부제"
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == ["투자 검토 요약"]
    assert _texts(doc).count("명시적 부제") == int(cover)


@pytest.mark.parametrize("profile", ["plain", "office-letter", "business-report", "meeting-minutes"])
@pytest.mark.parametrize("cover", [True, False])
def test_d1_general_profiles_do_not_infer_subtitle(profile: str, cover: bool) -> None:
    model = MarkdownParser(profile=profile).parse("# 보고서\n## 투자 검토 요약\n\n본문.\n")
    if profile == "office-letter":
        model.metadata.extra.update(recipients=["수신인"], sender={"organization": "발신기관"})
    doc = IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=True, include_cover=cover,
        include_toc=False, include_disclaimer=False,
    )).render(model)
    assert model.metadata.subtitle == ""
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == ["투자 검토 요약"]
    assert _texts(doc).count("투자 검토 요약") == 1


@pytest.mark.parametrize("cover", [True, False])
def test_d1_hand_built_model_does_not_hide_matching_heading(cover: bool) -> None:
    model = DocumentModel(
        metadata=DocumentMetadata(title="Report", subtitle="Scope"),
        elements=[Element(ElementType.HEADING_2, Heading(level=2, text="Scope"))],
    )
    doc = IBDocumentRenderer(options=RenderOptions(
        profile="ib-report", strict=True, include_cover=cover,
        include_toc=False, include_disclaimer=False,
    )).render(model)
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 2"] == ["Scope"]


@pytest.mark.parametrize("labels", [("Date", "Analyst"), ("작성일", "작성자")])
@pytest.mark.parametrize("cover", [True, False])
@pytest.mark.parametrize("strict", [True, False])
def test_d2_consecutive_header_metadata_has_separate_cover_cells(
    tmp_path: Path, labels: Tuple[str, str], cover: bool, strict: bool,
) -> None:
    date_label, analyst_label = labels
    model, doc = _render(
        f"# Title\n\n**{date_label}:** 2026-09-01\n**{analyst_label}:** Kim\n\nBody.\n",
        "ib-report", strict=strict, cover=cover,
    )
    target = tmp_path / "metadata.docx"
    doc.save(str(target))
    doc = Document(str(target))
    cells = {row.cells[0].text: row.cells[1].text for table in doc.tables for row in table.rows}
    if cover:
        assert cells["REPORT DATE"] == "2026-09-01"
        assert cells["PREPARED BY"] == "Kim"
    else:
        assert not cells
    assert model.metadata.extra["date"] == "2026-09-01"
    assert model.metadata.analyst == "Kim"
    assert _texts(doc).count("2026-09-01") == int(cover)
    assert _texts(doc).count("Kim") == int(cover)
    assert "Body." in [p.text for p in doc.paragraphs]
    assert not any(p.style.name.startswith("Heading") for p in doc.paragraphs)
    assert not any("**" in text for text in _texts(doc))


@pytest.mark.parametrize("unknown_index", [0, 1, 2])
@pytest.mark.parametrize("labels", [("Date", "Analyst", "Risk"), ("작성일", "작성자", "위험")])
@pytest.mark.parametrize("cover", [True, False])
def test_d2_unknown_labels_survive_among_consecutive_metadata(
    labels: Tuple[str, str, str], unknown_index: int, cover: bool,
) -> None:
    date_label, analyst_label, risk_label = labels
    lines = [f"**{date_label}:** 2026-09-01", f"**{analyst_label}:** Kim"]
    lines.insert(unknown_index, f"**{risk_label}:** Total loss is possible.")
    model, doc = _render("# Title\n\n" + "\n".join(lines) + "\n\nBody.\n", "ib-report", cover=cover)
    assert _texts(doc).count(f"{risk_label}: Total loss is possible.") == 1
    risk = next(p for p in doc.paragraphs if p.text == f"{risk_label}: Total loss is possible.")
    assert risk.runs[0].bold
    assert model.metadata.extra == {"date": "2026-09-01"}
    assert model.metadata.analyst == "Kim"
    assert "Body." in _texts(doc)
    assert not any("**" in text for text in _texts(doc))


@pytest.mark.parametrize("cover", [True, False])
def test_d2_metadata_stops_at_ordinary_body_text(cover: bool) -> None:
    model, doc = _render(
        "# Title\n\n**Date:** 2026-09-01\n**Analyst:** Kim\n"
        "Body starts here\nand continues.\n\n**Date:** body example\n",
        "ib-report", cover=cover,
    )
    assert "Body starts here and continues." in _texts(doc)
    assert "Date: body example" in _texts(doc)
    assert model.metadata.extra["date"] == "2026-09-01"
    assert model.metadata.analyst == "Kim"


@pytest.mark.parametrize("profile", ["plain", "office-letter", "business-report", "meeting-minutes"])
def test_d2_general_profiles_keep_consecutive_labels_in_body(profile: str) -> None:
    model = MarkdownParser(profile=profile).parse(
        "# Title\n\n**Date:** 2026-09-01\n**Analyst:** Kim\n\nBody.\n"
    )
    if profile == "office-letter":
        model.metadata.extra.update(recipients=["Reader"], sender={"organization": "Sender"})
    doc = IBDocumentRenderer(options=RenderOptions(profile=profile, strict=True)).render(model)
    assert "Date: 2026-09-01 Analyst: Kim" in _texts(doc)
    assert "date" not in model.metadata.extra
    assert model.metadata.analyst == ""


def test_d2_yaml_metadata_does_not_consume_body_labels() -> None:
    model, doc = _render(
        '---\ntitle: Title\ndate: "2026-09-02"\nanalyst: Lee\n---\n'
        "**Date:** 2026-09-01\n**Analyst:** Kim\n\nBody.\n", "ib-report",
    )
    assert "Date: 2026-09-01 Analyst: Kim" in _texts(doc)
    assert model.metadata.extra["date"] == "2026-09-02"
    assert model.metadata.analyst == "Lee"


@pytest.mark.parametrize("cover", [True, False])
def test_d2_unknown_label_keeps_soft_wrapped_body_paragraph(cover: bool) -> None:
    model, doc = _render(
        "# Title\n\n**Date:** 2026-09-01\n**Analyst:** Kim\n"
        "**Risk:** Total loss\nis possible.\n\nBody.\n", "ib-report", cover=cover,
    )
    assert "Risk: Total loss is possible." in _texts(doc)
    assert model.metadata.extra["date"] == "2026-09-01"
    assert model.metadata.analyst == "Kim"
