"""Parse-to-DOCX regressions for rendering and structural audit hardening."""

from dataclasses import asdict
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.styles.styles import Styles

from document_profiles import RenderOptions
from docx_audit import inspect_document
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser
from md_to_word import IBReportConverter

DIAGRAM = """```diagram:flow
title: Approval
boxes:
  - id: a
    label: Start
    pos: [0, 0]
  - id: b
    label: End
    pos: [3, 0]
arrows:
  - from: a
    to: b
```
"""


def render(markdown: str, profile: str = "plain", strict: bool = True) -> DocxDocument:
    """Parse and render with matching explicit profile settings."""
    model = MarkdownParser(profile=profile).parse(markdown)
    return IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=strict,
        include_cover=False, include_toc=False, include_disclaimer=False,
    )).render(model)


def test_numeric_cell_preserves_native_footnote(tmp_path: Path) -> None:
    """A footnote number must never become part of a financial amount."""
    doc = render(
        "| Metric | 2026 |\n|---|---:|\n| Revenue | 1234[^1] |\n\n"
        "[^1]: Audited amount.", profile="ib-report",
    )
    output = tmp_path / "numeric-footnote.docx"
    doc.save(output)
    reopened = Document(output)
    cell = reopened.tables[0].cell(1, 1)
    assert cell.text == "1,234"
    refs = cell._tc.xpath(".//w:footnoteReference")
    assert len(refs) == 1
    assert refs[0].get(qn("w:id")) == "1"
    assert b"Audited amount." in reopened.part.part_related_by(RT.FOOTNOTES).blob


def test_numeric_runs_keep_hyperlinks_spacing_and_explicit_value_scale(tmp_path: Path) -> None:
    """Presentation roles affect numeric runs without dropping links, notes or spaces."""
    doc = render(
        "---\ntables:\n  - type: financial\n    columns: [text, money, percent, bps, multiple]\n---\n"
        "| Metric | Amount | Rate | Spread | Multiple |\n|---|---|---|---|---|\n"
        "| Revenue | 1234 [source](https://example.com)[^1] | **12.50**[^1] | 100.00[^1] | 2.50[^1] |\n\n"
        "[^1]: Audited amount.", profile="ib-report",
    )
    output = tmp_path / "semantic-numbers.docx"
    doc.save(output)
    cells = Document(output).tables[0].rows[1].cells
    assert ["".join(cell._tc.xpath(".//w:t/text()")) for cell in cells] == [
        "Revenue", "1,234 source", "12.50%", "100.00 bps", "2.50x",
    ]
    assert len(cells[1]._tc.xpath(".//w:hyperlink")) == 1
    assert len(cells[1]._tc.xpath(".//w:footnoteReference")) == 1
    assert cells[2].paragraphs[0].runs[0].bold


def test_plain_numeric_cell_with_footnote_keeps_original_presentation() -> None:
    """The general profile must not infer financial presentation from numeric text."""
    doc = render("| Metric | 2026 |\n|---|---:|\n| Revenue | 1234[^1] |\n\n[^1]: Note.")
    cell = doc.tables[0].cell(1, 1)
    assert cell.text == "1234"
    assert len(cell._tc.xpath(".//w:footnoteReference")) == 1


@pytest.mark.parametrize("markdown, expression", [
    (r"$$ \notacommand{x} $$", r"\notacommand{x}"),
    (r"Before $\badcmd$ after.", r"\badcmd"),
    (r"## Heading $\badcmd$", r"\badcmd"),
    (r"- Item $\badcmd$", r"\badcmd"),
    ("| Value |\n|---|\n| $\\badcmd$ |", r"\badcmd"),
])
def test_equation_failure_is_recorded_and_strict_rejects(
    markdown: str, expression: str, tmp_path: Path,
) -> None:
    """Invalid math stays visible in non-strict output and rejects strict output."""
    model = MarkdownParser(profile="plain").parse(markdown)
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", strict=False))
    doc = renderer.render(model)
    assert renderer.errors
    output = tmp_path / "fallback.docx"
    doc.save(output)
    reopened = Document(output)
    assert expression in "".join(reopened.element.xpath(".//w:t/text()"))
    assert len(reopened.inline_shapes) == 0
    with pytest.raises(ValueError, match="LaTeX"):
        render(markdown)


@pytest.mark.parametrize("markdown", [r"$$ \badcmd $$", DIAGRAM])
def test_failed_image_render_releases_figures_and_temp_files(
    markdown: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A failed save must not leak a pyplot figure or even a partial PNG."""
    import tempfile

    import matplotlib.pyplot as plt
    from matplotlib.figure import Figure

    def broken_savefig(self, filename, *args, **kwargs) -> None:
        Path(filename).write_bytes(b"partial PNG")
        raise OSError("injected image write failure")

    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setattr(Figure, "savefig", broken_savefig)
    figures = plt.get_fignums()
    render(markdown, strict=False)
    assert plt.get_fignums() == figures
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("markdown", [r"$$ \notacommand{x} $$", r"Before $\badcmd$ after."])
@pytest.mark.parametrize("existing_output", [False, True])
def test_strict_equation_failure_never_saves_output(
    markdown: str, existing_output: bool, tmp_path: Path,
) -> None:
    """The real converter must reject before creating or overwriting a destination."""
    source, output = tmp_path / "invalid.md", tmp_path / "result.docx"
    source.write_text(markdown, encoding="utf-8")
    if existing_output:
        output.write_bytes(b"existing output must survive")
    with pytest.raises(RuntimeError, match="LaTeX"):
        IBReportConverter(str(source), str(output), RenderOptions(profile="plain", strict=True)).convert()
    if existing_output:
        assert output.read_bytes() == b"existing output must survive"
    else:
        assert not output.exists()


@pytest.mark.parametrize("markdown", [r"$$ x^2 $$", r"Before $x^2$ after."])
def test_equation_insertion_failure_is_visible_and_rejects_strict(
    markdown: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Insertion failures use the same diagnostic path and still remove the image."""
    import tempfile

    from docx.text.run import Run

    def broken_insert(*args, **kwargs) -> None:
        raise OSError("injected DOCX image insertion failure")

    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setattr(Run, "add_picture", broken_insert)
    model = MarkdownParser(profile="plain").parse(markdown)
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", strict=False))
    doc = renderer.render(model)
    assert renderer.errors
    assert "[x^2]" in "".join(doc.element.xpath(".//w:t/text()"))
    with pytest.raises(ValueError, match="LaTeX"):
        render(markdown)
    assert list(tmp_path.iterdir()) == []


def test_equation_diagnostics_are_reset_on_renderer_reuse() -> None:
    """A failed render cannot taint the next request handled by the same renderer."""
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", strict=True))
    with pytest.raises(ValueError, match="LaTeX"):
        renderer.render(MarkdownParser(profile="plain").parse(r"$\badcmd$"))
    doc = renderer.render(MarkdownParser(profile="plain").parse("Next document."))
    assert renderer.errors == []
    assert "Next document." in [p.text for p in doc.paragraphs]


def representative_audit_document() -> DocxDocument:
    """Include tables, style inheritance, local overrides and document defaults."""
    doc = render(
        "## Heading\n\nBody [source](https://example.com)[^1].\n\n"
        "- First\n    1. Nested\n- Second\n\n"
        "| Item | Value |\n|---|---|\n| A | 10 |\n\n[^1]: Note."
    )
    parent = doc.styles.add_style("Audit Parent", WD_STYLE_TYPE.PARAGRAPH)
    parent.paragraph_format.keep_together = True
    child = doc.styles.add_style("Audit Child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    doc.paragraphs[1].style = child
    cell = doc.tables[0].cell(1, 1).paragraphs[0]
    cell.style = child
    cell.paragraph_format.keep_together = False
    doc.styles["Normal"].paragraph_format.keep_with_next = True
    default = doc.styles.element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")[0]
    flag = OxmlElement("w:pageBreakBefore")
    default.append(flag)
    doc.paragraphs[1].paragraph_format.page_break_before = False
    return doc


def test_audit_preserves_results_without_per_paragraph_style_lookups(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    """Full audit observations stay stable while style resolution is indexed once."""
    doc = representative_audit_document()
    output = tmp_path / "audit.docx"
    doc.save(output)
    reopened = Document(output)
    # Snapshot captured with the original audit before optimizing style lookup.
    expected = {
        "paragraphs": 6, "tables": 1, "headings": 1, "images": 0,
        "numbered_paragraphs": 3, "issues": [],
        "warnings": [
            "Normal style has pagination constraints: these can show nonprinting square "
            "marks and create large page gaps. Keep such controls on necessary headings only."
        ],
        "pagination_marked_paragraphs": 10, "visual_review": "not_performed",
    }
    original = Styles.get_by_id
    calls = []

    def counted(self, *args, **kwargs):
        calls.append(args)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Styles, "get_by_id", counted)
    assert asdict(inspect_document(reopened)) == expected
    assert len(calls) <= 1


def test_render_skips_audit_observation_pass(monkeypatch: pytest.MonkeyPatch) -> None:
    """Rendering validates structure without resolving observational pagination."""
    import docx_audit

    def unexpected(*args, **kwargs):
        raise AssertionError("pagination observations should not run during rendering")

    monkeypatch.setattr(docx_audit, "_effective_pagination_flag", unexpected)
    doc = render("## Heading\n\nBody.\n\n| Item |\n|---|\n| A |")
    assert len(doc.tables) == 1


@pytest.mark.parametrize("literal", ["[Image: example]", "[Render Error: example]", "[Diagram: example]"])
def test_placeholder_literals_are_audit_warnings_not_render_errors(
    literal: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """User code containing diagnostic-like text must pass strict rendering and CLI audit."""
    import logging

    import docx_audit

    doc = render("```text\n" + literal + "\n```")
    output = tmp_path / "literal.docx"
    doc.save(output)
    report = inspect_document(Document(output))
    assert report.issues == []
    assert any(literal in warning for warning in report.warnings)
    monkeypatch.setattr("sys.argv", ["docx-audit", str(output)])
    with caplog.at_level(logging.INFO), pytest.raises(SystemExit) as exc:
        docx_audit.main()
    assert exc.value.code == 0
    assert '"issues": []' in caplog.text


def test_actual_image_failure_still_rejects_strict_render() -> None:
    """Ignoring placeholder-like user text must not hide real rendering failures."""
    with pytest.raises(ValueError, match="Image"):
        render("![Missing](definitely-not-an-existing-image.png)")


@pytest.mark.parametrize("location", ["body", "table", "style", "header"])
def test_audit_rejects_dangling_numbering_references(location: str, tmp_path: Path) -> None:
    """All paragraph numbering references, including other parts, must resolve."""
    doc = render("1. First\n\n| Item |\n|---|\n| A |")
    target = {
        "body": doc.paragraphs[0]._p,
        "table": doc.tables[0].cell(1, 0).paragraphs[0]._p,
        "style": doc.styles["Normal"].element,
        "header": doc.sections[0].header.paragraphs[0]._p,
    }[location]
    target.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 99999
    output = tmp_path / "dangling.docx"
    doc.save(output)
    assert any("99999" in issue for issue in inspect_document(Document(output)).issues)


@pytest.mark.parametrize("damage", ["missing-abstract", "missing-abstract-reference", "missing-part"])
def test_audit_rejects_broken_numbering_definitions(damage: str, tmp_path: Path) -> None:
    """A numId is usable only when its numbering part and abstract definition exist."""
    doc = render("1. First")
    num_id = doc.paragraphs[0]._p.pPr.numPr.numId.val
    numbering = doc.part.numbering_part.element
    num = next(node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == str(num_id))
    abstract_ref = num.find(qn("w:abstractNumId"))
    if damage == "missing-abstract":
        abstract_ref.set(qn("w:val"), "99999")
    elif damage == "missing-abstract-reference":
        num.remove(abstract_ref)
    else:
        relation = next(key for key, value in doc.part.rels.items() if value.reltype == RT.NUMBERING)
        doc.part.drop_rel(relation)
    output = tmp_path / "broken-numbering.docx"
    doc.save(output)
    assert inspect_document(Document(output)).issues


def test_audit_accepts_zero_numbering_and_valid_nested_lists(tmp_path: Path) -> None:
    """numId zero means no numbering; native multilevel references remain valid."""
    doc = render("1. First\n    1. Nested\n2. Next\n\nPlain paragraph.")
    doc.paragraphs[-1]._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    output = tmp_path / "valid-numbering.docx"
    doc.save(output)
    assert inspect_document(Document(output)).issues == []


@pytest.mark.parametrize("markdown, child_level, start", [
    ("- A\n    1. First\n- B\n    1. Second", 1, 1),
    ("1. A\n    3. First\n2. B\n    3. Second", 1, 3),
    ("- A\n    - Child A\n        7. First\n- B\n    - Child B\n        7. Second", 2, 7),
])
def test_nested_numbering_restarts_in_each_parent_scope(
    markdown: str, child_level: int, start: int, tmp_path: Path,
) -> None:
    """Separate parent items own distinct nested instances at the correct level."""
    doc = render(markdown)
    output = tmp_path / "scoped-lists.docx"
    doc.save(output)
    reopened = Document(output)
    children = [p for p in reopened.paragraphs if p.text in {"First", "Second"}]
    ids = [p._p.pPr.numPr.numId.val for p in children]
    assert len(set(ids)) == 2
    numbering = reopened.part.numbering_part.element
    for paragraph, num_id in zip(children, ids):
        assert paragraph._p.pPr.numPr.ilvl.val == child_level
        instance = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id))
        override = instance.find(qn("w:lvlOverride"))
        assert override is not None
        assert override.get(qn("w:ilvl")) == str(child_level)
        assert override.find(qn("w:startOverride")).get(qn("w:val")) == str(start)


def test_ordered_list_scope_preserves_sibling_and_top_level_continuity() -> None:
    """Nested siblings continue, and returning to a parent level does not restart it."""
    doc = render("3. A\n    1. First\n    2. Next\n4. B\n    1. Second")
    paragraphs = {p.text: p for p in doc.paragraphs}
    ids = {text: p._p.pPr.numPr.numId.val for text, p in paragraphs.items()}
    assert ids["A"] == ids["B"]
    assert ids["First"] == ids["Next"]
    assert ids["First"] != ids["Second"]
    instance = next(n for n in doc.part.numbering_part.element.findall(qn("w:num")) if n.get(qn("w:numId")) == str(ids["A"]))
    override = instance.find(qn("w:lvlOverride"))
    assert override.get(qn("w:ilvl")) == "0"
    assert override.find(qn("w:startOverride")).get(qn("w:val")) == "3"


@pytest.mark.parametrize("markdown", [
    DIAGRAM, r"$$ x^2 + y^2 = z^2 $$", r"$$ \text{가용자본} + \alpha $$",
    "---\ncharts: true\n---\n```chart\ntype: bar\nlabels: [상반기, 하반기]\n"
    "series: [{name: 매출, values: [100, -30]}]\n```",
])
def test_rendered_diagrams_and_equations_do_not_mutate_matplotlib(
    markdown: str, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Rendering must preserve the caller's backend, fonts and other rcParams."""
    import matplotlib

    monkeypatch.setitem(matplotlib.rcParams, "backend", "svg")
    monkeypatch.setitem(matplotlib.rcParams, "font.family", ["DejaVu Sans"])
    monkeypatch.setitem(matplotlib.rcParams, "axes.unicode_minus", True)
    before = dict(matplotlib.rcParams)
    doc = render(markdown)
    assert len(doc.inline_shapes) == 1
    assert dict(matplotlib.rcParams) == before
