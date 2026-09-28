"""Typed theme extension and request isolation through serialized DOCX output."""

import sys
from concurrent.futures import ThreadPoolExecutor
from dataclasses import fields
from io import BytesIO

import pytest
import yaml
from docx import Document
from docx.shared import Inches, Pt, RGBColor

import md_to_word
from document_profiles import RenderOptions, get_profile, load_style
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser
from render_styles import STYLE, IBStyle


def themed_document(path):
    model = MarkdownParser().parse("---\ntitle: Test\n---\n# Heading\n\n> [요약]\n> 내용\n\n- Bullet")
    doc = IBDocumentRenderer(options=RenderOptions(
        theme=str(path), include_cover=False, include_toc=False, include_disclaimer=False,
    )).render(model)
    data = BytesIO()
    doc.save(data)
    return Document(data)


def test_theme_heading_callout_and_bullet_honor_uppercase_fields(tmp_path):
    theme = tmp_path / "extended.yaml"
    theme.write_text('NAVY: "123456"\nH1_SIZE: 23\nBODY_FONT: Arial\nTOP_MARGIN: 0.6\n', encoding="utf-8")
    doc = themed_document(theme)
    heading = next(p for p in doc.paragraphs if p.text == "Heading")
    assert heading.runs[0].font.color.rgb == RGBColor.from_string("123456")
    assert heading.runs[0].font.size == Pt(23)
    assert doc.sections[0].top_margin == Inches(0.6)
    assert 'w:fill="123456"' in doc.tables[0]._tbl.xml
    bullet = next(p for p in doc.paragraphs if "Bullet" in p.text)
    assert bullet.runs[0].font.color.rgb == RGBColor.from_string("123456")
    assert STYLE.NAVY == IBStyle().NAVY


@pytest.mark.parametrize("field, value", [
    ("primary_color", 123456), ("NAVY", True), ("NAVY_HEX", "bad"),
    ("TABLE_HEADER_BG", 123456), ("BODY_SIZE", True), ("H1_SIZE", -1),
    ("H2_SIZE", float("inf")), ("BODY_LINE_SPACING", float("nan")),
    ("BODY_LINE_SPACING", 0), ("BODY_FONT", 123), ("KOREAN_FONT", ""),
    ("TOP_MARGIN", -1), ("TABLE_ZEBRA", "true"),
    ("FULL_LIST_INDENT_LEVELS", True), ("FULL_LIST_INDENT_LEVELS", -1),
    ("no_such_key", 1), (123, "bad"),
])
def test_invalid_theme_rejected_atomically(field, value, tmp_path):
    theme = tmp_path / "invalid.yaml"
    theme.write_text(yaml.safe_dump({"NAVY": "123456", field: value}), encoding="utf-8")
    with pytest.raises(ValueError, match=str(field)):
        themed_document(theme)
    assert STYLE.NAVY == IBStyle().NAVY


def test_all_presentation_fields_are_supported(tmp_path):
    style = IBStyle()
    data = {}
    for field in fields(style):
        if field.name.startswith("STYLE_") or field.name == "NATIVE_NUMBERING":
            continue  # Internal identifiers and profile numbering policy are not visual themes.
        value = getattr(style, field.name)
        if isinstance(value, RGBColor):
            value = str(value)
        elif isinstance(value, Pt):
            value = value.pt
        elif isinstance(value, Inches):
            value = value.inches
        data[field.name] = value
    path = tmp_path / "complete.yaml"
    path.write_text(yaml.safe_dump(data), encoding="utf-8")
    assert load_style(get_profile("ib-report"), str(path)) == style
    themed_document(path)


def test_code_background_honors_theme(tmp_path):
    theme = tmp_path / "code.yaml"
    theme.write_text('CODE_BG: "123456"\n', encoding="utf-8")
    doc = IBDocumentRenderer(options=RenderOptions(theme=str(theme))).render(
        MarkdownParser(profile="plain").parse("```text\nSource text.\n```")
    )
    data = BytesIO()
    doc.save(data)
    assert 'w:fill="123456"' in Document(data).tables[0]._tbl.xml


def test_legacy_red_theme_controls_waterfall_negative_color(tmp_path, monkeypatch):
    from matplotlib.colors import to_hex
    from matplotlib.figure import Figure

    theme = tmp_path / "negative.yaml"
    theme.write_text('RED: "AABBCC"\n', encoding="utf-8")
    colors = []
    saved = Figure.savefig

    def inspect(self, *args, **kwargs):
        colors.append(to_hex(self.axes[0].patches[1].get_facecolor()))
        return saved(self, *args, **kwargs)

    monkeypatch.setattr(Figure, "savefig", inspect)
    model = MarkdownParser().parse(
        "```chart\nchart_type: waterfall\nlabels: [A, B]\n"
        "series: [{name: Value, values: [100, -30]}]\n```"
    )
    doc = IBDocumentRenderer(options=RenderOptions(theme=str(theme), charts=True, strict=True)).render(model)
    data = BytesIO()
    doc.save(data)
    assert len(Document(data).inline_shapes) == 1
    assert colors == ["#aabbcc"]


def test_themes_remain_isolated_across_threads(tmp_path):
    paths = []
    for color in ("123456", "654321"):
        path = tmp_path / (color + ".yaml")
        path.write_text(f'NAVY: "{color}"\n', encoding="utf-8")
        paths.append(path)
    with ThreadPoolExecutor(max_workers=2) as pool:
        docs = list(pool.map(themed_document, paths * 2))
    for doc, color in zip(docs, ("123456", "654321") * 2):
        heading = next(p for p in doc.paragraphs if p.text == "Heading")
        assert str(heading.runs[0].font.color.rgb) == color
        assert f'w:fill="{color}"' in doc.tables[0]._tbl.xml
    assert STYLE.NAVY == IBStyle().NAVY


@pytest.mark.parametrize("invalid", [False, True])
def test_extended_theme_cli_overrides_yaml_and_preserves_failed_destination(tmp_path, monkeypatch, invalid):
    theme = tmp_path / "cli.yaml"
    theme.write_text('NAVY: "123456"\n' + ('H1_SIZE: true\n' if invalid else 'H1_SIZE: 21\n'), encoding="utf-8")
    source, output = tmp_path / "input.md", tmp_path / "output.docx"
    source.write_text("---\ntheme: mono\n---\n# Heading\n\nBody.", encoding="utf-8")
    output.write_bytes(b"keep existing")
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output), "--theme", str(theme),
                                     "--no-cover", "--no-toc", "--no-disclaimer", "--strict"])
    with pytest.raises(SystemExit) as exc:
        md_to_word.main()
    if invalid:
        assert exc.value.code != 0
        assert output.read_bytes() == b"keep existing"
    else:
        assert exc.value.code == 0
        heading = next(p for p in Document(output).paragraphs if p.text == "Heading")
        assert heading.runs[0].font.size == Pt(21)
        assert str(heading.runs[0].font.color.rgb) == "123456"
