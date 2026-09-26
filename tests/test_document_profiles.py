"""User-path regressions for Markdown-to-Word profiles and table integrity."""

import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from converters import get_default_registry
from document_model import DocumentModel, Element, ElementType
from document_profiles import PROFILES, RenderOptions
from docx_audit import audit_file, inspect_document
from ib_renderer import STYLE, DocumentStyler, IBDocumentRenderer, TableRenderer
from md_parser import MarkdownParser, TableParser, parse_markdown_file
from md_to_word import IBReportConverter

OFFICE_MD = """---
profile: office-letter
title: 자료 제출 요청
document_no: 기획-2026-015
date: "2026-09-14"
sender:
  organization: 예시회사
  department: 기획팀
  signatory: 대표이사
recipients:
  - 협력사 담당부서
cc:
  - 재무팀
attachments:
  - 제출 양식 1부
  - 작성 안내 1부
---
# 자료 제출 요청

1. 자료 제출
    1. 담당자 확인
2. 일정 확인
"""


def render(content, profile=None, **kwargs):
    model = MarkdownParser(profile=profile).parse(content)
    return IBDocumentRenderer(options=RenderOptions(**kwargs)).render(model)


def all_text(doc):
    return "\n".join(doc.element.xpath(".//w:t/text()"))


def test_blank_middle_cell_does_not_move_the_deadline():
    table = TableParser.parse(
        [
            "| Task | Owner | Deadline |",
            "|---|---|---|",
            "| Draft | | Friday |",
        ]
    )
    assert [cell.content for cell in table.rows[1].cells] == ["Draft", "", "Friday"]


def test_escaped_pipe_does_not_drop_the_next_column():
    table = TableParser.parse(["| Clause | Owner |", "|---|---|", r"| A\|B | Kim |"])
    assert [cell.content for cell in table.rows[1].cells] == ["A|B", "Kim"]


def test_financial_number_is_formatted_in_saved_document():
    table = TableParser.parse(["| Metric | 2026 |", "|---|---|", "| Revenue | **1234567** |"])
    doc = Document()
    DocumentStyler(doc).create_styles()
    TableRenderer(doc).render(table)
    payload = BytesIO()
    doc.save(payload)
    payload.seek(0)
    saved = Document(payload)
    assert saved.tables[0].cell(1, 1).text == "1,234,567"
    assert saved.tables[0].cell(1, 1).paragraphs[0].runs[0].bold


@pytest.mark.parametrize("profile", list(PROFILES))
def test_each_profile_produces_a_valid_document(profile):
    content = OFFICE_MD if profile == "office-letter" else "# Example\n\nBody."
    doc = render(content, profile=profile)
    assert not inspect_document(doc).issues
    assert "Body." in all_text(doc) or "자료 제출" in all_text(doc)


def test_office_fields_and_native_korean_numbering():
    doc = render(OFFICE_MD)
    text = all_text(doc)
    for expected in ("협력사 담당부서", "재무팀", "기획-2026-015", "제출 양식 1부", "대표이사"):
        assert expected in text
    assert text.count("자료 제출 요청") == 1
    assert "Korea Development Bank" not in text
    assert "CONFIDENTIAL" not in "".join(p.text for p in doc.sections[0].header.paragraphs)
    assert doc.sections[0].page_width.mm == pytest.approx(210, abs=0.01)
    numbers = doc.element.xpath(".//w:pPr/w:numPr")
    assert len(numbers) == 5
    assert [n.ilvl.val for n in numbers] == [0, 1, 0, 0, 0]
    assert 'w:val="ganada"' in doc.part.numbering_part.element.xml
    assert "1. 자료 제출" not in text


def test_plain_preserves_references_as_content_and_does_not_infer_financial_table():
    model = MarkdownParser(profile="plain").parse(
        "## References\n\n1. Source material\n\n| Item | 2026 |\n|---|---|\n| Code | 001234 |"
    )
    assert model.elements[1].element_type == ElementType.NUMBERED_LIST
    assert model.elements[2].content.table_type.name == "GENERIC"
    doc = IBDocumentRenderer().render(model)
    assert doc.tables[0].cell(1, 1).text == "001234"


def test_structured_metadata_is_not_stringified():
    model = MarkdownParser().parse(OFFICE_MD)
    assert model.metadata.extra["sender"]["organization"] == "예시회사"
    assert model.metadata.extra["recipients"] == ["협력사 담당부서"]


def test_cli_profile_overrides_frontmatter_without_ib_defaults(tmp_path):
    source = tmp_path / "input.md"
    source.write_text(
        "---\nprofile: ib-report\ntitle: Neutral\n---\n# Neutral\n\nBody.", encoding="utf-8"
    )
    model = parse_markdown_file(str(source), profile="plain")
    assert model.metadata.profile == "plain"
    assert model.metadata.company == ""
    converter = IBReportConverter(str(source), render_options=RenderOptions(profile="plain"))
    assert "당행" not in all_text(converter._render(model))


def test_cli_library_registry_use_identical_composition(tmp_path):
    source = tmp_path / "source.md"
    source.write_text(OFFICE_MD, encoding="utf-8")
    model = parse_markdown_file(str(source))
    options = RenderOptions(strict=True)
    cli = IBReportConverter(str(source), render_options=options)._render(model)
    api = IBDocumentRenderer(options=options).render(model)
    registry = get_default_registry().convert(model, output_format="docx", render_options=options)
    assert cli.element.xml == api.element.xml == registry.element.xml
    assert cli.sections[0].header._element.xml == api.sections[0].header._element.xml


def test_disclaimer_override_removes_both_cover_and_final_disclaimer():
    content = "---\nlayout:\n  toc: false\n---\n# Report\n\nBody."
    doc = render(content, include_disclaimer=False)
    assert "당행" not in all_text(doc)
    assert "DISCLAIMER" not in all_text(doc)
    assert "TABLE OF CONTENTS" not in all_text(doc)


def test_frontmatter_and_cli_flags_have_precedence():
    content = "---\nprofile: plain\nlayout:\n  toc: true\n---\n# Example\n\nBody."
    assert "목차" in all_text(render(content))
    assert "목차" not in all_text(render(content, include_toc=False))


@pytest.mark.parametrize(
    "metadata",
    [
        "profile: missing",
        "profile: plain\nlayout: nope",
        "profile: plain\nlayout:\n  toc: 'false'",
        "profile: plain\nlayout:\n  typo: true",
    ],
)
def test_invalid_configuration_is_rejected(metadata):
    with pytest.raises(ValueError):
        render("---\n" + metadata + "\n---\n# Example")


def test_office_letter_requires_sender_and_recipient():
    with pytest.raises(ValueError, match="recipients"):
        render("# Letter", profile="office-letter")


def test_links_and_numeric_markdown_footnotes_are_native():
    doc = render(
        "See [source](https://example.com).[^1]\n\n[^1]: Supporting note.", profile="plain"
    )
    assert len(doc.element.xpath(".//w:hyperlink")) == 1
    assert len(doc.element.xpath(".//w:footnoteReference")) == 1
    assert "[^1]" not in all_text(doc)
    assert not inspect_document(doc).issues


def test_explicit_financial_roles_units_sources_and_landscape():
    content = """---
profile: ib-memo
tables:
  - type: financial
    columns: [text, money, percent, code]
    caption: 실적 요약
    unit: 백만원
    as_of: "2026-06-30"
    source: "[회사 자료](https://example.com)"
    landscape: true
---
# 실적

| 항목 | 금액 | 비율 | 코드 |
|---|---|---|---|
| 매출 | **1234567** | 12.5 | 001234 |

후속 본문.
"""
    doc = render(content)
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["매출", "1,234,567", "12.5%", "001234"]
    assert "단위: 백만원" in all_text(doc)
    assert "기준일: 2026-06-30" in all_text(doc)
    assert len(doc.sections) == 3
    assert doc.sections[1].page_width > doc.sections[1].page_height
    assert doc.sections[2].page_width < doc.sections[2].page_height
    assert len(doc.tables[0]._tbl.xpath(".//w:tblHeader")) == 1
    assert len(doc.tables[0]._tbl.xpath(".//w:cantSplit")) == 2


def test_explicit_sensitivity_base_case_uses_coordinates():
    model = MarkdownParser().parse(
        "---\ntables:\n  - type: sensitivity\n    base_case: {row: 1, column: 2}\n---\n| Rate | Value |\n|---|---|\n| 1 | 10 |\n| 2 | 20 |"
    )
    table = model.elements[0].content
    assert table.rows[1].cells[1].is_base_case
    assert sum(c.is_base_case for row in table.rows for c in row.cells) == 1


def test_strict_input_loss_does_not_save_successfully(tmp_path):
    source, output = tmp_path / "broken.md", tmp_path / "broken.docx"
    source.write_text("| A | B |\n|---|---|\n| x | y | LOST |", encoding="utf-8")
    with pytest.raises(RuntimeError, match="validation"):
        IBReportConverter(str(source), str(output), RenderOptions(strict=True)).convert()
    assert not output.exists()


def test_strict_renderer_errors_are_not_silently_successful():
    model = DocumentModel(elements=[Element(ElementType.TABLE, None)])
    with pytest.raises(ValueError, match="validation"):
        IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)


def test_style_scope_is_restored_after_exception(tmp_path):
    theme = tmp_path / "theme.yaml"
    theme.write_text("body_font: Arial\nprimary_color: '123456'", encoding="utf-8")
    model = DocumentModel(elements=[Element(ElementType.TABLE, None)])
    with pytest.raises(ValueError):
        IBDocumentRenderer(options=RenderOptions(theme=str(theme), strict=True)).render(model)
    assert STYLE.BODY_FONT == "Calibri"
    assert STYLE.NAVY_HEX == "003366"


def test_profile_styles_are_isolated_across_concurrent_requests():
    def font(profile):
        doc = render(
            "Body.",
            profile=profile,
            include_cover=False,
            include_toc=False,
            include_disclaimer=False,
        )
        return doc.styles["IB Body"].font.name

    profiles = ["plain", "ib-report"] * 4
    with ThreadPoolExecutor(max_workers=4) as executor:
        actual = list(executor.map(font, profiles))
    assert actual == ["Malgun Gothic", "Calibri"] * 4


def test_legacy_style_proxy_cannot_be_mutated():
    with pytest.raises(AttributeError):
        STYLE.BODY_FONT = "Injected global font"


def test_relative_theme_is_resolved_from_markdown_directory(tmp_path):
    (tmp_path / "style.yaml").write_text("body_font: Arial", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("---\nprofile: plain\ntheme: style.yaml\n---\nBody.", encoding="utf-8")
    doc = IBDocumentRenderer().render(parse_markdown_file(str(source)))
    assert doc.styles["IB Body"].font.name == "Arial"


def test_removed_conversion_directions_are_not_registered():
    registry = get_default_registry()
    assert registry.find_converter("retired.docx") is None
    assert registry.find_converter(DocumentModel(), output_format="md") is None


def test_office_letter_cli_creates_inspectable_document(tmp_path):
    source, output = tmp_path / "공문.md", tmp_path / "공문.docx"
    source.write_text(OFFICE_MD, encoding="utf-8")
    result = subprocess.run(
        [sys.executable, "md_to_word.py", str(source), str(output), "--strict"],
        cwd=str(Path(__file__).resolve().parent.parent),
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0, result.stderr.decode("utf-8", errors="replace")
    assert not audit_file(str(output)).issues


def test_plain_reference_list_is_not_duplicated_as_endnotes():
    doc = render("## References\n\n1. Source material", profile="plain")
    assert all_text(doc).count("Source material") == 1
    assert "ENDNOTES" not in all_text(doc)


def test_all_empty_table_row_is_preserved():
    table = TableParser.parse(["| A | B |", "|---|---|", "| | |"])
    assert len(table.rows) == 2
    assert [c.content for c in table.rows[1].cells] == ["", ""]


def test_sensitivity_does_not_guess_a_base_case():
    table = TableParser.parse(
        ["| Sensitivity | Value |", "|---|---|", "| Base | 10 |", "| Up | 20 |"]
    )
    assert not any(c.is_base_case for row in table.rows for c in row.cells)


def test_explicit_risk_table_styles_valid_levels():
    doc = render(
        "---\nprofile: plain\ntables:\n  - type: risk\n---\n| 항목 | 등급 |\n|---|---|\n| 운영 | 높음 |"
    )
    run = doc.tables[0].cell(1, 1).paragraphs[0].runs[0]
    assert run.bold
    assert str(run.font.color.rgb) == str(STYLE.RED)


def test_bold_financial_negative_keeps_colour():
    doc = render(
        "| Metric | 2026 |\n|---|---|\n| Income | **-1234** |",
        include_cover=False,
        include_toc=False,
        include_disclaimer=False,
    )
    cell = doc.tables[0].cell(1, 1)
    assert cell.text == "-1,234"
    assert str(cell.paragraphs[0].runs[0].font.color.rgb) == str(STYLE.RED)


def test_explicit_code_role_does_not_receive_negative_financial_styling():
    doc = render(
        "---\nprofile: ib-memo\ntables:\n  - type: financial\n    columns: [text, code]\n---\n| Name | Code |\n|---|---|\n| Item | -001234 |"
    )
    cell = doc.tables[0].cell(1, 1)
    assert cell.text == "-001234"
    assert cell.paragraphs[0].runs[0].font.color.rgb is None


def test_strict_empty_diagram_is_not_silently_dropped():
    from document_model import Diagram

    model = DocumentModel(elements=[Element(ElementType.DIAGRAM, Diagram())])
    with pytest.raises(ValueError, match="validation"):
        IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)


@pytest.mark.parametrize("text", ["Missing.[^9]", "Note.[^1]\n\n[^1]: First\n[^1]: Second"])
def test_strict_rejects_invalid_footnote_contract(text):
    with pytest.raises(ValueError, match="validation"):
        render(text, profile="plain", strict=True)


def test_footnote_zero_cannot_collide_with_word_separator():
    with pytest.raises(ValueError, match="positive"):
        render("Note.[^0]\n\n[^0]: Invalid", profile="plain")


def test_output_signature_tracks_profile():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = render("Body", profile="meeting-minutes")
    signature = doc.part.package.part_related_by(RT.CUSTOM_PROPERTIES).blob.decode()
    assert "meeting-minutes" in signature
    assert "2.0.0" in signature


def test_ib_body_mentioning_sources_does_not_hide_following_table():
    model = MarkdownParser().parse(
        "---\nprofile: ib-memo\n---\n실적과 출처를 확인합니다.\n\n| 지표 | 금액 |\n|---|---|\n| 매출 | 1234 |"
    )
    assert any(e.element_type == ElementType.TABLE for e in model.elements)


@pytest.mark.parametrize(
    "source",
    sorted((Path(__file__).resolve().parent.parent / "samples" / "profiles").glob("*.md")),
    ids=lambda p: p.stem,
)
def test_shipped_examples_render_strictly(source):
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(
        parse_markdown_file(str(source))
    )
    assert not inspect_document(doc).issues
