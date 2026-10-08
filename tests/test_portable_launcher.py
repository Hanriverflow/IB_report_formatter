"""Portable entry point checks against the actual parsing/rendering/save path."""

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document

from scripts import build_portable
from scripts.build_portable import package_inputs
from tools import portable_launcher


def test_conversion_preserves_yaml_and_korean_spaced_paths(tmp_path):
    source_dir = tmp_path / "한글 입력 폴더"
    source_dir.mkdir()
    (source_dir / "house.yaml").write_text(
        (Path(__file__).resolve().parents[1] / "samples/profiles/term-sheet-house.yaml").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    source = source_dir / "첫 텀시트.md"
    source.write_text(
        '---\nprofile: term-sheet\nhouse: house.yaml\ntitle: "{{borrower}} 조건표"\n'
        'terms:\n  borrower: 가상기업\n---\n# 주요 조건\n\n'
        '| 항목 | 내용 |\n|---|---|\n| 차주 | {{borrower}} |\n', encoding="utf-8"
    )
    result = portable_launcher.convert_document(str(source), str(tmp_path / "한글 출력 폴더"))
    assert result.is_file()
    document = Document(result)
    xml = document._element.xml
    assert "가상기업" in xml
    assert "{{borrower}}" not in xml
    assert "주요 조건" in xml


def test_conversion_does_not_overwrite_existing_result(tmp_path):
    source = tmp_path / "sample.md"
    source.write_text("---\nprofile: plain\n---\n새 결과\n", encoding="utf-8")
    existing = tmp_path / "sample.docx"
    existing.write_bytes(b"existing document")
    saved = portable_launcher.convert_document(str(source), str(tmp_path))
    assert saved != existing
    assert existing.read_bytes() == b"existing document"
    assert Document(saved).paragraphs[0].text == "새 결과"


def test_strict_failure_does_not_save_partial_docx(tmp_path):
    source = tmp_path / "누락 이미지.md"
    source.write_text("---\nprofile: plain\n---\n![누락](absent.png)\n", encoding="utf-8")
    with pytest.raises(RuntimeError, match="Rendering failed"):
        portable_launcher.convert_document(str(source), str(tmp_path / "출력"))
    assert not list(tmp_path.rglob("*.docx"))


@pytest.mark.parametrize("source, output, message", [
    ("", ".", "선택"), ("a.md", "", "폴더"), ("a.docx", ".", "Markdown"),
])
def test_actionable_input_errors(source, output, message):
    with pytest.raises(ValueError, match=message):
        portable_launcher.convert_document(source, output)


def test_engine_saved_path_is_returned_without_reconstruction(tmp_path, monkeypatch):
    source = tmp_path / "source.md"
    source.write_text("---\nprofile: plain\n---\n본문", encoding="utf-8")
    actual = tmp_path / "lock fallback.docx"
    def save_fallback(doc, requested):
        doc.save(actual)
        return actual
    monkeypatch.setattr("md_to_word.safe_save", save_fallback)
    assert portable_launcher.convert_document(str(source), str(tmp_path)) == actual


def test_frozen_root_follows_executable_not_cwd(tmp_path, monkeypatch):
    monkeypatch.setattr(portable_launcher.sys, "frozen", True, raising=False)
    monkeypatch.setattr(portable_launcher.sys, "executable", str(tmp_path / "문서변환기.exe"))
    assert portable_launcher.application_root() == tmp_path


def test_package_allowlist_only_includes_public_guide_license_and_samples():
    inputs = package_inputs()
    assert len(inputs) == 13
    assert {destination.parts[0] for _, destination in inputs} == {"시작하기.html", "라이선스", "예제"}
    assert all(".." not in destination.parts and not destination.is_absolute() for _, destination in inputs)
    assert not any(source.suffix == ".docx" for source, _ in inputs)
    assert (Path("예제/빠른시작/minimal-house.yaml")) in {dest for _, dest in inputs}


def test_portable_archive_filename_and_matching_root(tmp_path, monkeypatch):
    source_root = tmp_path / "source"
    source_root.mkdir()
    (source_root / "pyproject.toml").write_text(
        '[project]\nversion = "9.8.7"\n', encoding="utf-8"
    )
    license_file = source_root / "LICENSE"
    license_file.write_text("test license", encoding="utf-8")
    app_dir = tmp_path / "frozen-app"
    app_dir.mkdir()
    executable = f"{build_portable.APP_NAME}.exe"
    (app_dir / executable).write_bytes(b"test frozen executable")
    monkeypatch.setattr(build_portable, "ROOT", source_root)
    monkeypatch.setattr(build_portable.sys, "platform", "win32")
    monkeypatch.setattr(
        build_portable, "package_inputs",
        lambda: [(license_file, Path("라이선스/LICENSE.txt"))],
    )
    monkeypatch.setattr(build_portable, "collect_notices", lambda destination: {})
    monkeypatch.setattr(build_portable.sys, "path", list(build_portable.sys.path))

    archive = build_portable.build(tmp_path / "release", app_dir)

    assert archive.name == "ib-report-formatter-windows-x64-9.8.7.zip"
    prefix = "ib-report-formatter-windows-x64-9.8.7/"
    with zipfile.ZipFile(archive) as bundle:
        assert all(name.startswith(prefix) for name in bundle.namelist())
        assert bundle.read(prefix + executable) == b"test frozen executable"
        assert bundle.read(prefix + "라이선스/LICENSE.txt") == b"test license"
        assert prefix + "입력/" in bundle.namelist()
        assert prefix + "출력/" in bundle.namelist()
        assert json.loads(bundle.read(prefix + "manifest.json"))["version"] == "9.8.7"
