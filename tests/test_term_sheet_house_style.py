"""House style options (`style:`) for the term-sheet profile; all data fictional."""

from dataclasses import fields
from io import BytesIO
from pathlib import Path
from typing import List

import pytest
import yaml
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image as PILImage

from document_model import Table
from document_profiles import RenderOptions, get_profile, load_style
from ib_renderer import IBDocumentRenderer, TableRenderer
from md_parser import MarkdownParser, parse_markdown_file
from term_sheet import TermSheetStyle

NAVY_HEX = load_style(get_profile("term-sheet"), None).NAVY_HEX
GREY_HEX = load_style(get_profile("term-sheet"), None).TS_CONFIDENTIAL_HEX

TABLE = "| 구 분 | 내 용 |\n|---|---|\n| 차 주 | 가나다머티리얼즈㈜ |\n"
BODY = "## 1. 개요\n\n" + TABLE


def _logo(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    PILImage.new("RGB", (240, 60), "navy").save(path)
    return path


def _house(folder: Path, style: dict) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "house.yaml"
    data = {"prepared_by": "라마바은행 자본시장부", "disclaimer": "가상 조건 검토용입니다.\n두 번째 줄입니다.", "style": style}
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    return path


def _markdown(house: Path, body: str = BODY, **fields) -> str:
    data = {"profile": "term-sheet", "title": "가나다머티리얼즈㈜ 구조화금융", "house": str(house), **fields}
    return "---\n" + yaml.safe_dump(data, allow_unicode=True) + "---\n" + body


def _render(markdown: str, strict: bool = True):
    doc = IBDocumentRenderer(options=RenderOptions(strict=strict)).render(MarkdownParser().parse(markdown))
    payload = BytesIO()
    doc.save(payload)
    return Document(payload)


def _attr(element, path: str, name: str = "w:val") -> List[str]:
    return [node.get(qn(name)) for node in element.xpath(path)]


def _heading_index(doc) -> int:
    return next(index for index, paragraph in enumerate(doc.paragraphs) if paragraph.text == "1. 개요")


def test_default_style_keeps_the_standard_layout(tmp_path: Path) -> None:
    doc = _render(_markdown(_house(tmp_path, {})))
    before = doc.paragraphs[:_heading_index(doc)]
    assert not any(paragraph._p.xpath('.//w:br[@w:type="page"]') for paragraph in before)
    assert not doc.element.body.xpath(".//w:drawing")
    header = doc.sections[0].header.paragraphs[0]
    assert not header._p.xpath("./w:pPr/w:pBdr")
    footer = doc.sections[0].footer.paragraphs[0]
    assert "NUMPAGES" in "".join(footer._p.xpath(".//w:instrText/text()"))
    table = doc.tables[0]._tbl
    assert _attr(table, "./w:tblPr/w:tblBorders/w:left") == ["single"]


def test_cover_page_logo_and_boxed_disclaimer(tmp_path: Path) -> None:
    folder = tmp_path / "house"
    _logo(folder / "assets" / "logo.png")
    doc = _render(_markdown(_house(folder, {"cover": "page", "logo": "assets/logo.png", "disclaimer": "box"})))
    opening = doc.paragraphs[:_heading_index(doc)]
    assert opening[0].paragraph_format.space_before == pytest.approx(Mm(60), abs=635)
    assert any(paragraph._p.xpath(".//w:drawing") for paragraph in opening)
    boxed = [paragraph for paragraph in opening if paragraph.text in ("가상 조건 검토용입니다.", "두 번째 줄입니다.")]
    assert len(boxed) == 2
    for paragraph in boxed:
        assert [node.tag.split("}")[1] for node in paragraph._p.xpath("./w:pPr/w:pBdr/*")] == [
            "top", "left", "bottom", "right",
        ]
    assert opening[-1]._p.xpath('.//w:br[@w:type="page"]')


def test_a_landscape_table_right_after_the_cover_turns_the_page_once(tmp_path: Path) -> None:
    house = _house(tmp_path, {"cover": "page"})
    doc = _render(_markdown(house, body=TABLE + "\n뒤 문단.\n", tables=[{"landscape": True}]))
    body = doc.element.body
    # The landscape section break already starts a new page; the cover's page
    # break before it is redundant (Word ignores it, other viewers may not).
    assert not body.xpath('.//w:br[@w:type="page"]')
    assert [section.page_width > section.page_height for section in doc.sections] == [False, True, False]
    cover_end = body.xpath("./w:p[w:pPr/w:sectPr]")[0]
    assert "".join(cover_end.getprevious().xpath(".//w:t/text()")) == "두 번째 줄입니다."


def test_a_heading_between_the_cover_and_a_landscape_table_keeps_the_page_break(tmp_path: Path) -> None:
    doc = _render(_markdown(_house(tmp_path, {"cover": "page"}), tables=[{"landscape": True}]))
    opening = doc.paragraphs[:_heading_index(doc)]
    assert opening[-1]._p.xpath('.//w:br[@w:type="page"]')
    assert len(doc.sections) == 3


def test_a_paragraph_with_two_page_breaks_before_a_landscape_table_is_kept() -> None:
    # Two breaks in one paragraph ask for a blank page; only a lone break is redundant.
    doc = Document()
    doc.add_page_break().add_run().add_break(WD_BREAK.PAGE)
    table = MarkdownParser().parse(TABLE).elements[0].content
    assert isinstance(table, Table)
    table.landscape = True
    TableRenderer(doc).render(table)
    assert len(doc.sections) == 3
    assert len(doc.element.body.xpath('./w:p/w:r/w:br[@w:type="page"]')) == 2


def test_header_and_footer_style(tmp_path: Path) -> None:
    style = {
        "label_color": "#C00000", "header_rule": True, "footer_rule": True,
        "page_number": "- {page} -", "page_number_align": "center",
    }
    doc = _render(_markdown(_house(tmp_path, style)))
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    assert set(_attr(header._p, ".//w:r/w:rPr/w:color")) == {"C00000"}
    assert _attr(header._p, "./w:pPr/w:pBdr/w:bottom") == ["single"]
    footer = section.footer.paragraphs[0]
    assert _attr(footer._p, "./w:pPr/w:pBdr/w:top") == ["single"]
    assert _attr(footer._p, "./w:pPr/w:tabs/w:tab") == ["center"]
    width = (section.page_width - section.left_margin - section.right_margin) // 2
    assert int(_attr(footer._p, "./w:pPr/w:tabs/w:tab", "w:pos")[0]) == pytest.approx(width / 635, abs=1)
    assert [text.strip() for text in footer._p.xpath(".//w:instrText/text()")] == ["PAGE"]
    assert "".join(footer._p.xpath(".//w:t/text()")) == "-  -"  # "- ", PAGE, " -"


def test_dark_table_header_and_open_sides(tmp_path: Path) -> None:
    doc = _render(_markdown(_house(tmp_path, {"table_header": "dark", "table_sides": "open"})))
    table = doc.tables[0]._tbl
    header_cells = table.xpath("./w:tr[1]/w:tc")
    assert {fill for tc in header_cells for fill in _attr(tc, "./w:tcPr/w:shd", "w:fill")} == {NAVY_HEX}
    assert set(_attr(table, "./w:tr[1]//w:r/w:rPr/w:color")) == {"FFFFFF"}
    assert _attr(table, "./w:tblPr/w:tblBorders/w:left") == ["nil"]
    assert _attr(table, "./w:tblPr/w:tblBorders/w:right") == ["nil"]
    assert _attr(table, "./w:tblPr/w:tblBorders/w:insideV") == ["single"]


def test_frontmatter_style_overrides_the_house_key_by_key(tmp_path: Path) -> None:
    house = _house(tmp_path, {"cover": "page", "table_header": "dark"})
    doc = _render(_markdown(house, style={"cover": "inline"}))
    before = doc.paragraphs[:_heading_index(doc)]
    assert not any(paragraph._p.xpath('.//w:br[@w:type="page"]') for paragraph in before)
    assert set(_attr(doc.tables[0]._tbl, "./w:tr[1]/w:tc/w:tcPr/w:shd", "w:fill")) == {NAVY_HEX}


@pytest.mark.parametrize("style", [
    {"cover": "full"}, {"unknown": 1}, {"header_rule": "yes"}, {"label_color": "red"},
    {"page_number": "{page"}, {"page_number": "Page"}, {"page_number": "{page} of {total}"},
    {"page_number": "{page}\n{pages}"}, {"page_number": "{page}\r/ {pages}"},
    {"page_number": "{page}\u2028"}, {"logo_width_mm": 0}, {"logo": "  "}, {"logo": 7},
])
def test_invalid_style_settings_are_rejected(tmp_path: Path, style: dict) -> None:
    with pytest.raises(ValueError, match="style"):
        _render(_markdown(_house(tmp_path, style)))


@pytest.mark.parametrize("strict", [True, False])
@pytest.mark.parametrize("key", [field.name for field in fields(TermSheetStyle)])
def test_a_null_style_value_is_a_setting_error(tmp_path: Path, key: str, strict: bool) -> None:
    # Omitting a key keeps its default; a YAML null used to pass validation and
    # then break the footer (page_number) or drop the logo (logo_width_mm).
    house = _house(tmp_path / "house", {})
    for source in (_markdown(_house(tmp_path, {key: None})), _markdown(house, style={key: None})):
        with pytest.raises(ValueError, match=f"style.{key} must not be null"):
            _render(source, strict=strict)


def test_an_empty_frontmatter_logo_removes_the_house_logo(tmp_path: Path) -> None:
    folder = tmp_path / "house"
    _logo(folder / "logo.png")
    house = _house(folder, {"logo": "logo.png"})
    assert _render(_markdown(house)).element.body.xpath(".//w:drawing")
    assert not _render(_markdown(house, style={"logo": ""})).element.body.xpath(".//w:drawing")


def test_an_empty_label_color_keeps_the_default_grey(tmp_path: Path) -> None:
    doc = _render(_markdown(_house(tmp_path, {"label_color": "#C00000"}), style={"label_color": ""}))
    header = doc.sections[0].header.paragraphs[0]
    assert header.text
    assert set(_attr(header._p, ".//w:r/w:rPr/w:color")) == {GREY_HEX}


def test_a_missing_logo_is_a_render_diagnostic(tmp_path: Path) -> None:
    source = _markdown(_house(tmp_path, {"logo": "missing.png"}))
    with pytest.raises(ValueError, match="House logo could not be rendered"):
        _render(source)
    assert not _render(source, strict=False).element.body.xpath(".//w:drawing")


def test_a_frontmatter_logo_is_relative_to_the_markdown_file(tmp_path: Path) -> None:
    _logo(tmp_path / "doc" / "logo.png")
    house = _house(tmp_path / "house", {})
    source = tmp_path / "doc" / "sheet.md"
    source.write_text(_markdown(house, style={"logo": "logo.png"}), encoding="utf-8")
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(parse_markdown_file(str(source)))
    assert doc.element.body.xpath(".//w:drawing")
    with pytest.raises(ValueError, match="absolute"):
        _render(_markdown(house, style={"logo": "logo.png"}))
