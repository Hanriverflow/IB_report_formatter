"""Table span parsing contracts through synthetic Markdown and real rendering."""

from io import BytesIO

import pytest
from docx import Document

from document_model import DocumentModel, Table
from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser


def parse_table(body: str, spec: str = "{spans: true}", profile: str = "plain") -> tuple[DocumentModel, Table]:
    """Parse one synthetic table with the requested ordered specification."""
    metadata = (
        f"profile: {profile}\ntitle: 가나다머티리얼즈㈜ 조건 검토\n"
        "prepared_by: 라마바은행 자본시장부\ndisclaimer: 가상 검토용 문구입니다.\n"
    )
    if spec:
        metadata += f"tables:\n  - {spec}\n"
    model = MarkdownParser().parse(f"---\n{metadata}---\n{body}")
    tables = [element.content for element in model.elements if isinstance(element.content, Table)]
    assert len(tables) == 1
    return model, tables[0]


def merges(table: Table) -> list[list[str | None]]:
    """Expose model span directives without depending on renderer internals."""
    return [[cell.merge for cell in row.cells] for row in table.rows]


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ("| 구분 | << | 내용 |\n|---|---|---|\n| 항목 | << | 값 |", [[None, "left", None], [None, "left", None]]),
        ("| 구분 | 내용 |\n|---|---|\n| 항목 | 값 |\n| ^^ | 다음 |", [[None, None], [None, None], ["up", None]]),
        ("| 구분 | 항목 | 내용 |\n|---|---|---|\n| 기준 | << | 값 |\n| ^^ | << | 다음 |", [[None, None, None], [None, "left", None], ["up", "left", None]]),
        ("| 구분 | 항목 | 내용 |\n|---|---|---|\n| 기준 | << | 값 |\n| ^^ | ^^ | 다음 |", [[None, None, None], [None, "left", None], ["up", "up", None]]),
    ],
    ids=["horizontal", "vertical", "rectangle_left", "rectangle_up"],
)
def test_valid_span_groups_resolve_without_warnings(body: str, expected: list[list[str | None]]) -> None:
    model, table = parse_table(body)
    assert merges(table) == expected
    assert not model.warnings


@pytest.mark.parametrize(
    "body",
    [
        "| 구분 | 항목 | 내용 |\n|---|---|---|\n| 기준 | << | 값 |\n| ^^ | 별도 | 다음 |",
        "| 구분 | 항목 | 내용 |\n|---|---|---|\n| 기준 | << | 값 |\n| ^^ | | 다음 |",
        "| 구분 | 내용 |\n|---|---|\n| ^^ | 값 |",
        "| ^^ | 내용 |\n|---|---|\n| 항목 | 값 |",
        "| 구분 | 내용 |\n|---|---|\n| << | 값 |",
        "| << | << | 내용 |\n|---|---|---|\n| 항목 | 값 | 다음 |",
    ],
    ids=["l_shape", "blank_hole", "header_body_boundary", "top_boundary", "left_boundary", "anchorless_chain"],
)
def test_invalid_span_group_rolls_back_to_literal_markers(body: str) -> None:
    model, table = parse_table(body)
    assert all(cell.merge is None for row in table.rows for cell in row.cells)
    assert len(model.warnings) == 1
    assert "span" in model.warnings[0].lower()
    for row in table.rows:
        for cell in row.cells:
            if cell.content in {"^^", "<<"}:
                assert "".join(run.text for run in cell.runs) == cell.content


def test_invalid_group_does_not_roll_back_an_adjacent_valid_group() -> None:
    model, table = parse_table(
        "| 가 | 나 | 다 | 라 |\n|---|---|---|---|\n"
        "| 기준 | << | 정상 | << |\n| ^^ | 별도 | ^^ | << |"
    )
    assert merges(table) == [
        [None, None, None, None],
        [None, None, None, "left"],
        [None, None, "up", "left"],
    ]
    assert len(model.warnings) == 1


def test_each_invalid_span_group_has_its_own_warning() -> None:
    model, table = parse_table(
        "| 가 | 나 | 다 |\n|---|---|---|\n| << | 별도 | ^^ |"
    )
    assert merges(table) == [[None, None, None], [None, None, None]]
    assert len(model.warnings) == 2


def test_invalid_outer_rectangle_preserves_an_independent_inner_group() -> None:
    model, table = parse_table(
        "| 가 | 나 | 다 | 라 |\n|---|---|---|---|\n"
        "| 외부 | << | << | 값 |\n| ^^ | 내부 | << | 값 |\n| ^^ | << | << | 값 |"
    )
    assert merges(table) == [
        [None, None, None, None],
        [None, None, None, None],
        [None, None, "left", None],
        [None, None, None, None],
    ]
    assert len(model.warnings) == 1


def test_body_reference_rolls_back_its_connected_header_group_only() -> None:
    model, table = parse_table(
        "| 구분 | << | 내용 |\n|---|---|---|\n"
        "| ^^ | << | 값 |\n| 별도 | << | 다음 |"
    )
    assert merges(table) == [
        [None, None, None], [None, None, None], [None, "left", None]
    ]
    assert table.label_columns == 1
    assert len(model.warnings) == 1


def test_raw_content_controls_classification_instead_of_visible_runs() -> None:
    model, table = parse_table(
        "| 구분 | 내용 |\n|---|---|\n| 항목 | 값 |\n| **^^** | `<<` |"
    )
    assert merges(table) == [[None, None], [None, None], [None, None]]
    assert table.rows[2].cells[0].content == "**^^**"
    assert "".join(run.text for run in table.rows[2].cells[0].runs) == "^^"
    assert table.rows[2].cells[1].content == "`<<`"
    assert not model.warnings


def test_active_span_escapes_fix_raw_content_and_runs_without_merging() -> None:
    model, table = parse_table(
        "| 구분 | 내용 |\n|---|---|\n" + r"| \^^ | \<< |"
    )
    assert [cell.content for cell in table.rows[1].cells] == ["^^", "<<"]
    assert ["".join(run.text for run in cell.runs) for cell in table.rows[1].cells] == ["^^", "<<"]
    assert merges(table) == [[None, None], [None, None]]
    assert not model.warnings


def test_escaped_literal_can_anchor_an_explicit_span() -> None:
    model, table = parse_table("| 구분 | 내용 |\n|---|---|\n" + r"| \^^ | << |")
    assert [cell.content for cell in table.rows[1].cells] == ["^^", "<<"]
    assert merges(table) == [[None, None], [None, "left"]]
    assert not model.warnings


@pytest.mark.parametrize("profile", ["plain", "term-sheet"])
def test_disabled_spans_preserve_marker_content_and_legacy_escape_runs(profile: str) -> None:
    model, table = parse_table(
        "| 구분 | 내용 |\n|---|---|\n| ^^ | << |\n" + r"| \^^ | \<< |",
        "{spans: false}",
        profile,
    )
    assert merges(table) == [[None, None], [None, None], [None, None]]
    assert [cell.content for cell in table.rows[2].cells] == [r"\^^", r"\<<"]
    assert ["".join(run.text for run in cell.runs) for cell in table.rows[2].cells] == ["^^", r"\<<"]
    assert not model.warnings


def test_blank_cells_are_literals_and_can_anchor_explicit_merges() -> None:
    model, table = parse_table(
        "| 구분 | 내용 |\n|---|---|\n| | 값 |\n| ^^ | |"
    )
    assert table.rows[1].cells[0].content == ""
    assert table.rows[2].cells[1].content == ""
    assert merges(table) == [[None, None], [None, None], ["up", None]]
    assert not model.warnings


@pytest.mark.parametrize(("profile", "expected"), [("plain", None), ("term-sheet", "left")])
def test_profile_span_default_applies_without_specification(profile: str, expected: str | None) -> None:
    model, table = parse_table("| 구분 | << | 내용 |\n|---|---|---|\n| 항목 | 값 | 다음 |", "", profile)
    assert table.rows[0].cells[1].merge == expected
    assert not model.warnings


def test_ordered_placeholder_and_unspecified_tables_keep_profile_defaults() -> None:
    model = MarkdownParser().parse(
        "---\nprofile: term-sheet\ntitle: 가나다머티리얼즈㈜ 조건 검토\n"
        "prepared_by: 라마바은행 자본시장부\ndisclaimer: 가상 검토용 문구입니다.\n"
        "tables:\n  - {}\n  - {spans: false}\n---\n"
        + "\n\n".join(["| 구분 | << | 내용 |\n|---|---|---|\n| 항목 | 값 | 다음 |"] * 3)
    )
    tables = [element.content for element in model.elements if isinstance(element.content, Table)]
    assert [table.rows[0].cells[1].merge for table in tables] == ["left", None, "left"]


@pytest.mark.parametrize("value", ["1", "0", "null", '"true"', "[]", "{}"])
def test_spans_spec_requires_a_boolean(value: str) -> None:
    with pytest.raises(ValueError, match="spans"):
        parse_table("| 구분 | 내용 |\n|---|---|\n| 항목 | 값 |", f"{{spans: {value}}}")


@pytest.mark.parametrize("value", ["true", "false", "-1", "2", "1.0", '"1"', "null", "[]"])
def test_label_columns_spec_requires_an_integer_in_bounds(value: str) -> None:
    with pytest.raises(ValueError, match="label_columns"):
        parse_table("| 구분 | 내용 |\n|---|---|\n| 항목 | 값 |", f"{{label_columns: {value}}}")


@pytest.mark.parametrize(
    ("header", "delimiter", "body", "expected"),
    [
        ("| 구분 | 내용 |", "|---|---|", "| 항목 | 값 |", 1),
        ("| 구분 | << | 내용 |", "|---|---|---|", "| 항목 | 세부 | 값 |", 2),
        ("| 구분 | << | << |", "|---|---|---|", "| 항목 | 세부 | 값 |", 2),
        ("| 내용 |", "|---|", "| 값 |", 0),
        ("| << | << | 내용 |", "|---|---|---|", "| 항목 | 세부 | 값 |", 1),
    ],
    ids=["one_label", "two_labels", "clamped", "one_column", "invalid_header_rollback"],
)
def test_label_inference_uses_validated_first_header_span(
    header: str, delimiter: str, body: str, expected: int
) -> None:
    _, table = parse_table("\n".join([header, delimiter, body]), "", "term-sheet")
    assert table.label_columns == expected


@pytest.mark.parametrize("count", [0, 1, 2])
def test_explicit_label_columns_override_inference(count: int) -> None:
    _, table = parse_table(
        "| 구분 | << | 내용 |\n|---|---|---|\n| 항목 | 세부 | 값 |",
        f"{{label_columns: {count}}}",
        "term-sheet",
    )
    assert table.label_columns == count


def test_legacy_profile_without_span_spec_keeps_label_inference_inactive() -> None:
    _, table = parse_table("| 구분 | 내용 |\n|---|---|\n| 항목 | 값 |", "", "plain")
    assert table.label_columns is None


def test_invalid_span_markers_survive_saved_document_and_strict_rejects() -> None:
    model, _ = parse_table("| 구분 | 내용 |\n|---|---|\n| ^^ | << |")
    doc = IBDocumentRenderer().render(model)
    payload = BytesIO()
    doc.save(payload)
    payload.seek(0)
    saved = Document(payload)
    assert saved.tables[0].cell(1, 0).text == "^^"
    assert saved.tables[0].cell(1, 1).text == "<<"
    with pytest.raises(ValueError, match="(?i)(span|strict)"):
        IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
