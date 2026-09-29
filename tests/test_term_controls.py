"""Term values rendered as Word plain-text content controls with a generation snapshot."""

import json
import re
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from lxml import etree

from converters import get_default_registry
from document_profiles import RenderOptions
from docx_audit import inspect_document
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser
from md_to_word import build_parser, run_conversion

VALUES = {"amount": "500억원", "tenor": "3년"}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def xp(element, path: str) -> List:
    """XPath with explicit namespaces; python-docx has no class for `w:sdt`."""
    return etree.XPath(path, namespaces=NS)(element)


def front(values: Dict[str, str], extra: str = "") -> str:
    """Frontmatter with exact YAML string values (JSON is valid YAML)."""
    lines = "".join(
        f"  {key}: {json.dumps(value, ensure_ascii=False)}\n" for key, value in values.items()
    )
    return "---\n" + extra + "terms:\n" + lines + "---\n"


def render(
    markdown: str, profile: str = "plain", renderer: Optional[IBDocumentRenderer] = None,
    **options,
) -> DocxDocument:
    """Parse and render strictly through the production path, then reopen the file."""
    model = MarkdownParser(profile=profile).parse(markdown)
    renderer = renderer or IBDocumentRenderer(options=RenderOptions(
        profile=profile, strict=True, **options,
    ))
    return reopen(renderer.render(model))


def reopen(doc: DocxDocument) -> DocxDocument:
    """Save and reopen, as Word or docx-audit would read the output."""
    payload = BytesIO()
    doc.save(payload)
    payload.seek(0)
    return Document(payload)


def controls(doc: DocxDocument) -> List:
    """Run-level structured document tags in the body, in document order."""
    return xp(doc.element.body, ".//w:sdt")


def local(element) -> str:
    """Tag without namespace."""
    return etree.QName(element).localname


def describe(sdt) -> Tuple[str, str, int, str]:
    """Alias, tag, id and text of one control."""
    properties = sdt.find(qn("w:sdtPr"))
    return (
        properties.find(qn("w:alias")).get(qn("w:val")),
        properties.find(qn("w:tag")).get(qn("w:val")),
        int(properties.find(qn("w:id")).get(qn("w:val"))),
        "".join(xp(sdt, "./w:sdtContent//w:t/text()")),
    )


def xml_text(element) -> str:
    """Text of an element including content-control content."""
    return "".join(xp(element, ".//w:t/text()"))


def custom_properties(doc: DocxDocument) -> Dict[str, Tuple[int, str]]:
    """Custom document properties as name -> (pid, text)."""
    part = doc.part.package.part_related_by(RT.CUSTOM_PROPERTIES)
    root = etree.fromstring(part.blob)
    return {prop.get("name"): (int(prop.get("pid")), "".join(prop.itertext())) for prop in root}


def term_properties(doc: DocxDocument) -> Dict[str, str]:
    """Only the term snapshot properties."""
    return {
        name[len("ibrep.term."):]: value
        for name, (_, value) in custom_properties(doc).items()
        if name.startswith("ibrep.term.")
    }


# ─────────────────────────────────────────────────────────────────────────────
# Tagging
# ─────────────────────────────────────────────────────────────────────────────


def test_value_is_wrapped_in_an_unlocked_plain_text_control() -> None:
    doc = render(front(VALUES) + "Deal {{amount}} closes.")
    [sdt] = controls(doc)
    assert [local(child) for child in sdt.find(qn("w:sdtPr"))] == ["alias", "tag", "id", "text"]
    alias, tag, control_id, text = describe(sdt)
    assert (alias, tag, text) == ("amount", "ibrep:term:amount", "500억원")
    assert control_id > 0
    assert not xp(sdt, ".//w:lock | .//w:dataBinding | .//w:showingPlcHdr")
    assert len(xp(sdt, "./w:sdtContent/w:r")) == 1
    paragraph = sdt.getparent()
    assert local(paragraph) == "p"
    assert xml_text(paragraph) == "Deal 500억원 closes."


def test_every_occurrence_is_tagged_with_a_unique_id() -> None:
    doc = render(
        front(VALUES)
        + "## Tenor {{tenor}}\n\nDeal {{amount}} and {{amount}}.\n\n- Bullet {{amount}}\n\n"
        "| {{tenor}} | Value |\n|---|---|\n| Amount | {{amount}} |\n\n"
        "> [참고] Quote {{amount}}\n\n```text\ncode\n```\n"
    )
    described = [describe(sdt) for sdt in controls(doc)]
    assert [(alias, text) for alias, _, _, text in described] == [
        ("tenor", "3년"), ("amount", "500억원"), ("amount", "500억원"), ("amount", "500억원"),
        ("tenor", "3년"), ("amount", "500억원"), ("amount", "500억원"),
    ]
    ids = [control_id for _, _, control_id, _ in described]
    assert len(set(ids)) == len(ids) and min(ids) > 0
    other_ids = {int(value) for value in xp(doc.element, "//@w:id") if value.lstrip("-").isdigit()}
    assert other_ids and not other_ids & set(ids)


def test_headings_callouts_lists_and_header_cells_use_carried_runs() -> None:
    doc = render(
        front(VALUES)
        + "## Tenor {{tenor}}\n\n> [참고] Amount {{amount}}\n\n- Bullet {{amount}}\n\n"
        "| {{tenor}} | B |\n|---|---|\n| 1 | 2 |"
    )
    heading = next(p for p in doc.paragraphs if p.style.name == "Heading 2")
    assert xml_text(heading._p) == "Tenor 3년"
    [heading_control] = xp(heading._p, "./w:sdt")
    assert xp(heading_control, "./w:sdtContent/w:r/w:rPr/w:b")
    callout = doc.tables[0].cell(0, 0)
    assert [xml_text(p._p) for p in callout.paragraphs] == ["ℹ 참고", "Amount 500억원"]
    assert len(xp(callout._tc, ".//w:sdt")) == 1
    bullet = next(p for p in doc.paragraphs if xml_text(p._p) == "Bullet 500억원")
    assert xp(bullet._p, "./w:sdt")
    header = doc.tables[1].cell(0, 0)
    [header_control] = xp(header._tc, ".//w:sdt")
    assert describe(header_control)[3] == "3년"
    assert xp(header_control, "./w:sdtContent/w:r/w:rPr/w:b")


def test_run_formatting_survives_inside_the_control() -> None:
    doc = render(front(VALUES) + '**{{amount}}** and <span style="color:#C00000">{{tenor}}</span>')
    bold, colored = controls(doc)
    assert xp(bold, "./w:sdtContent/w:r/w:rPr/w:b")
    assert xp(colored, "./w:sdtContent/w:r/w:rPr/w:color/@w:val") == ["C00000"]


def test_table_styles_apply_to_values_and_numbers_keep_their_written_form() -> None:
    doc = render(
        front({"loss": "(1234)", "level": "High", "base": "7"}, extra=(
            "tables:\n"
            "  - type: financial\n    columns: [text, money]\n"
            "  - type: risk\n"
            "  - type: sensitivity\n    base_case: {row: 1, column: 2}\n"
        ))
        + "| Item | Amount |\n|---|---|\n| Loss | {{loss}} |\n| Typed | (1234) |\n\n"
        "| Risk | Impact |\n|---|---|\n| Rate | {{level}} |\n\n"
        "| BEP | Case |\n|---|---|\n| A | {{base}} |",
        profile="ib-memo",
    )
    financial, risk, sensitivity = doc.tables
    loss = financial.cell(1, 1)._tc
    assert [describe(sdt)[3] for sdt in xp(loss, ".//w:sdt")] == ["(1234)"]
    assert xp(loss, ".//w:sdt/w:sdtContent/w:r/w:rPr/w:color/@w:val") == ["C00000"]
    assert xml_text(financial.cell(2, 1)._tc) == "(1,234)"
    level = risk.cell(1, 1)._tc
    assert xp(level, ".//w:sdt/w:sdtContent/w:r/w:rPr/w:color/@w:val") == ["C00000"]
    assert xp(level, ".//w:sdt/w:sdtContent/w:r/w:rPr/w:b")
    base = sensitivity.cell(1, 1)._tc
    assert xp(base, ".//w:sdt/w:sdtContent/w:r/w:rPr/w:b")
    assert "FFFF00" in xp(base, "./w:tcPr/w:shd/@w:fill")


def test_hyperlinked_value_is_substituted_but_never_tagged() -> None:
    doc = render(front(VALUES) + "See [{{amount}}](https://example.com) and {{tenor}}.")
    [link] = xp(doc.element.body, ".//w:hyperlink")
    assert xml_text(link) == "500억원"
    assert not xp(link, ".//w:sdt") and not xp(link, "ancestor::w:sdt")
    assert [describe(sdt)[0] for sdt in controls(doc)] == ["tenor"]
    assert term_properties(doc) == {"tenor": "3년"}


def test_value_converted_to_a_legacy_footnote_reference_is_not_wrapped() -> None:
    doc = render(
        front({"note": "1", "amount": "500억원"})
        + "Claim {{amount}}^{{note}}^.\n\n## References\n\n1. Source document.",
        profile="ib-memo",
    )
    assert len(xp(doc.element.body, ".//w:footnoteReference")) == 1
    assert [describe(sdt)[0] for sdt in controls(doc)] == ["amount"]
    assert not xp(doc.element.body, ".//w:sdt//w:footnoteReference")
    assert term_properties(doc) == {"amount": "500억원"}


def test_empty_value_becomes_an_empty_control() -> None:
    doc = render(front({"blank": ""}) + "A{{blank}}B")
    [sdt] = controls(doc)
    assert describe(sdt)[:2] == ("blank", "ibrep:term:blank")
    assert describe(sdt)[3] == ""
    assert xml_text(sdt.getparent()) == "AB"
    assert term_properties(doc) == {"blank": ""}


def test_business_report_opening_uses_substituted_untagged_strings() -> None:
    doc = render(
        front(VALUES, extra='title: "가나다 {{amount}}"\ndate: "{{tenor}} 후"\n') + "Body {{tenor}}.",
        profile="business-report",
    )
    title = next(p for p in doc.paragraphs if p.style.name == "Title")
    assert title.text == "가나다 500억원" and not title._p.xpath(".//w:sdt")
    assert "작성일: 3년 후" in [p.text for p in doc.paragraphs]
    assert [describe(sdt)[0] for sdt in controls(doc)] == ["tenor"]


# ─────────────────────────────────────────────────────────────────────────────
# Snapshot
# ─────────────────────────────────────────────────────────────────────────────


def test_snapshot_records_tagged_keys_and_preserves_generator_properties() -> None:
    doc = render(
        front({**VALUES, "quote": 'A<&>"\'B', "spare": "unused"})
        + "Deal {{amount}} [{{tenor}}](https://example.com) {{quote}}."
    )
    properties = custom_properties(doc)
    assert properties["generator"][1] == "ib_report_formatter"
    assert properties["generator_profile"][1] == "plain"
    assert "generator_version" in properties
    assert term_properties(doc) == {"amount": "500억원", "quote": 'A<&>"\'B'}
    pids = [pid for pid, _ in properties.values()]
    assert len(set(pids)) == len(pids) and min(pids) >= 2
    [quote] = [sdt for sdt in controls(doc) if describe(sdt)[0] == "quote"]
    assert describe(quote)[3] == 'A<&>"\'B'


@pytest.mark.parametrize("profile", ["plain", "business-report"])
def test_unreferenced_terms_leave_the_document_unchanged(profile: str) -> None:
    body = (
        "# Title\n\n## Scope {not a term}\n\nBody **bold** [link](https://example.com)[^1].\n\n"
        "- Item\n    1. Nested\n\n> [참고] Quote\n\n| A | B |\n|---|---:|\n| x | 1234 |\n\n"
        "```text\ncode\n```\n\n[^1]: Note."
    )
    bookmark = re.compile(r'w:name="_ibrep_[^"]*"')
    outputs = []
    for markdown in ("---\ntitle: T\n---\n" + body, front(VALUES, extra="title: T\n") + body):
        doc = render(markdown, profile=profile)
        outputs.append((bookmark.sub("", doc.element.xml), custom_properties(doc)))
    assert outputs[0] == outputs[1]


def test_documents_without_tagged_values_have_no_snapshot_or_controls() -> None:
    for markdown in ["No terms here.", front(VALUES) + "Unused values only.", "---\nterms: {}\n---\nText."]:
        doc = render(markdown)
        assert controls(doc) == []
        assert term_properties(doc) == {}
        assert "generator" in custom_properties(doc)


# ─────────────────────────────────────────────────────────────────────────────
# Tags off and precedence
# ─────────────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("layout, option, tagged", [
    ("", None, True),
    ("layout:\n  term_tags: false\n", None, False),
    ("layout:\n  term_tags: false\n", True, True),
    ("layout:\n  term_tags: true\n", False, False),
    ("", False, False),
])
def test_term_tags_precedence_between_yaml_and_api(
    layout: str, option: Optional[bool], tagged: bool,
) -> None:
    doc = render(front(VALUES, extra=layout) + "Deal {{amount}} closes.", term_tags=option)
    assert bool(controls(doc)) is tagged
    assert bool(term_properties(doc)) is tagged
    assert xml_text(doc.element.body).startswith("Deal 500억원 closes.")
    if not tagged:
        assert "Deal 500억원 closes." in [p.text for p in doc.paragraphs]


@pytest.mark.parametrize("flags, tagged", [([], True), (["--no-term-tags"], False)])
def test_cli_no_term_tags_flag(tmp_path: Path, flags: List[str], tagged: bool) -> None:
    source, output = tmp_path / "terms.md", tmp_path / "terms.docx"
    source.write_text(front(VALUES) + "Deal {{amount}} closes.", encoding="utf-8")
    args = build_parser().parse_args([str(source), str(output), "--profile", "plain", "--strict", *flags])
    assert run_conversion(source, args) == 0
    doc = Document(output)
    assert bool(controls(doc)) is tagged
    assert bool(term_properties(doc)) is tagged


def test_registry_passes_term_tags_through() -> None:
    model = MarkdownParser(profile="plain").parse(front(VALUES) + "Deal {{amount}}.")
    registry = get_default_registry()
    tagged = reopen(registry.convert(model, output_format="docx", profile="plain"))
    plain = reopen(registry.convert(model, output_format="docx", profile="plain", term_tags=False))
    assert len(controls(tagged)) == 1 and controls(plain) == []


# ─────────────────────────────────────────────────────────────────────────────
# Render lifecycle
# ─────────────────────────────────────────────────────────────────────────────


def test_renderer_reuse_starts_each_render_with_an_empty_collector() -> None:
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", strict=True))
    first = render(front(VALUES) + "First {{amount}}.", renderer=renderer)
    second = render(front(VALUES) + "Second {{tenor}}.", renderer=renderer)
    third = render("Third without terms.", renderer=renderer)
    assert [describe(sdt)[0] for sdt in controls(first)] == ["amount"]
    assert [describe(sdt)[0] for sdt in controls(second)] == ["tenor"]
    assert term_properties(second) == {"tenor": "3년"}
    assert controls(third) == [] and term_properties(third) == {}


def test_tagged_document_passes_strict_rendering_and_structural_audit() -> None:
    doc = render(
        front(VALUES) + "Deal {{amount}}.[^1]\n\n| A | B |\n|---|---|\n| {{tenor}} | 2 |\n\n[^1]: Note.",
    )
    assert inspect_document(doc).issues == []
    assert controls(doc)
