"""Input-loss fixes found on a converted real term sheet (synthetic data only)."""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image as PILImage

from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, TextParser, parse_markdown_file

TERM_SHEET = (
    "---\nprofile: term-sheet\ntitle: 가나다머티리얼즈㈜ 조건 검토\n"
    "prepared_by: 라마바은행 자본시장부\ndisclaimer: 가상 조건 검토용입니다.\n---\n\n"
)


def _render_model(model, strict: bool = True):
    doc = IBDocumentRenderer(options=RenderOptions(strict=strict)).render(model)
    payload = BytesIO()
    doc.save(payload)
    return Document(payload)


def _render(markdown: str, profile: str = "business-report", strict: bool = True):
    return _render_model(MarkdownParser(profile=profile).parse(markdown), strict)


def _source(tmp_path: Path, body: str, front: str = "") -> Path:
    path = tmp_path / "memo.md"
    front = front or "---\nprofile: business-report\ntitle: 가상 문서\n---\n\n"
    path.write_text(front + body + "\n", encoding="utf-8")
    return path


def _png(path: Path, size=(40, 20)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    PILImage.new("RGB", size, "navy").save(path)
    return path


def _texts(doc) -> str:
    return "".join(doc.element.body.xpath(".//w:t/text()"))


def _cell_text(cell) -> str:
    return "".join(run.text for run in cell.runs)


# ═══════════════════════════════════════════════════════════════════════════════
# EMPHASIS DELIMITERS
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("source,expected", [
    ("금액 2 * 3 * 4", [("금액 2 * 3 * 4", False, False)]),
    ("매출처* 앞 채권<br>* 주요 매출처 :", [("매출처* 앞 채권\n* 주요 매출처 :", False, False)]),
    ("값 ** 2 ** 3", [("값 ** 2 ** 3", False, False)]),
    ("a _ b _ c", [("a _ b _ c", False, False)]),
    ("*기울임* 과 _밑줄_ 과 **굵게**", [
        ("기울임", True, False), (" 과 ", False, False), ("밑줄", True, False),
        (" 과 ", False, False), ("굵게", False, True),
    ]),
    ("*a*", [("a", True, False)]),
])
def test_emphasis_needs_flanking_delimiters(source: str, expected: list) -> None:
    for parse in (TextParser.parse_runs, TextParser.parse_runs_plain):
        assert [(run.text, run.italic, run.bold) for run in parse(source)] == expected


def test_callout_keeps_spaced_asterisks_literal() -> None:
    doc = _render("# 메모\n\n> 금액 2 * 3 * 4 확인\n")
    assert "금액 2 * 3 * 4 확인" in _texts(doc)
    italic = [
        "".join(run.xpath("./w:t/text()"))
        for run in doc.element.body.xpath(".//w:r[w:rPr/w:i[not(@w:val) or @w:val!='0']]")
    ]
    assert not any("3" in text for text in italic)


# ═══════════════════════════════════════════════════════════════════════════════
# IMAGES: PERCENT-ENCODED PATHS, TABLE CELLS, INLINE
# ═══════════════════════════════════════════════════════════════════════════════


def test_percent_encoded_image_path_finds_the_file(tmp_path: Path) -> None:
    _png(tmp_path / "img dir" / "a.png")
    _png(tmp_path / "50%.png")
    source = _source(tmp_path, "![그림](img%20dir/a.png)\n\n![비율](50%.png)")
    doc = _render_model(parse_markdown_file(str(source)))
    assert len(doc.element.body.xpath(".//w:drawing")) == 2


def test_image_in_a_table_cell_is_inserted_and_fits_the_cell(tmp_path: Path) -> None:
    _png(tmp_path / "img" / "구조도.png", size=(2400, 900))
    source = _source(tmp_path, "| 구분 | 내용 |\n|---|---|\n| 구조 | ![구조도](<img/구조도.png>)<br>① 설명 |")
    doc = _render_model(parse_markdown_file(str(source)))
    table = doc.tables[0]
    cell = table.cell(1, 1)
    assert cell._tc.xpath(".//w:drawing")
    assert cell._tc.xpath(".//wp:docPr/@descr") == ["구조도"]
    assert "![" not in cell.text and "① 설명" in cell.text
    grid = [int(width) for width in table._tbl.xpath("./w:tblGrid/w:gridCol/@w:w")]
    assert int(cell._tc.xpath(".//wp:extent/@cx")[0]) <= grid[1] * 635


def test_term_sheet_cell_image_keeps_the_following_marker_line(tmp_path: Path) -> None:
    _png(tmp_path / "구조도.png", size=(1200, 500))
    source = _source(
        tmp_path, "| 구 분 | 내 용 |\n|---|---|\n| 금융구조 | ![구조도](구조도.png)<br>① 매출대금 집금 |",
        front=TERM_SHEET,
    )
    doc = _render_model(parse_markdown_file(str(source)))
    paragraphs = doc.tables[0].cell(1, 1).paragraphs
    assert paragraphs[0]._p.xpath(".//w:drawing")
    assert paragraphs[1].text == "① 매출대금 집금"


def test_missing_cell_image_is_rejected_in_strict_and_marked_otherwise(tmp_path: Path) -> None:
    source = _source(tmp_path, "| 구분 | 내용 |\n|---|---|\n| 구조 | ![없음](missing.png) |")
    with pytest.raises(ValueError, match="[Ii]mage"):
        _render_model(parse_markdown_file(str(source)))
    doc = _render_model(parse_markdown_file(str(source)), strict=False)
    assert doc.tables[0].cell(1, 1).text == "[Image: 없음]"


def test_inline_image_in_a_paragraph_and_escaped_image_syntax(tmp_path: Path) -> None:
    _png(tmp_path / "icon.png")
    source = _source(tmp_path, "앞 ![아이콘](icon.png) 뒤\n\n글자 !\\[x](icon.png) 유지")
    doc = _render_model(parse_markdown_file(str(source)))
    paragraph = next(p for p in doc.paragraphs if p.text.startswith("앞"))
    assert paragraph._p.xpath(".//w:drawing") and paragraph.text == "앞  뒤"
    assert "글자 ![x](icon.png) 유지" in _texts(doc)


def test_standalone_html_img_is_an_image(tmp_path: Path) -> None:
    _png(tmp_path / "logo.png")
    doc = _render_model(parse_markdown_file(str(_source(tmp_path, '<img src="logo.png" alt="로고">'))))
    assert len(doc.element.body.xpath(".//w:drawing")) == 1
    assert "<img" not in _texts(doc)


# ═══════════════════════════════════════════════════════════════════════════════
# HTML TABLES
# ═══════════════════════════════════════════════════════════════════════════════

HTML_TABLE = """<table>
<tr><th>구 분</th><th colspan="2">내 용</th></tr>
<tr><td rowspan="2">참여기관</td><td>• 차주</td><td>: 가나다머티리얼즈㈜</td></tr>
<tr><td>• 대주</td><td>: 라마바은행 &amp; 신탁부</td></tr>
<tr><td>비고</td><td colspan="2">첫 줄<br>둘째 줄 2 * 3 | <b>굵게</b> &lt;br&gt; ^^</td></tr>
</table>"""


@pytest.mark.parametrize("profile", ["plain", "business-report", "ib-report"])
def test_html_table_becomes_a_word_table_with_merges(profile: str) -> None:
    doc = _render("# 표\n\n" + HTML_TABLE + "\n\n본문 계속\n", profile=profile, strict=profile != "ib-report")
    table = next(table for table in doc.tables if "참여기관" in table._tbl.xml)
    assert (len(table.rows), len(table.columns)) == (4, 3)
    assert table._tbl.xpath("./w:tr[1]/w:tc[2]/w:tcPr/w:gridSpan/@w:val") == ["2"]
    assert table._tbl.xpath("./w:tr[4]/w:tc[2]/w:tcPr/w:gridSpan/@w:val") == ["2"]
    assert len(table._tbl.xpath(".//w:vMerge")) == 2
    assert table.cell(2, 2).text == ": 라마바은행 & 신탁부"
    assert table.cell(3, 1).text == "첫 줄\n둘째 줄 2 * 3 | 굵게 <br> ^^"
    assert table._tbl.xpath(".//w:r[w:rPr/w:b][w:t='굵게']")
    assert "<t" not in _texts(doc) and "본문 계속" in _texts(doc)


SCHEDULE = """<table>
<tr><th rowspan="2">회 차</th><th>기 간</th><th colspan="2">대출</th></tr>
<tr><td>(개월)</td><td>상환액</td><td>잔액</td></tr>
<tr><td>1</td><td>3</td><td>-</td><td>300</td></tr>
<tr><td colspan="2">합 계</td><td>300</td><td>-</td></tr>
</table>"""


def test_multi_row_html_header_is_flattened_into_one_header_row() -> None:
    table = MarkdownParser(profile="plain").parse(SCHEDULE).elements[0].content
    assert [_cell_text(cell) for cell in table.rows[0].cells] == ["회 차", "기 간\n(개월)", "대출\n상환액", "대출\n잔액"]
    assert [row.is_header for row in table.rows] == [True, False, False]
    doc = _render(SCHEDULE, profile="plain")
    assert doc.tables[0]._tbl.xpath("./w:tr[3]/w:tc[1]/w:tcPr/w:gridSpan/@w:val") == ["2"]


def test_nested_single_cell_table_with_an_image_becomes_a_cell_image(tmp_path: Path) -> None:
    _png(tmp_path / "images" / "a b.png")
    body = (
        "<table>\n<tr><th>구 분</th><th>내 용</th></tr>\n"
        '<tr><td>구조</td><td><table>\n<tr><th><img src="images/a%20b.png" alt="구조도"></th></tr>\n'
        "</table><br>① 설명</td></tr>\n</table>"
    )
    doc = _render_model(parse_markdown_file(str(_source(tmp_path, body))))
    assert len(doc.tables) == 1
    cell = doc.tables[0].cell(1, 1)
    assert cell._tc.xpath(".//w:drawing") and cell.text.strip() == "① 설명"


def test_html_table_in_term_sheet_uses_label_columns() -> None:
    doc = _render(TERM_SHEET + "## 1. 개요\n\n" + HTML_TABLE + "\n", profile="term-sheet")
    assert doc.tables[0]._tbl.xpath("./w:tr[2]/w:tc[1]/w:tcPr/w:shd/@w:fill")


@pytest.mark.parametrize("source,message", [
    ("<table><tr><td>가</td></tr>\n</table><table><tr><td>나</td></tr></table>\n", "further HTML table"),
    ("<table>캡션<tr><td>가</td></tr></table>\n", "outside the HTML table cells"),
])
def test_html_text_that_cannot_be_placed_is_reported(source: str, message: str) -> None:
    model = MarkdownParser(profile="plain").parse(source)
    assert any(message in warning for warning in model.warnings)
    assert _cell_text(model.elements[0].content.rows[0].cells[0]) == "가"


# ═══════════════════════════════════════════════════════════════════════════════
# REVIEW FOLLOW-UP: IMAGES ARE ONE UNIT, HTML TEXT IS LITERAL
# ═══════════════════════════════════════════════════════════════════════════════


def _png_bytes() -> bytes:
    buffer = BytesIO()
    PILImage.new("RGB", (4, 4), "navy").save(buffer, format="PNG")
    return buffer.getvalue()


def test_image_syntax_inside_a_link_destination_stays_part_of_the_link() -> None:
    runs = TextParser.parse_runs("[링크](<https://example.com/![x](a.png)>)")
    assert [(run.text, run.image, run.hyperlink) for run in runs] == [
        ("링크", None, "https://example.com/![x](a.png)"),
    ]


def test_dollar_signs_in_image_fields_are_not_math(tmp_path: Path) -> None:
    _png(tmp_path / "cash$2026$x.png")
    model = parse_markdown_file(str(_source(tmp_path, "앞 ![로고 $1$](cash$2026$x.png) 뒤")))
    runs = model.elements[-1].content.runs
    assert [(run.image.alt_text, run.image.path) for run in runs if run.image] == [("로고 $1$", "cash$2026$x.png")]
    assert not any(run.is_latex for run in runs)
    assert len(_render_model(model).element.body.xpath(".//w:drawing")) == 1


def test_emphasis_around_an_image_keeps_its_formatting() -> None:
    runs = TextParser.parse_runs("문단 **앞 ![x](icon.png) 뒤** 끝")
    assert [(run.text, run.bold, run.image is not None) for run in runs] == [
        ("문단 ", False, False), ("앞 ", True, False), ("", True, True), (" 뒤", True, False), (" 끝", False, False),
    ]


def test_term_references_in_image_fields_stay_literal() -> None:
    front = "---\nprofile: business-report\ntitle: 가상 문서\nterms:\n  key: VALUE\n---\n\n"
    model = MarkdownParser().parse(front + "글 ![{{key}}](icon.png) 과 ![대체]({{key}}.png) 끝, {{key}}\n")
    runs = model.elements[-1].content.runs
    assert [(run.image.alt_text, run.image.path) for run in runs if run.image] == [
        ("{{key}}", "icon.png"), ("대체", "{{key}}.png"),
    ]
    assert [run.text for run in runs if run.term_key] == ["VALUE"]


@pytest.mark.parametrize("source,shape,last_row", [
    ('<table><tr><td rowspan="2">A</td><td>B</td></tr><tr></tr><tr><td>C</td><td>D</td></tr></table>',
     (2, 2), ["C", "D"]),
    ('<table><tr><th>H1</th><th>H2</th></tr><tr><td rowspan="2">A</td><td>B</td></tr><tr></tr>'
     "<tr><td>C</td><td>D</td></tr></table>", (4, 2), ["C", "D"]),
])
def test_html_empty_rows_keep_spans_in_their_columns(source: str, shape: tuple, last_row: list) -> None:
    table = _render(source, profile="plain").tables[0]
    assert (len(table.rows), len(table.columns)) == shape
    assert [cell.text for cell in table.rows[-1].cells] == last_row
    # The body row span (second case) is a real vertical merge, not literal markers.
    assert len(table._tbl.xpath(".//w:vMerge")) == (2 if shape[0] == 4 else 0)
    assert "^^" not in "".join(table._tbl.xpath(".//w:t/text()"))


def test_html_literal_tags_and_markdown_characters_stay_text() -> None:
    source = (
        "<table><tr><td>셀</td></tr><tr><td>&lt;span style=\"color:#ff0000\"&gt;literal&lt;/span&gt; "
        "&lt;br&gt; **x** [a](b.pdf) {{none}}</td></tr></table>"
    )
    doc = _render(source, profile="plain")
    assert doc.tables[0].cell(1, 0).text == '<span style="color:#ff0000">literal</span> <br> **x** [a](b.pdf) {{none}}'


def test_html_block_boundaries_and_nested_tables_break_lines() -> None:
    source = (
        "<table><tr><td>셀</td></tr><tr><td>A<div>B</div>C<p>D</p>E</td></tr>"
        "<tr><td><table><tr><td>n1</td><td>n2</td></tr><tr><td>n3</td></tr></table>after</td></tr></table>"
    )
    table = MarkdownParser(profile="plain").parse(source).elements[0].content
    assert _cell_text(table.rows[1].cells[0]) == "A\nB\nC\nD\nE"
    assert _cell_text(table.rows[2].cells[0]) == "n1 n2\nn3\nafter"


def test_html_nested_formatting_combines() -> None:
    source = (
        "<table><tr><td>셀</td></tr><tr><td><b><i>x</i></b> <b>a <b>b</b> c</b> "
        '<a href="https://example.com"><b>링크</b></a></td></tr></table>'
    )
    runs = MarkdownParser(profile="plain").parse(source).elements[0].content.rows[1].cells[0].runs
    assert [(run.text, run.bold, run.italic, run.hyperlink) for run in runs] == [
        ("x", True, True, None), (" ", False, False, None), ("a b c", True, False, None),
        (" ", False, False, None), ("링크", True, False, "https://example.com"),
    ]


def test_html_cells_substitute_terms() -> None:
    front = "---\nprofile: business-report\ntitle: 가상 문서\nterms:\n  company: 가나다머티리얼즈㈜\n---\n\n"
    model = MarkdownParser().parse(front + "<table><tr><td>구분</td></tr><tr><td>차주: {{company}}</td></tr></table>\n")
    table = next(element.content for element in model.elements if element.element_type.name == "TABLE")
    assert [(run.text, run.term_key) for run in table.rows[1].cells[0].runs] == [
        ("차주: ", None), ("가나다머티리얼즈㈜", "company"),
    ]


def test_data_uri_images_in_html(tmp_path: Path) -> None:
    import base64

    uri = "data:image/png;base64," + base64.b64encode(_png_bytes()).decode()
    body = f'<img src="{uri}" alt="x">\n\n<table><tr><td>셀</td></tr><tr><td><img src="{uri}"></td></tr></table>'
    doc = _render_model(parse_markdown_file(str(_source(tmp_path, body))))
    assert len(doc.element.body.xpath(".//w:drawing")) == 2


TERMS_FRONT = "---\nprofile: business-report\ntitle: 가상 문서\nterms:\n  mark: \" << \"\n  empty: \"\"\n  key: VALUE\n---\n\n"


def _html_cell_runs(cell_html: str, front: str = TERMS_FRONT) -> list:
    model = MarkdownParser().parse(front + f"<table><tr><td>구분</td><td>내용</td></tr><tr><td>x</td><td>{cell_html}</td></tr></table>\n")
    table = next(element.content for element in model.elements if element.element_type.name == "TABLE")
    return table.rows[1].cells[1].runs


def test_html_cell_text_that_looks_like_a_span_marker_stays_literal() -> None:
    runs = _html_cell_runs("{{mark}}")
    assert [(run.text, run.term_key) for run in runs] == [(" << ", "mark")]
    doc = _render(TERMS_FRONT + "<table><tr><td>A</td><td>B</td></tr><tr><td><b>^^</b></td><td>{{mark}}</td></tr></table>\n")
    table = doc.tables[0]
    assert not table._tbl.xpath(".//w:gridSpan") and not table._tbl.xpath(".//w:vMerge")
    assert table._tbl.xpath(".//w:r[w:rPr/w:b][w:t='^^']")
    with pytest.raises(ValueError, match="[Ii]mage"):
        _render(TERMS_FRONT + '<table><tr><td>A</td></tr><tr><td>^^<img src="missing.png" alt="없음"></td></tr></table>\n')


def test_html_runs_split_by_neutral_tags_are_whole_for_numbers_and_terms() -> None:
    doc = _render(
        "<table><tr><th>항목</th><th>2026</th></tr><tr><td>x</td><td><span>1234</span>5678</td></tr></table>\n",
        profile="ib-report", strict=False,
    )
    table = next(table for table in doc.tables if "항목" in table._tbl.xml)
    assert table.cell(1, 1).text == "12,345,678"
    assert [(run.text, run.term_key) for run in _html_cell_runs("<span>{{</span><span>key}}</span>")] == [("VALUE", "key")]


def test_html_empty_term_value_keeps_its_run() -> None:
    runs = _html_cell_runs("A{{empty}}B")
    assert [(run.text, run.term_key) for run in runs] == [("A", None), ("", "empty"), ("B", None)]


def test_html_token_shaped_text_is_not_a_term() -> None:
    runs = _html_cell_runs("&#xE000;TERM0&#xE001; {{key}}")
    assert "".join(run.text for run in runs) == "TERM0 VALUE"
    assert [run.term_key for run in runs if run.term_key] == ["key"]


@pytest.mark.parametrize("source,alt", [("a ![a\\]b](a.png) z", "a]b"), ("a \\\\![x](a.png) z", "x")])
def test_escapes_around_images_follow_backslash_parity(source: str, alt: str) -> None:
    assert [run.image.alt_text for run in TextParser.parse_runs(source) if run.image] == [alt]


def test_image_in_a_link_label_is_a_linked_image() -> None:
    runs = TextParser.parse_runs("a [![x](a.png)](b.pdf) z")
    assert [(run.text, run.image is not None, run.hyperlink) for run in runs] == [
        ("a ", False, None), ("", True, "b.pdf"), (" z", False, None),
    ]


def test_unclosed_html_table_is_reported_and_strict_rejects() -> None:
    model = MarkdownParser(profile="plain").parse("<table>\n<tr><td>가</td></tr>\n")
    assert any("HTML table" in warning for warning in model.warnings)
    with pytest.raises(ValueError):
        _render_model(model)
