"""Named section bundle precedence through parsing, rendering and the CLI."""

import logging
import sys
from dataclasses import FrozenInstanceError
from io import BytesIO

import pytest
from docx import Document

import md_to_word
from document_profiles import PROFILES, RenderOptions, resolve_options
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser


def compose(profile="ib-report", frontmatter="", **kwargs):
    model = MarkdownParser(profile=profile).parse(
        "---\ntitle: Bundle test\nsender: {organization: Example}\nrecipients: [Team]\n"
        + frontmatter + "\n---\n# Bundle test\n\nBody text."
    )
    options = RenderOptions(**kwargs)
    resolved = resolve_options(model.metadata, options)
    doc = IBDocumentRenderer(options=options).render(model)
    output = BytesIO()
    doc.save(output)
    reopened = Document(output)
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert ("TABLE OF CONTENTS" in text or "목차" in text) == resolved.toc
    assert ("면책 조항" in text) == resolved.disclaimer
    return resolved


@pytest.mark.parametrize("profile", PROFILES)
@pytest.mark.parametrize("preset, expected", [
    ("ib-report", None), ("termsheet", (False, False, False)),
    ("legal-memo", (False, True, False)), ("lecture-note", (True, True, False)),
])
def test_all_presets_on_all_profiles(profile, preset, expected):
    resolved = compose(profile, preset=preset)
    actual = (resolved.cover, resolved.toc, resolved.disclaimer)
    defaults = PROFILES[profile]
    assert actual == (expected or (defaults.cover, defaults.toc, defaults.disclaimer))


@pytest.mark.parametrize("section", ["cover", "toc", "disclaimer"])
@pytest.mark.parametrize("enabled", [False, True])
def test_each_preset_precedence_edge(section, enabled):
    override = {"include_" + section: enabled}
    yaml_flag = f"layout: {{{section}: {str(not enabled).lower()}}}"
    # Explicit field > caller preset > YAML field > YAML preset > profile.
    result = compose(frontmatter="preset: lecture-note\n" + yaml_flag,
                     preset="termsheet", **override)
    assert getattr(result, section) == enabled
    result = compose(frontmatter="preset: lecture-note\nlayout: {cover: true, toc: true, disclaimer: true}",
                     preset="termsheet")
    assert not getattr(result, section)
    result = compose(frontmatter="preset: termsheet\n" + yaml_flag)
    assert getattr(result, section) == (not enabled)
    result = compose(frontmatter="preset: termsheet")
    assert not getattr(result, section)
    # The no-op preset supplies no values at any precedence level.
    result = compose(frontmatter=yaml_flag, preset="ib-report")
    assert getattr(result, section) == (not enabled)


def test_presets_are_immutable():
    from document_profiles import PRESETS

    with pytest.raises(FrozenInstanceError):
        PRESETS["termsheet"].include_cover = True
    with pytest.raises(TypeError):
        PRESETS["new"] = RenderOptions()


def test_noop_caller_preset_preserves_frontmatter_bundle():
    resolved = compose(frontmatter="preset: termsheet", preset="ib-report")
    assert (resolved.cover, resolved.toc, resolved.disclaimer) == (False, False, False)


def test_cli_presets_listing_and_flags(tmp_path, monkeypatch, caplog):
    monkeypatch.setattr(md_to_word, "configure_logging", lambda verbose: None)
    caplog.set_level(logging.INFO)
    monkeypatch.setattr(sys, "argv", ["md-to-word", "--list-presets"])
    md_to_word.main()
    assert all(name in caplog.text for name in ("ib-report", "termsheet", "legal-memo", "lecture-note"))
    source, output = tmp_path / "source.md", tmp_path / "result.docx"
    source.write_text("---\npreset: termsheet\nlayout: {toc: false}\n---\n# Title\n\nBody.", encoding="utf-8")
    for extra, expected in [([], True), (["--no-toc"], False)]:
        monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output),
                                         "--preset", "legal-memo"] + extra)
        with pytest.raises(SystemExit) as exc:
            md_to_word.main()
        assert exc.value.code == 0
        assert ("TABLE OF CONTENTS" in [p.text for p in Document(output).paragraphs]) == expected


@pytest.mark.parametrize("use_cli", [False, True])
def test_unknown_preset_rejected_without_output(tmp_path, monkeypatch, use_cli):
    source, output = tmp_path / "bad.md", tmp_path / "bad.docx"
    source.write_text("Body." if use_cli else "---\npreset: typo\n---\nBody.", encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output)]
                        + (["--preset", "typo"] if use_cli else []))
    with pytest.raises(SystemExit) as exc:
        md_to_word.main()
    assert exc.value.code != 0
    assert not output.exists()


@pytest.mark.parametrize("flag, preset, text", [
    ("--no-cover", "lecture-note", "CONFIDENTIAL - FOR INSTITUTIONAL USE ONLY"),
    ("--no-toc", "lecture-note", "TABLE OF CONTENTS"),
    ("--no-disclaimer", "ib-report", "면책 조항"),
])
def test_explicit_cli_section_flags_beat_preset(tmp_path, monkeypatch, flag, preset, text):
    source, output = tmp_path / "input.md", tmp_path / "output.docx"
    source.write_text("---\nlayout: {cover: true, toc: true, disclaimer: true}\n---\n# Title\n\nBody.",
                      encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output), "--preset", preset, flag])
    with pytest.raises(SystemExit) as exc:
        md_to_word.main()
    assert exc.value.code == 0
    doc = Document(output)
    if flag == "--no-cover":
        # Cover-only large title run; body headings and TOC use smaller fonts.
        assert not any(run.font.size and run.font.size.pt >= 24 for p in doc.paragraphs for run in p.runs)
    else:
        assert text not in [p.text for p in doc.paragraphs]
