"""Term variables inside the term-sheet profile, through the real parser and renderer."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser

W_TAG = qn("w:tag")
HOUSE = "prepared_by: 라마바은행 자본시장부\ndisclaimer: 가상 조건 검토 자료입니다.\n"


def _render(markdown: str, strict: bool = True):
    model = MarkdownParser().parse(markdown)
    doc = IBDocumentRenderer(options=RenderOptions(strict=strict)).render(model)
    payload = BytesIO()
    doc.save(payload)
    return model, Document(payload)


def _tags(element) -> list:
    return [tag.get(qn("w:val")) for tag in element.iter(W_TAG)]


def test_opening_and_merged_cells_tag_term_values():
    md = (
        '---\nprofile: term-sheet\ntitle: "가나다머티리얼즈㈜ ABCP {{amount}}"\n'
        'date: "{{issue_date}}"\n' + HOUSE
        + 'terms:\n  amount: "300억원"\n  issue_date: "2026. 10."\n  cap: "[1.10]%p"\n---\n\n'
        "## 1. 조건\n\n"
        "| 구 분 | << | 내 용 |\n|---|---|---|\n"
        "| ABCP | 금액 | • {{amount}}<br>• CAP {{cap}} |\n"
        "| ^^ | 만기 | • 3년 |\n"
    )
    model, doc = _render(md)
    assert not model.warnings
    body = doc.element.body
    table = body.find(qn("w:tbl"))
    opening = []
    for child in body:
        if child is table:
            break
        opening.append(child)
    assert [tag for child in opening for tag in _tags(child)] == [
        "ibrep:term:amount", "ibrep:term:issue_date",
    ]
    assert _tags(table) == ["ibrep:term:amount", "ibrep:term:cap"]
    # The content cell keeps one Word paragraph per authored line.
    content_cell = doc.tables[0].cell(1, 2)
    assert len(content_cell._tc.findall(qn("w:p"))) == 2
    assert table.xpath(".//w:vMerge") and table.xpath(".//w:gridSpan")


def test_table_note_substitutes_terms_and_reports_undefined_keys():
    md = (
        "---\nprofile: plain\nterms:\n  amount: \"500\"\n"
        'tables:\n  - {note: "Amount {{amount}}; {{missing}}"}\n---\n\n'
        "| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    model = MarkdownParser().parse(md)
    table = model.elements[0].content
    assert table.note == "Amount 500; {{missing}}"
    assert any("missing" in warning for warning in model.warnings)
    with pytest.raises(ValueError):
        IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)


@pytest.mark.parametrize("value", ["", "   "])
def test_blank_substituted_subtitle_falls_back_without_stale_runs(value):
    md = (
        '---\nprofile: term-sheet\ntitle: 가상 조건표\nsubtitle: "{{s}}"\nversion: v1\n'
        + HOUSE + f'terms:\n  s: "{value}"\n---\n\n본문.\n'
    )
    model, doc = _render(md)
    assert model.metadata.subtitle == "Term Sheet"
    assert "subtitle" not in model.metadata.display_runs
    texts = [paragraph.text for paragraph in doc.paragraphs[:3]]
    assert "Term Sheet" in texts
    assert doc.core_properties.subject == "Term Sheet"


def test_title_length_is_checked_after_substitution():
    key = "k" * 40
    title = " ".join("{{" + key + "}}" for _ in range(6))
    assert len(title) > 255
    ok = (
        f'---\nprofile: term-sheet\ntitle: "{title}"\n' + HOUSE
        + f'terms:\n  {key}: "A"\n---\n\n본문.\n'
    )
    model = MarkdownParser().parse(ok)
    assert model.metadata.title == "A A A A A A"
    too_long = (
        '---\nprofile: term-sheet\ntitle: "{{t}}"\n' + HOUSE
        + 'terms:\n  t: "' + "가" * 256 + '"\n---\n\n본문.\n'
    )
    with pytest.raises(ValueError, match="255"):
        MarkdownParser().parse(too_long)
