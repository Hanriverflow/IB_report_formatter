"""Tests for YAML theme resolution, conversion, validation, and restoration.

Coverage pins the bundled default theme to every public ``IBStyle`` attribute
and verifies isolated overrides for colors and measurements.
"""

from pathlib import Path

import pytest
import yaml
from docx import Document
from docx.shared import Pt, RGBColor

import ib_renderer
from ib_renderer import IBStyle
from md_parser import Blockquote, DocumentModel, Element, ElementType, Heading, ListItem
from theme_loader import load_theme, restore_style, snapshot_style

THEMES_DIR = Path(__file__).resolve().parents[1] / "themes"


@pytest.fixture(autouse=True)
def preserve_ib_style():
    """Prevent style mutations from leaking between tests."""
    snapshot = snapshot_style()
    try:
        yield
    finally:
        restore_style(snapshot)


def _public_style_names():
    """Return public, non-callable attribute names using loader semantics."""
    return {
        name
        for name in dir(IBStyle)
        if not name.startswith("_") and not callable(getattr(IBStyle, name))
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Bundled default theme
# ═══════════════════════════════════════════════════════════════════════════════


def test_default_theme_matches_hardcoded_style():
    """The bundled default profile is semantically identical to ``IBStyle``."""
    expected = snapshot_style()

    applied = load_theme("default")

    assert set(applied) == set(expected)
    for name, original_value in expected.items():
        assert getattr(ib_renderer.STYLE, name) == original_value


def test_default_theme_contains_every_public_style_attribute():
    """The default YAML has no missing or extra public style fields."""
    theme_path = THEMES_DIR / "default.yaml"
    with theme_path.open("r", encoding="utf-8") as stream:
        theme = yaml.safe_load(stream)

    assert isinstance(theme, dict)
    assert set(theme) == _public_style_names()


# ═══════════════════════════════════════════════════════════════════════════════
# Custom themes and validation
# ═══════════════════════════════════════════════════════════════════════════════


def test_custom_theme_overrides_and_restores_typed_values(tmp_path):
    """Custom RGB and point values are converted and can be restored."""
    original = snapshot_style()
    theme_path = tmp_path / "custom.yaml"
    theme_path.write_text('NAVY: "112233"\nH1_SIZE: 20.0\n', encoding="utf-8")

    load_theme(str(theme_path))

    assert ib_renderer.STYLE.NAVY == RGBColor(0x11, 0x22, 0x33)
    assert ib_renderer.STYLE.H1_SIZE == Pt(20)

    restore_style(original)
    assert ib_renderer.STYLE.NAVY == original["NAVY"]
    assert ib_renderer.STYLE.H1_SIZE == original["H1_SIZE"]


def test_unknown_theme_key_raises_value_error(tmp_path):
    """Unknown keys are rejected as style-profile drift."""
    theme_path = tmp_path / "unknown.yaml"
    theme_path.write_text('NOT_A_REAL_FIELD: "x"\n', encoding="utf-8")

    with pytest.raises(ValueError, match="NOT_A_REAL_FIELD"):
        load_theme(str(theme_path))


def test_missing_theme_raises_file_not_found_error():
    """A missing path and bundled name produce a clear resolution error."""
    with pytest.raises(FileNotFoundError, match="this-theme-does-not-exist"):
        load_theme("this-theme-does-not-exist")


# ═══════════════════════════════════════════════════════════════════════════════
# END-TO-END RENDERING PROOF
# ═══════════════════════════════════════════════════════════════════════════════


def test_loaded_theme_changes_rendered_bullet_color(tmp_path):
    """A loaded theme changes the serialized bullet-character run color."""
    theme_path = tmp_path / "bullet-theme.yaml"
    theme_path.write_text('NAVY: "1A2B3C"\n', encoding="utf-8")
    load_theme(str(theme_path))

    model = DocumentModel(
        elements=[
            Element(
                element_type=ElementType.BULLET_LIST,
                content=ListItem(text="Bullet item", indent_level=0),
            )
        ]
    )
    doc = ib_renderer.IBDocumentRenderer().render(model)
    output_path = tmp_path / "themed-bullet.docx"
    doc.save(str(output_path))

    reopened = Document(str(output_path))
    matching_paragraphs = [
        paragraph
        for paragraph in reopened.paragraphs
        if "Bullet item" in paragraph.text
    ]
    assert len(matching_paragraphs) == 1

    run_colors = [
        run.font.color.rgb
        for run in matching_paragraphs[0].runs
        if run.font.color.rgb is not None
    ]
    assert RGBColor(0x1A, 0x2B, 0x3C) in run_colors
    assert RGBColor(0x00, 0x33, 0x66) not in run_colors


def test_loaded_theme_changes_rendered_heading_color(tmp_path):
    """A loaded theme changes the serialized heading run color."""
    theme_path = tmp_path / "heading-theme.yaml"
    theme_path.write_text('NAVY: "1A2B3C"\n', encoding="utf-8")
    load_theme(str(theme_path))

    model = DocumentModel(
        elements=[
            Element(
                element_type=ElementType.HEADING_1,
                content=Heading(text="Themed Heading", level=1),
            ),
        ]
    )
    doc = ib_renderer.IBDocumentRenderer().render(model)
    output_path = tmp_path / "themed-heading.docx"
    doc.save(str(output_path))

    reopened = Document(str(output_path))
    matching_paragraphs = [
        paragraph
        for paragraph in reopened.paragraphs
        if paragraph.text == "Themed Heading"
        and paragraph.style is not None
        and paragraph.style.name == "Heading 1"
    ]
    assert len(matching_paragraphs) == 1

    run_colors = [
        run.font.color.rgb
        for run in matching_paragraphs[0].runs
        if run.font.color.rgb is not None
    ]
    assert RGBColor(0x1A, 0x2B, 0x3C) in run_colors
    assert RGBColor(0x00, 0x33, 0x66) not in run_colors


def test_loaded_theme_changes_rendered_callout_color(tmp_path):
    """A loaded theme changes the serialized callout title run color."""
    theme_path = tmp_path / "callout-theme.yaml"
    theme_path.write_text('NAVY: "1A2B3C"\n', encoding="utf-8")
    load_theme(str(theme_path))

    model = DocumentModel(
        elements=[
            Element(
                element_type=ElementType.BLOCKQUOTE,
                content=Blockquote(title="KEY INSIGHT", text="Some insight text."),
            )
        ]
    )
    doc = ib_renderer.IBDocumentRenderer().render(model)
    output_path = tmp_path / "themed-callout.docx"
    doc.save(str(output_path))

    reopened = Document(str(output_path))
    matching_paragraphs = [
        paragraph
        for table in reopened.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if "KEY INSIGHT" in paragraph.text
    ]
    assert len(matching_paragraphs) == 1

    run_colors = [
        run.font.color.rgb
        for run in matching_paragraphs[0].runs
        if run.font.color.rgb is not None
    ]
    assert RGBColor(0x1A, 0x2B, 0x3C) in run_colors
    assert RGBColor(0x00, 0x33, 0x66) not in run_colors
