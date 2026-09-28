"""Installed raster fonts and loss diagnostics through parsing and saved DOCX."""

import sys
import warnings
from io import BytesIO
from uuid import UUID

import pytest
from docx import Document
from docx.oxml.ns import qn
from matplotlib import font_manager
from matplotlib.figure import Figure
from matplotlib.text import Text

import md_to_word
import render_styles
from document_profiles import RenderOptions
from ib_renderer import FontPolicy, IBDocumentRenderer
from md_parser import MarkdownParser

CHART = """```chart
type: bar
title: 가나
labels: [가, 나]
series:
  - name: 가
    values: [1, 2]
  - name: 나
    values: [2, 3]
y_label: 가
source: 나
```
"""
DIAGRAM = """```diagram:flow
title: 가
boxes:
  - {id: a, label: 가, pos: [0, 0]}
  - {id: b, label: 나, pos: [3, 0]}
arrows:
  - {from: a, to: b, label: 가}
notes: [나]
```
"""
RASTERS = [CHART, r"$$ \text{가} = 1 $$", r"Before $\text{가}$ after.", DIAGRAM]


def render(markdown: str, **options):
    """Return diagnostics and a reopened DOCX from the production render path."""
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", charts=True, **options))
    doc = renderer.render(MarkdownParser(profile="plain").parse(markdown))
    buffer = BytesIO()
    doc.save(buffer)
    return renderer, Document(buffer)


@pytest.fixture
def latin_only(monkeypatch: pytest.MonkeyPatch) -> None:
    """Simulate a machine with only Matplotlib's bundled DejaVu font."""
    path = font_manager.findfont("DejaVu Sans", fallback_to_default=False)
    monkeypatch.setattr(render_styles, "_installed_raster_font",
                        lambda family: path if family == "DejaVu Sans" else None, raising=False)
    # Loss is intentional in these scenarios and asserted via the public diagnostic.
    warnings.filterwarnings("ignore", message=r"Glyph .* missing from font", category=UserWarning)


@pytest.mark.parametrize("markdown", RASTERS)
def test_missing_cjk_font_warns_visibly_and_strict_rejects(markdown, latin_only, caplog):
    renderer, doc = render(markdown, strict=False)
    diagnostic = "; ".join(renderer.errors)
    assert "CJK raster font" in diagnostic
    assert "Malgun Gothic" in diagnostic and "NanumGothic" in diagnostic
    assert "CJK raster font" in caplog.text
    assert "CJK raster font" in "".join(doc.element.xpath(".//w:t/text()"))
    assert doc.inline_shapes
    with pytest.raises(ValueError, match="CJK raster font"):
        render(markdown, strict=True)


@pytest.mark.parametrize("text", ["漢字", "かな", "カナ", "ᄀ", "ㄱ", "𠀀"])
def test_other_cjk_text_is_also_diagnosed(text, latin_only):
    with pytest.raises(ValueError, match="CJK raster font"):
        render(CHART.replace("가", text).replace("나", text), strict=True)


@pytest.mark.parametrize("markdown", RASTERS)
def test_latin_only_images_need_no_cjk_font(markdown, latin_only):
    renderer, doc = render(markdown.replace("가", "A").replace("나", "B"), strict=True)
    assert not renderer.errors
    assert doc.inline_shapes


@pytest.mark.parametrize("existing", [False, True])
def test_cli_missing_font_rejects_before_save(tmp_path, monkeypatch, latin_only, existing, capsys):
    source, output = tmp_path / "chart.md", tmp_path / "chart.docx"
    source.write_text(CHART, encoding="utf-8")
    if existing:
        output.write_bytes(b"keep existing output")
    monkeypatch.setattr(sys, "argv", ["md-to-word", str(source), str(output),
                                    "--profile", "plain", "--charts", "--strict"])
    with pytest.raises(SystemExit) as exc:
        md_to_word.main()
    assert exc.value.code == 1
    assert "CJK raster font" in capsys.readouterr().out
    if existing:
        assert output.read_bytes() == b"keep existing output"
    else:
        assert not output.exists()


@pytest.fixture
def installed_test_fonts(tmp_path, monkeypatch):
    """Synthetic glyphs test font routing without depending on OS fonts.

    These deliberately map Hangul to a test glyph; real Hangul appearance is
    covered separately by the environment test, not by this fixture.
    """
    from fontTools.ttLib import TTFont

    monkeypatch.setattr(font_manager.fontManager, "ttflist", list(font_manager.fontManager.ttflist))
    source = font_manager.findfont("DejaVu Sans", fallback_to_default=False)
    paths = {}
    for family in ("NanumGothic", "Test Theme CJK"):
        path = tmp_path / (family.replace(" ", "-") + ".ttf")
        with TTFont(source) as font:
            for table in font["cmap"].tables:
                if table.isUnicode():
                    for character in "가나※":
                        table.cmap[ord(character)] = "A"
            for record in font["name"].names:
                if record.nameID in {1, 3, 4, 6, 16}:
                    record.string = family.encode(record.getEncoding())
            font.save(path)
        font_manager.fontManager.addfont(path)
        paths[family] = str(path)
    monkeypatch.setattr(render_styles, "_installed_raster_font", paths.get, raising=False)
    yield paths
    font_manager.fontManager._findfont_cached.cache_clear()


@pytest.mark.parametrize("markdown", RASTERS)
def test_nanum_fallback_is_drawn_without_missing_glyph_warnings(markdown, installed_test_fonts, monkeypatch):
    families = []
    saved = Figure.savefig

    def inspect(figure, *args, **kwargs):
        families.extend(artist.get_fontfamily() for artist in figure.findobj(match=Text)
                        if "가" in artist.get_text() or "나" in artist.get_text())
        return saved(figure, *args, **kwargs)

    monkeypatch.setattr(Figure, "savefig", inspect)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        renderer, doc = render(markdown, strict=True)
    assert not renderer.errors
    assert doc.inline_shapes
    assert families and all(family == ["NanumGothic"] for family in families)
    assert not [item for item in caught if "missing from font" in str(item.message)]


@pytest.mark.parametrize("markdown", RASTERS)
def test_theme_preferred_per_request_and_docx_font_declarations_unchanged(
    installed_test_fonts, tmp_path, monkeypatch, markdown,
):
    monkeypatch.setattr("ib_renderer.uuid4", lambda: UUID(int=1))
    theme = tmp_path / "theme.yaml"
    theme.write_text("KOREAN_FONT: Test Theme CJK\n", encoding="utf-8")
    families = []
    saved = Figure.savefig

    def inspect(figure, *args, **kwargs):
        families.append(next(artist.get_fontfamily() for artist in figure.findobj(match=Text)
                             if "가" in artist.get_text()))
        return saved(figure, *args, **kwargs)

    monkeypatch.setattr(Figure, "savefig", inspect)
    _, preferred_doc = render("가\n\n" + markdown, theme=str(theme), strict=True)
    monkeypatch.setattr(render_styles, "_installed_raster_font",
                        lambda family: installed_test_fonts.get(family) if family == "NanumGothic" else None)
    _, fallback_doc = render("가\n\n" + markdown, theme=str(theme), strict=True)
    _, default_doc = render("가\n\n" + markdown, strict=True)
    assert families == [["Test Theme CJK"], ["NanumGothic"], ["NanumGothic"]]
    # Equation/diagram image names contain independently generated temp filenames, and
    # content-sized rasters change extent with the font's metrics; neither is a font declaration.
    for doc in (preferred_doc, fallback_doc):
        for picture in doc.element.xpath(".//pic:cNvPr"):
            picture.set("name", "raster.png")
        for extent in doc.element.xpath(".//wp:extent | .//a:ext[@cx]"):
            extent.set("cx", "0")
            extent.set("cy", "0")
    assert preferred_doc.element.xml == fallback_doc.element.xml
    assert preferred_doc.styles.element.xml == fallback_doc.styles.element.xml
    assert preferred_doc.part.numbering_part.element.xml == fallback_doc.part.numbering_part.element.xml
    assert preferred_doc.paragraphs[0].runs[0]._r.rPr.rFonts.get(qn("w:eastAsia")) == "Test Theme CJK"
    assert default_doc.paragraphs[0].runs[0]._r.rPr.rFonts.get(qn("w:eastAsia")) == FontPolicy.resolve_korean_font()


def test_installed_latin_theme_does_not_count_as_cjk(tmp_path, latin_only):
    theme = tmp_path / "latin.yaml"
    theme.write_text("KOREAN_FONT: DejaVu Sans\n", encoding="utf-8")
    with pytest.raises(ValueError, match="CJK raster font"):
        render(CHART, theme=str(theme), strict=True)


def test_font_diagnostics_reset_on_renderer_reuse(latin_only):
    renderer = IBDocumentRenderer(options=RenderOptions(profile="plain", charts=True, strict=True))
    with pytest.raises(ValueError, match="CJK raster font"):
        renderer.render(MarkdownParser(profile="plain").parse(CHART))
    doc = renderer.render(MarkdownParser(profile="plain").parse("Next document"))
    assert not renderer.errors
    assert doc.paragraphs[0].text == "Next document"


def test_installed_lookup_is_cached_without_capturing_preference(monkeypatch):
    calls = []

    def findfont(properties, *, fallback_to_default):
        assert fallback_to_default is False
        family = properties.get_family()[0]
        calls.append(family)
        if family == "Unavailable Test Family":
            raise ValueError("not installed")
        return family + ".ttf"

    monkeypatch.setattr(font_manager, "findfont", findfont)
    render_styles._installed_raster_font.cache_clear()
    try:
        for _ in range(2):
            assert render_styles.RasterFontPolicy.resolve("First Test Family", "ASCII") == "First Test Family"
            assert render_styles.RasterFontPolicy.resolve("Second Test Family", "ASCII") == "Second Test Family"
            assert render_styles.RasterFontPolicy.resolve("Unavailable Test Family", "ASCII") == "Malgun Gothic"
        assert calls == ["First Test Family", "Second Test Family", "Unavailable Test Family", "Malgun Gothic"]
    finally:
        render_styles._installed_raster_font.cache_clear()


def test_real_installed_cjk_font_renders_hangul_without_glyph_loss():
    from matplotlib.ft2font import FT2Font

    candidates = ("Malgun Gothic", "Apple SD Gothic Neo", "AppleGothic", "NanumGothic",
                  "NanumBarunGothic", "Noto Sans CJK KR", "Noto Sans KR", "Source Han Sans KR", "UnDotum")
    for family in candidates:
        try:
            path = font_manager.findfont(font_manager.FontProperties(family=[family]), fallback_to_default=False)
            if all(ord(character) in FT2Font(path).get_charmap() for character in "가나"):
                break
        except ValueError:
            continue
    else:
        pytest.skip("No installed CJK font with Hangul coverage")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        renderer, doc = render(CHART, strict=True)
    assert doc.inline_shapes and not renderer.errors
    assert not [item for item in caught if "missing from font" in str(item.message)]
