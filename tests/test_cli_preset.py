"""Preset-option tests for the Markdown-to-Word CLI."""

import pytest
from docx import Document

import md_to_word
import theme_loader
from preset_loader import PresetConfig


@pytest.fixture(autouse=True)
def preserve_ib_style():
    """Restore all shared IB style values after every test."""
    snapshot = theme_loader.snapshot_style()
    yield
    theme_loader.restore_style(snapshot)


def test_merge_preset_explicit_cli_values_win():
    """Explicit section, separator, and theme flags should beat preset values."""
    args = md_to_word.build_parser().parse_args(
        [
            "report.md",
            "--no-cover",
            "--separator-mode",
            "rule",
            "--theme",
            "navy",
        ]
    )
    preset = PresetConfig(
        theme="ib-report",
        include_cover=True,
        separator_mode="page-break",
    )

    merged = md_to_word.merge_preset(args, preset)

    assert merged.no_cover is True
    assert merged.theme == "navy"
    assert merged.separator_mode == "rule"


def test_merge_preset_values_beat_built_in_defaults():
    """Preset values should supply fields left unspecified on the CLI."""
    args = md_to_word.build_parser().parse_args(["report.md"])
    preset = PresetConfig(
        theme="navy",
        include_cover=False,
        separator_mode="rule",
    )

    merged = md_to_word.merge_preset(args, preset)

    assert merged.no_cover is True
    assert merged.theme == "navy"
    assert merged.separator_mode == "rule"


@pytest.mark.parametrize("preset", [None, PresetConfig()])
def test_merge_preset_restores_separator_default_when_values_are_absent(preset):
    """No separator override should retain the observable auto default."""
    args = md_to_word.build_parser().parse_args(["report.md"])

    merged = md_to_word.merge_preset(args, preset)

    assert args.separator_mode is None
    assert merged.separator_mode == "auto"


def test_preset_cannot_force_section_on_against_explicit_skip():
    """A preset include value should not undo an explicit no-toc flag."""
    args = md_to_word.build_parser().parse_args(["report.md", "--no-toc"])
    preset = PresetConfig(include_toc=True)

    merged = md_to_word.merge_preset(args, preset)

    assert merged.no_toc is True


def test_termsheet_preset_suppresses_toc_and_disclaimer(tmp_path):
    """Termsheet conversion should omit sections that default conversion renders."""
    input_path = tmp_path / "report.md"
    preset_output = tmp_path / "termsheet.docx"
    default_output = tmp_path / "default.docx"
    input_path.write_text("# Test Report\n\nOne paragraph.\n", encoding="utf-8")

    preset_args = md_to_word.build_parser().parse_args(
        [str(input_path), str(preset_output), "--preset", "termsheet"]
    )
    preset_exit_code, preset = md_to_word.load_document_preset(preset_args.preset)
    assert preset_exit_code == 0
    merged_preset_args = md_to_word.merge_preset(preset_args, preset)
    assert md_to_word.run_conversion(input_path, merged_preset_args) == 0

    preset_paragraphs = [
        paragraph.text for paragraph in Document(str(preset_output)).paragraphs
    ]
    assert "TABLE OF CONTENTS" not in preset_paragraphs
    assert "면책 조항" not in preset_paragraphs

    default_args = md_to_word.build_parser().parse_args(
        [str(input_path), str(default_output)]
    )
    default_exit_code, default_preset = md_to_word.load_document_preset(
        default_args.preset
    )
    assert default_exit_code == 0
    merged_default_args = md_to_word.merge_preset(default_args, default_preset)
    assert merged_default_args.separator_mode == "auto"
    assert md_to_word.run_conversion(input_path, merged_default_args) == 0

    default_paragraphs = [
        paragraph.text for paragraph in Document(str(default_output)).paragraphs
    ]
    assert "TABLE OF CONTENTS" in default_paragraphs
    assert "면책 조항" in default_paragraphs


def test_unknown_preset_returns_nonzero_exit_code():
    """An unknown preset should fail loading without returning a config."""
    exit_code, preset = md_to_word.load_document_preset(
        "does_not_exist_preset_xyz"
    )

    assert exit_code == 1
    assert preset is None
