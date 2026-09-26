"""Release regressions for real office and IB authoring inputs."""

import base64
from pathlib import Path

import pytest

from document_profiles import RenderOptions
from docx_audit import inspect_document
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, parse_markdown_file


def test_valid_long_yaml_metadata_keeps_selected_profile():
    title = "문서 표준화 " * 30
    model = MarkdownParser().parse(
        "---\nprofile: business-report\ntitle: '" + title + "'\n---\n본문."
    )
    assert model.metadata.profile == "business-report"
    assert model.metadata.title == title


def test_declared_invalid_yaml_is_not_rendered_as_an_ib_report():
    with pytest.raises(ValueError, match="frontmatter"):
        MarkdownParser().parse(
            "---\nprofile: office-letter\nsender:\n  organization: [invalid\n---\n본문."
        )


def test_file_parser_profile_override_wins_even_over_unknown_yaml_profile(tmp_path):
    source = tmp_path / "input.md"
    source.write_text("---\nprofile: retired-profile\n---\n본문.", encoding="utf-8")
    model = parse_markdown_file(str(source), profile="plain")
    assert model.metadata.profile == "plain"


def test_relative_image_is_resolved_from_markdown_not_process_directory(tmp_path):
    image = tmp_path / "mark.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO7Z8eQAAAAASUVORK5CYII="
        )
    )
    source = tmp_path / "report.md"
    source.write_text("---\nprofile: plain\n---\n![표식](mark.png)", encoding="utf-8")
    model = parse_markdown_file(str(source))
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    assert len(doc.inline_shapes) == 1


def test_plain_superscript_is_not_mistaken_for_a_footnote():
    model = MarkdownParser(profile="plain").parse("면적 m^2^와 실제 각주.[^2]\n\n[^2]: 각주 설명.")
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    assert len(doc.element.xpath(".//w:footnoteReference")) == 1
    assert "2" in doc.element.xpath(".//w:r[w:rPr/w:vertAlign]/w:t/text()")


def test_explicit_footnote_in_heading_is_native():
    model = MarkdownParser(profile="plain").parse("## 근거[^1]\n\n[^1]: 설명.")
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    assert len(doc.element.xpath(".//w:footnoteReference")) == 1


def test_explicit_footnote_in_table_header_is_native():
    model = MarkdownParser(profile="plain").parse(
        "| 지표[^1] | 값 |\n|---|---|\n| 항목 | 10 |\n\n[^1]: 단위 설명."
    )
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(model)
    assert len(doc.element.xpath(".//w:footnoteReference")) == 1


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "euc-kr", "cp949"])
def test_korean_file_text_is_decoded_without_statistical_corruption(tmp_path, encoding):
    source = tmp_path / "report.md"
    source.write_bytes("# 한글 제목\n\n표식".encode(encoding))
    model = parse_markdown_file(str(source), profile="plain")
    assert model.metadata.title == "한글 제목"
    assert model.elements[-1].content.text == "표식"


def test_toc_preview_is_a_field_result_not_a_duplicate_outside_the_field():
    model = MarkdownParser().parse("---\ntitle: Report\n---\n## Section Alpha\n\nBody.")
    doc = IBDocumentRenderer(options=RenderOptions(include_disclaimer=False)).render(model)
    nodes = list(doc.element.iter())
    from docx.oxml.ns import qn

    toc_instruction = next(
        i
        for i, node in enumerate(nodes)
        if node.tag == qn("w:instrText") and "TOC " in (node.text or "")
    )
    toc_end = next(
        i
        for i, node in enumerate(nodes[toc_instruction:], toc_instruction)
        if node.tag == qn("w:fldChar") and node.get(qn("w:fldCharType")) == "end"
    )
    preview = next(
        i
        for i, node in enumerate(nodes[toc_instruction:], toc_instruction)
        if node.tag == qn("w:t") and node.text == "Section Alpha"
    )
    assert preview < toc_end
    toc_title = next(p for p in doc.paragraphs if p.text == "TABLE OF CONTENTS")
    assert not toc_title.style.name.startswith("Heading ")


def test_office_title_uses_neutral_profile_colour():
    doc = IBDocumentRenderer().render(MarkdownParser(profile="business-report").parse("본문."))
    assert str(doc.styles["Title"].font.color.rgb) == "202020"
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert not doc.styles["Title"].element.xpath(
        "./w:rPr/w:rFonts/@w:eastAsiaTheme | ./w:rPr/w:rFonts/@w:asciiTheme"
    )


def test_plain_empty_header_has_no_separator():
    doc = IBDocumentRenderer().render(MarkdownParser(profile="plain").parse("본문."))
    assert not doc.sections[0].header._element.xpath(".//w:pBdr")


def test_memo_spacing_is_compact_without_changing_legacy_report():
    from document_profiles import get_profile, load_style

    memo = load_style(get_profile("ib-memo"))
    report = load_style(get_profile("ib-report"))
    assert memo.BODY_SPACE_AFTER.pt == 6
    assert report.BODY_SPACE_AFTER.pt == 8
    assert memo.TOP_MARGIN < report.TOP_MARGIN


@pytest.mark.parametrize(
    "source",
    sorted((Path(__file__).resolve().parent.parent / "samples" / "qa").glob("*.md")),
    ids=lambda path: path.stem,
)
def test_pagination_fixtures_render_without_structural_errors(source):
    doc = IBDocumentRenderer(options=RenderOptions(strict=True)).render(
        parse_markdown_file(str(source))
    )
    assert not inspect_document(doc).issues


def test_header_right_tab_tracks_each_landscape_section_width():
    source = Path(__file__).resolve().parent.parent / "samples" / "qa" / "landscape.md"
    doc = IBDocumentRenderer().render(parse_markdown_file(str(source)))
    for section in doc.sections:
        tabs = section.header.paragraphs[0].paragraph_format.tab_stops
        assert len(section.header.paragraphs[0].style.paragraph_format.tab_stops) == 0
        assert len(tabs) == 1
        # Word serializes dimensions in twips (635 EMU).
        assert (
            abs(
                tabs[0].position - (section.page_width - section.left_margin - section.right_margin)
            )
            <= 635
        )
