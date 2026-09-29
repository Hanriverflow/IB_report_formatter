"""Improvements found on a real internal memo: inline code, local file links, Korean spacing."""

import logging
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, TextParser, parse_markdown_file


def _render(markdown: str, profile: str = "business-report", strict: bool = True):
    model = MarkdownParser(profile=profile).parse(markdown)
    doc = IBDocumentRenderer(options=RenderOptions(profile=profile, strict=strict)).render(model)
    payload = BytesIO()
    doc.save(payload)
    return Document(payload)


def _hyperlink_targets(doc) -> list:
    targets = []
    for link in doc.element.body.iter(qn("w:hyperlink")):
        relationship = doc.part.rels[link.get(qn("r:id"))]
        assert relationship.reltype == RT.HYPERLINK and relationship.is_external
        # python-docx's run elements repeat text in itertext(); read w:t nodes instead.
        targets.append(("".join(link.xpath(".//w:t/text()")), relationship.target_ref))
    return targets


# ═══════════════════════════════════════════════════════════════════════════════
# INLINE CODE
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("source,expected", [
    ("계좌 `022-1` 참고", [("계좌 ", False), ("022-1", True), (" 참고", False)]),
    ("`` a`b ``", [("a`b", True)]),
    (r"\`literal\`", [("`literal`", False)]),
])
def test_code_spans_become_literal_code_runs(source: str, expected: list) -> None:
    assert [(run.text, run.code) for run in TextParser.parse_runs(source)] == expected


def test_code_run_keeps_surrounding_emphasis_and_link() -> None:
    bold = TextParser.parse_runs("**`코드`** 끝")[0]
    assert (bold.text, bold.code, bold.bold) == ("코드", True, True)
    link = TextParser.parse_runs("[`x`](https://example.com)")[0]
    assert (link.text, link.code, link.hyperlink) == ("x", True, "https://example.com")


def test_rendered_code_has_no_backticks_and_uses_the_code_font() -> None:
    doc = _render("# 메모\n\n계좌번호는 `022-9200-0000-000`입니다.\n")
    paragraph = next(p for p in doc.paragraphs if "022-9200-0000-000" in p.text)
    assert "`" not in paragraph.text
    code_run = next(run for run in paragraph.runs if run.text == "022-9200-0000-000")
    assert code_run.font.name == "Consolas"


def test_code_in_a_money_column_keeps_its_literal_digits() -> None:
    markdown = (
        "---\nprofile: business-report\ntitle: 가상 메모\n"
        "tables:\n  - {columns: [text, money]}\n---\n\n"
        "| 항목 | 금액 |\n|---|---:|\n| 코드 | `1234567` |\n| 숫자 | 1234567 |\n"
    )
    cells = [cell.text for cell in _render(markdown).tables[0].columns[1].cells]
    assert cells[1:] == ["1234567", "1,234,567"]


# ═══════════════════════════════════════════════════════════════════════════════
# LOCAL FILE LINKS
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("source,target", [
    ("[웹](https://example.com/a)", "https://example.com/a"),
    ("[정리](03_정리.md)", "03_%EC%A0%95%EB%A6%AC.md"),
    ("[꺾쇠](<C:/가상 폴더/정리.md>)", "file:///C:/%EA%B0%80%EC%83%81%20%ED%8F%B4%EB%8D%94/%EC%A0%95%EB%A6%AC.md"),
    (r"[윈도](C:\docs\a.pdf)", "file:///C:/docs/a.pdf"),
    ("[상위](../자료/표.xlsx)", "../%EC%9E%90%EB%A3%8C/%ED%91%9C.xlsx"),
])
def test_local_and_web_links_become_hyperlinks(source: str, target: str) -> None:
    runs = TextParser.parse_runs(source)
    assert len(runs) == 1 and runs[0].hyperlink == target


@pytest.mark.parametrize("source", [
    "[참고](별첨)", "[주](Co.Ltd)", "[v](1.0)", "[도메인](www.example.com/a)", "[공백](./a b.md)",
])
def test_non_path_parentheses_stay_literal(source: str) -> None:
    runs = TextParser.parse_runs(source)
    assert "".join(run.text for run in runs) == source
    assert not any(run.hyperlink for run in runs)


def test_file_links_inside_the_folder_become_relative(tmp_path: Path, caplog) -> None:
    inside = tmp_path / "sub" / "가상 정리.md"
    outside = "C:/elsewhere/other.md"
    source = tmp_path / "memo.md"
    source.write_text(
        "# 메모\n\n[안쪽][in] · [바깥][out] · [인라인](<" + inside.as_posix() + ">)\n\n"
        "[in]: <" + inside.as_posix() + ">\n[out]: <" + outside + ">\n",
        encoding="utf-8",
    )
    with caplog.at_level(logging.WARNING):
        model = parse_markdown_file(str(source))
    runs = model.elements[-1].content.runs
    assert [(run.text, run.hyperlink) for run in runs if run.hyperlink] == [
        ("안쪽", "sub/%EA%B0%80%EC%83%81%20%EC%A0%95%EB%A6%AC.md"),
        ("바깥", "file:///C:/elsewhere/other.md"),
        ("인라인", "sub/%EA%B0%80%EC%83%81%20%EC%A0%95%EB%A6%AC.md"),
    ]
    assert any("outside the document folder" in record.message for record in caplog.records)
    assert not model.warnings


def test_rendered_local_link_is_an_external_hyperlink_showing_only_its_label() -> None:
    doc = _render("# 메모\n\n자세한 내용은 [계좌 정리](03_계좌_정리.md)를 참고.\n")
    assert _hyperlink_targets(doc) == [("계좌 정리", "03_%EA%B3%84%EC%A2%8C_%EC%A0%95%EB%A6%AC.md")]
    assert "03_계좌_정리.md" not in "".join(p.text for p in doc.paragraphs)


# ═══════════════════════════════════════════════════════════════════════════════
# KOREAN SPACING
# ═══════════════════════════════════════════════════════════════════════════════


def _auto_spacing(doc) -> dict:
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
    return {
        tag: (defaults.find(qn("w:" + tag)).get(qn("w:val")) if defaults.find(qn("w:" + tag)) is not None else None)
        for tag in ("autoSpaceDE", "autoSpaceDN")
    }


@pytest.mark.parametrize("profile", ["business-report", "meeting-minutes", "plain"])
def test_general_profiles_turn_off_korean_auto_spacing(profile: str) -> None:
    doc = _render("# 메모\n\nSPC는 제2종 300억원을 보유합니다.\n", profile=profile)
    assert _auto_spacing(doc) == {"autoSpaceDE": "0", "autoSpaceDN": "0"}


@pytest.mark.parametrize("profile", ["ib-report", "ib-memo"])
def test_ib_profiles_keep_word_default_spacing(profile: str) -> None:
    doc = _render("# Memo\n\nBody.\n", profile=profile, strict=False)
    assert _auto_spacing(doc) == {"autoSpaceDE": None, "autoSpaceDN": None}
