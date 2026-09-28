"""Cover-free titles through the production parser, renderer and CLI."""

import sys
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import List

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Pt

import md_to_word
from document_model import Heading
from document_profiles import RenderOptions
from ib_renderer import IBDocumentRenderer
from md_parser import MarkdownParser, parse_markdown_file
from render_styles import STYLE, IBStyle

ROOT = Path(__file__).resolve().parents[1]
YAML_TITLE = (
    '---\ntitle: Report title\nsubtitle: Report subtitle\n'
    'date: "2026-09-29"\nanalyst: Kim\n---\n'
)
INFERRED_TITLE = (
    '# Report title\n## Report subtitle\n\n'
    '**Date:** 2026-09-29\n**Analyst:** Kim\n\n'
)
BODY = 'Introduction.\n\n## Body scope\n\nBody text.\n'


def _reopen(doc: DocxDocument) -> DocxDocument:
    """Inspect the saved DOCX, including its resolved styles and fields."""
    payload = BytesIO()
    doc.save(payload)
    return Document(payload)


def _assert_non_outline(paragraph) -> None:
    """Titles must stay out of Word's outline-based TOC on field refresh."""
    assert not paragraph._p.xpath('./w:pPr/w:outlineLvl')
    style = paragraph.style
    while style is not None:
        assert not style.element.xpath('./w:pPr/w:outlineLvl')
        style = style.base_style


@pytest.mark.parametrize('source', [YAML_TITLE, INFERRED_TITLE], ids=['yaml', 'inferred'])
@pytest.mark.parametrize('toc', [False, True])
@pytest.mark.parametrize('strict', [False, True])
def test_cover_off_title_subtitle_metadata_once(source: str, toc: bool, strict: bool) -> None:
    model = MarkdownParser().parse(source + BODY)
    original = deepcopy(model)
    renderer = IBDocumentRenderer(options=RenderOptions(
        include_cover=False, include_toc=toc, include_disclaimer=False, strict=strict,
    ))
    doc = _reopen(renderer.render(model))
    texts = [p.text for p in doc.paragraphs]
    assert texts.count('Report title') == 1
    assert texts.count('Report subtitle') == 1
    title_index = texts.index('Report title')
    assert texts[title_index:title_index + 5] == [
        'Report title', 'Report subtitle', '작성일: 2026-09-29', '작성자: Kim',
        'TABLE OF CONTENTS' if toc else 'Introduction.',
    ]
    assert doc.paragraphs[title_index].style.name == 'Title'
    for paragraph in doc.paragraphs[title_index:title_index + 4]:
        _assert_non_outline(paragraph)
    assert texts.count('Body scope') == (2 if toc else 1)
    assert [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')] == ['Body scope']
    assert doc.element.xpath('.//w:instrText/text()') == (
        ['TOC \\o "1-4" \\h \\z \\u'] if toc else []
    )
    assert [n.get(qn('w:fldCharType')) for n in doc.element.xpath('.//w:fldChar')] == (
        ['begin', 'separate', 'end'] if toc else []
    )
    assert title_index == 0
    if toc:
        body_index = texts.index('Introduction.')
        assert doc.paragraphs[body_index - 1]._p.xpath('.//w:br[@w:type="page"]')
    assert model == original
    assert not renderer.errors


@pytest.mark.parametrize('toc', [False, True])
def test_matching_body_h1_is_not_duplicated_or_left_in_toc(toc: bool) -> None:
    model = MarkdownParser().parse(YAML_TITLE + '# Report title\n\n' + BODY)
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(
        include_cover=False, include_toc=toc, include_disclaimer=False, strict=True,
    )).render(model))
    assert [p.text for p in doc.paragraphs].count('Report title') == 1
    assert next(p for p in doc.paragraphs if p.text == 'Report title').style.name == 'Title'
    assert [p.text for p in doc.paragraphs if p.style.name == 'Heading 2'] == ['Body scope']
    # Only the private render copy loses the duplicate H1.
    assert any(isinstance(e.content, Heading) and e.content.text == 'Report title'
               for e in model.elements)


@pytest.mark.parametrize('title', ['', '   '])
def test_empty_title_does_not_create_an_empty_block(title: str) -> None:
    source = f'---\ntitle: "{title}"\nsubtitle: ""\nanalyst: Kim\n---\nBody text.'
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(
        include_cover=False, include_toc=False, include_disclaimer=False, strict=True,
    )).render(MarkdownParser().parse(source)))
    assert [p.text for p in doc.paragraphs] == ['Body text.']


def test_absent_subtitle_does_not_add_a_spacer() -> None:
    model = MarkdownParser().parse('---\ntitle: Report title\nanalyst: ""\n---\nBody text.')
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(preset='termsheet')).render(model))
    assert [p.text for p in doc.paragraphs] == ['Report title', 'Body text.']


@pytest.mark.parametrize('setting,toc', [
    ('layout: {cover: false, toc: false, disclaimer: false}', False),
    ('preset: termsheet', False), ('preset: legal-memo', True),
])
def test_frontmatter_cover_selection(setting: str, toc: bool) -> None:
    source = YAML_TITLE.replace('---\n', '---\n' + setting + '\n', 1) + BODY
    doc = _reopen(IBDocumentRenderer().render(MarkdownParser().parse(source)))
    texts = [p.text for p in doc.paragraphs]
    assert texts.count('Report title') == texts.count('Report subtitle') == 1
    assert ('TABLE OF CONTENTS' in texts) == toc


@pytest.mark.parametrize('flags,toc', [
    (['--no-cover'], True), (['--preset', 'termsheet'], False),
    (['--preset', 'legal-memo'], True),
])
def test_cli_cover_off_title_block(tmp_path: Path, monkeypatch, flags: List[str], toc: bool) -> None:
    source, target = tmp_path / 'report.md', tmp_path / 'report.docx'
    source.write_text(YAML_TITLE + BODY, encoding='utf-8')
    monkeypatch.setattr(sys, 'argv', ['md-to-word', str(source), str(target), '--strict'] + flags)
    with pytest.raises(SystemExit) as result:
        md_to_word.main()
    assert result.value.code == 0
    doc = Document(target)
    texts = [p.text for p in doc.paragraphs]
    assert texts.count('Report title') == texts.count('Report subtitle') == 1
    assert ('TABLE OF CONTENTS' in texts) == toc
    assert doc.sections[0].header.paragraphs[0].text
    assert [text.strip() for text in doc.sections[0].footer._element.xpath('.//w:instrText/text()')] == [
        'PAGE', 'NUMPAGES',
    ]


@pytest.mark.parametrize('flags', [['--no-cover'], ['--preset', 'legal-memo']],
                         ids=['no-cover', 'legal-memo'])
@pytest.mark.parametrize('source', [YAML_TITLE, INFERRED_TITLE], ids=['yaml', 'inferred'])
def test_cli_title_block_precedes_toc_on_first_page(
    tmp_path: Path, monkeypatch, flags: List[str], source: str,
) -> None:
    """Keep title metadata before the TOC and the existing break before body text."""
    input_path, output = tmp_path / 'order.md', tmp_path / 'order.docx'
    input_path.write_text(source + BODY, encoding='utf-8')
    monkeypatch.setattr(sys, 'argv', ['md-to-word', str(input_path), str(output), '--strict'] + flags)
    with pytest.raises(SystemExit) as result:
        md_to_word.main()
    assert result.value.code == 0
    doc = Document(output)
    paragraphs = doc.paragraphs
    texts = [p.text for p in paragraphs]
    assert texts[:5] == [
        'Report title', 'Report subtitle', '작성일: 2026-09-29', '작성자: Kim', 'TABLE OF CONTENTS',
    ]
    assert doc.element.body[0] is paragraphs[0]._p
    field_index = next(index for index, paragraph in enumerate(paragraphs)
                       if paragraph._p.xpath('.//w:instrText'))
    body_index = texts.index('Introduction.')
    assert 4 < field_index < body_index
    # No explicit page/section break or inherited page-break-before separates title and TOC.
    for paragraph in paragraphs[:field_index + 1]:
        assert not paragraph._p.xpath('.//w:br[@w:type="page"] | ./w:pPr/w:sectPr')
        assert not paragraph.paragraph_format.page_break_before
        style = paragraph.style
        while style is not None:
            assert not style.paragraph_format.page_break_before
            style = style.base_style
    assert paragraphs[body_index - 1]._p.xpath('.//w:br[@w:type="page"]')
    for paragraph in paragraphs[:4]:
        _assert_non_outline(paragraph)
    assert texts.count('Report title') == texts.count('Report subtitle') == 1
    assert texts.count('Body scope') == 2
    assert doc.element.xpath('.//w:instrText/text()') == ['TOC \\o "1-4" \\h \\z \\u']


@pytest.mark.parametrize('failure', ['Undefined footnote', 'Image'])
def test_strict_cover_off_still_rejects_loss_before_save(tmp_path: Path, monkeypatch, failure: str) -> None:
    source, target = tmp_path / 'bad.md', tmp_path / 'existing.docx'
    source.write_text(YAML_TITLE + ('Unknown[^9].' if failure == 'Undefined footnote'
                                    else '![Missing](absent.png)'), encoding='utf-8')
    target.write_bytes(b'Existing document')
    monkeypatch.setattr(sys, 'argv', ['md-to-word', str(source), str(target),
                                     '--preset', 'termsheet', '--strict'])
    with pytest.raises(SystemExit) as result:
        md_to_word.main()
    assert result.value.code != 0
    assert target.read_bytes() == b'Existing document'


def test_title_block_reuses_memo_typography_with_isolated_themes(tmp_path: Path) -> None:
    paths = []
    for color in ('123456', '654321'):
        path = tmp_path / (color + '.yaml')
        path.write_text(f'NAVY: "{color}"\nH1_SIZE: 23\nBODY_SIZE: 12\n'
                        'BODY_FONT: Arial\nKOREAN_FONT: Batang\n', encoding='utf-8')
        paths.append(path)

    def render_title(path: Path) -> DocxDocument:
        return _reopen(IBDocumentRenderer(options=RenderOptions(
            theme=str(path), preset='termsheet', strict=True,
        )).render(MarkdownParser().parse(YAML_TITLE + BODY)))

    with ThreadPoolExecutor(max_workers=2) as pool:
        docs = list(pool.map(render_title, paths * 2))
    for index, doc in enumerate(docs):
        title = doc.paragraphs[0]
        assert title.text == 'Report title'
        assert title.style.name == 'Title'
        assert str(title.style.font.color.rgb) == paths[index % 2].stem
        assert title.runs[0].font.size == Pt(23)
        assert title.runs[0].font.name == 'Arial'
        assert title.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia')) == 'Batang'
        assert doc.paragraphs[1].runs[0].font.size == Pt(12)
        memo = _reopen(IBDocumentRenderer(options=RenderOptions(
            theme=str(paths[index % 2]), preset='termsheet', strict=True,
        )).render(MarkdownParser(profile='ib-memo').parse(YAML_TITLE + BODY)))
        assert title._p.xml == memo.paragraphs[0]._p.xml
    assert STYLE.NAVY == IBStyle().NAVY


@pytest.mark.parametrize('profile,style', [
    ('ib-memo', 'Title'), ('plain', 'Heading 1'), ('office-letter', 'Office Subject'),
    ('business-report', 'Title'), ('meeting-minutes', 'Title'),
])
def test_existing_profile_body_titles_keep_their_styles(profile: str, style: str) -> None:
    model = parse_markdown_file(str(ROOT / 'samples/profiles' / (profile + '.md')))
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(strict=True)).render(model))
    matches = [p for p in doc.paragraphs if p.text == model.metadata.title
               or p.text == '제목\t' + model.metadata.title]
    assert len(matches) == 1
    assert matches[0].style.name == style


def test_cover_off_title_preserves_landscape_and_continuous_page_fields() -> None:
    model = parse_markdown_file(str(ROOT / 'samples/qa/landscape.md'), profile='ib-report')
    doc = _reopen(IBDocumentRenderer(options=RenderOptions(
        profile='ib-report', preset='termsheet', strict=True,
    )).render(model))
    assert doc.paragraphs[0].text == model.metadata.title
    assert doc.paragraphs[0].style.name == 'Title'
    assert [s.orientation for s in doc.sections] == [
        WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT,
    ]
    for section in doc.sections:
        assert not section._sectPr.xpath('./w:pgNumType[@w:start]')
        assert [text.strip() for text in section.footer._element.xpath('.//w:instrText/text()')] == [
            'PAGE', 'NUMPAGES',
        ]
